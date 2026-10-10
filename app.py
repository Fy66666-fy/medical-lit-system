import sys
import os
import io
import re
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import streamlit as st

from core import (appraisal, cache, cite, feedback, health, http, jobs, library, locate,
                  logger, mesh, pdfdoc, pubmed, quota, review, summarizer, storage, translate)
from version import APP_VERSION  # 版本单一来源（v2.5.0）：发版只需改 version.py

# ---- 运行日志与异常兜底（v2.4.0）----
logger.setup(APP_VERSION)   # 按天落盘；写入失败静默忽略，绝不拖垮主流程
logger.install_excepthook()  # 未捕获异常写入日志，便于事后定位

# ---- P1：会话级数据隔离 ----
# 云端所有访客共用同一进程，若不按会话分片，A 用户的检索记录/收藏会出现在 B 用户界面上。
# 桌面版由 desktop_app.py 设置 MEDLIT_SCOPE=local（沿用旧数据路径），云端按会话 id 分片。
def _init_scope() -> str:
    forced = os.environ.get("MEDLIT_SCOPE", "").strip()
    if forced:
        return forced
    try:
        from streamlit.runtime.scriptrunner import get_script_run_ctx

        ctx = get_script_run_ctx()
        sid = getattr(ctx, "session_id", "") if ctx else ""
        if sid:
            return "s" + str(sid).replace("-", "")[:12]
    except Exception:
        pass
    return "local"


storage.set_scope(_init_scope())

# ---- NCBI API Key（速率 3 → 10 req/s）----
# 优先级：侧边栏输入 > st.secrets > 环境变量 NCBI_API_KEY。
# 不配置也能用，只是并发检索时会更慢。
try:  # st.secrets 仅在存在 secrets 文件 / 云端配置时可用
    _ncbi_secret = (st.secrets.get("NCBI_API_KEY") or "").strip()
except Exception:
    _ncbi_secret = ""
if _ncbi_secret:
    http.configure_ncbi(_ncbi_secret)


def _diag_digest() -> str:
    """生成一段可读的诊断摘要：用户报错时附上它，能省掉大半来回追问。

    刻意**不含**检索内容、密钥与身份信息——只保留环境与运行指标。
    """
    hs = http.stats()
    cs = cache.stats()
    q = quota.limits()
    lines = [
        f"版本 {APP_VERSION} · 作用域 {storage.current_scope()}",
        f"外部请求 {hs['requests']} 次 / 重试 {hs['retries']} / 失败 {hs['failures']}"
        f" / 平均 {hs['avg_ms']}ms",
        f"NCBI 速率 {'%.0f req/s（已配 Key）' % (1 / http.ncbi_rate_limit()) if http.ncbi_api_key() else '约 3 req/s（未配 Key）'}",
        f"缓存 {cache.format_summary()}",
        f"依赖健康 {health.format_summary(7)}",
        f"翻译配额 会话 {q['trans_chars']['session']} 字符 / 每日 {q['trans_chars']['daily']} 字符",
        f"LLM 配额 会话 {q['llm_calls']['session']} 次 / 每日 {q['llm_calls']['daily']} 次",
    ]
    blk = quota.last_block()
    if blk.get("message"):
        lines.append(f"最近配额拦截（{blk.get('time')}）：{blk['message'][:120]}")
    return "\n".join(lines)


def _build_issue_text(kind: str, content: str) -> str:
    """留档时复制给用户的一段纯文本，可直接粘给作者。"""
    return "\n".join([
        f"[类型] {kind}",
        f"[描述] {content}",
        "[诊断]",
        _diag_digest(),
    ])


def _resolve_tencent_creds() -> tuple[str, str]:
    """腾讯云翻译凭据解析：侧边栏输入 > st.secrets > 环境变量"""
    sid = os.environ.get("TENCENT_SECRET_ID", "")
    skey = os.environ.get("TENCENT_SECRET_KEY", "")
    try:  # st.secrets 仅在存在 secrets 文件/云端配置时可用
        if st.secrets.get("TENCENT_SECRET_ID") and st.secrets.get("TENCENT_SECRET_KEY"):
            sid = st.secrets["TENCENT_SECRET_ID"]
            skey = st.secrets["TENCENT_SECRET_KEY"]
    except Exception:
        pass
    sid = st.session_state.get("tencent_sid") or sid
    skey = st.session_state.get("tencent_skey") or skey
    return (sid or "").strip(), (skey or "").strip()


_sid, _skey = _resolve_tencent_creds()
translate.configure(_sid, _skey)

# ---- 可选插件（streamlit 生态组件，缺失时自动降级为原生控件）----
try:
    from streamlit_option_menu import option_menu

    HAS_OPTION_MENU = True
except Exception:  # pragma: no cover - 插件可选
    HAS_OPTION_MENU = False

st.set_page_config(
    page_title="医学文献智能摘要与检索系统",
    page_icon="🩺",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---------------- 全局样式（简约清新主题） ----------------
st.markdown(
    """
    <style>
    /* ---------- 主题变量 ---------- */
    :root {
        --ink:  #14313c;
        --ink2: #5b7280;
        --ink3: #93a4ae;
        --teal: #2e9e8f;
        --teal-d: #217a6e;
        --teal-s: #eaf7f4;
        --sky:  #3d8fd1;
        --line: #e6f1ef;
    }
    .stApp { background: #f7fbfa; }
    .block-container { padding-top: 2rem; padding-bottom: 3rem; max-width: 1300px; }
    [data-testid="stSidebar"] { background: #ffffff; border-right: 1px solid var(--line); }
    [data-testid="stHeader"] { background: transparent; }

    /* ---------- 内页页头 ---------- */
    .page-head {
        background: linear-gradient(120deg, #f1fbf8 0%, #f4f9fd 100%);
        border: 1px solid #e3f1ee; border-radius: 16px;
        padding: 1rem 1.35rem; margin-bottom: 1rem;
    }
    .page-head h1 { font-size: 1.24rem; margin: 0; color: var(--ink); font-weight: 700; }
    .page-head p { margin: 0.28rem 0 0; color: var(--ink2); font-size: 0.87rem; }

    /* ---------- 首页 Hero（整体卡片，由容器 key 承载） ---------- */
    div[class*="st-key-hero"] {
        background: linear-gradient(125deg, #f0fbf8 0%, #f4f9fd 58%, #f7f5fd 100%);
        border: 1px solid #e3f1ee; border-radius: 24px;
        padding: 1.9rem 2rem 1.4rem;
        box-shadow: 0 6px 26px rgba(46, 158, 143, 0.06);
        margin-bottom: 1.15rem;
    }
    .hero-pill {
        display: inline-block; background: #ffffff; color: var(--teal);
        border: 1px solid #cdece5; border-radius: 999px;
        padding: 4px 14px; font-size: 0.78rem; font-weight: 600; letter-spacing: 0.3px;
    }
    .hero-wrap h1 {
        font-size: 1.98rem; line-height: 1.32; margin: 0.9rem 0 0.6rem;
        color: #12303a; font-weight: 800; letter-spacing: -0.3px;
    }
    .hero-sub { color: var(--ink2); font-size: 0.97rem; line-height: 1.8; margin: 0 0 1rem; max-width: 520px; }
    .hero-tags span {
        display: inline-block; background: rgba(255, 255, 255, 0.85); color: #4a6070;
        border: 1px solid #e0efec; border-radius: 999px;
        padding: 3px 11px; margin: 0 6px 6px 0; font-size: 0.775rem;
    }
    .hero-svg-box { display: flex; align-items: center; justify-content: center; }
    .hero-svg-box svg { max-width: 100%; height: auto; }

    /* ---------- 统计卡 ---------- */
    .stat-card {
        background: #ffffff; border: 1px solid var(--line); border-radius: 16px;
        padding: 0.9rem 1.1rem; height: 100%;
        box-shadow: 0 2px 10px rgba(46, 158, 143, 0.045);
    }
    .stat-card .k { color: var(--ink3); font-size: 0.78rem; letter-spacing: 0.3px; }
    .stat-card .v { color: #12303a; font-size: 1.58rem; font-weight: 700; line-height: 1.3; margin-top: 0.12rem; }
    .stat-card .u { color: var(--ink2); font-size: 0.75rem; }

    /* ---------- 功能卡 ---------- */
    div[class*="st-key-featcard"] {
        background: #ffffff; border: 1px solid var(--line); border-radius: 18px;
        padding: 1.15rem 1.15rem 0.8rem; height: 100%;
        box-shadow: 0 3px 14px rgba(46, 158, 143, 0.05);
        transition: transform 0.18s ease, box-shadow 0.18s ease, border-color 0.18s ease;
    }
    div[class*="st-key-featcard"]:hover {
        transform: translateY(-3px);
        box-shadow: 0 12px 28px rgba(46, 158, 143, 0.12);
        border-color: #c8ebe4;
    }
    .fc-icon {
        width: 40px; height: 40px; border-radius: 12px;
        display: flex; align-items: center; justify-content: center;
        font-size: 1.15rem; margin-bottom: 0.7rem;
    }
    .fc-title { font-size: 0.99rem; font-weight: 700; color: #12303a; margin-bottom: 0.3rem; }
    .fc-desc { font-size: 0.835rem; color: var(--ink2); line-height: 1.72; min-height: 4.3rem; }
    div[class*="st-key-featcard"] .stButton > button {
        width: 100%; font-size: 0.84rem; padding: 0.28rem 0; margin-top: 0.15rem;
    }

    /* ---------- 步骤条 ---------- */
    .step {
        background: #ffffff; border: 1px dashed #d9ece8; border-radius: 16px;
        padding: 0.95rem 1.1rem; height: 100%;
    }
    .step .num {
        display: inline-flex; align-items: center; justify-content: center;
        width: 26px; height: 26px; border-radius: 50%;
        background: var(--teal-s); color: var(--teal-d);
        font-size: 0.82rem; font-weight: 700; margin-bottom: 0.5rem;
    }
    .step .t { font-weight: 650; color: #12303a; font-size: 0.91rem; margin-bottom: 0.18rem; }
    .step .d { color: var(--ink2); font-size: 0.815rem; line-height: 1.7; }

    /* ---------- 区块标题 ---------- */
    .sec-title { display: flex; align-items: baseline; gap: 9px; margin: 1.15rem 0 0.8rem; }
    .sec-title .bar { width: 4px; height: 17px; border-radius: 3px; background: var(--teal); }
    .sec-title .tx { font-size: 1.05rem; font-weight: 700; color: #12303a; }
    .sec-title .hint { font-size: 0.79rem; color: var(--ink3); }

    /* ---------- 按钮 ---------- */
    .stButton > button, .stDownloadButton > button {
        border-radius: 11px; border: 1px solid #d9eae7; font-weight: 600;
        color: #33505c; transition: all 0.16s ease;
    }
    .stButton > button:hover, .stDownloadButton > button:hover {
        border-color: var(--teal); color: var(--teal-d);
        background: #f3fbf9; transform: translateY(-1px);
    }
    .stButton > button[kind="primary"],
    .stButton > button[data-testid="stBaseButton-primary"] {
        background: var(--teal); border-color: var(--teal); color: #ffffff;
        box-shadow: 0 4px 14px rgba(46, 158, 143, 0.24);
    }
    .stButton > button[kind="primary"]:hover,
    .stButton > button[data-testid="stBaseButton-primary"]:hover {
        background: var(--teal-d); border-color: var(--teal-d); color: #ffffff;
    }

    /* ---------- 表单控件协同（与主页同一套边框 / 圆角 / 主题色） ---------- */
    [data-testid="stTextInput"] input,
    [data-testid="stTextArea"] textarea {
        border-radius: 10px; border: 1px solid #d9eae7;
    }
    [data-testid="stTextInput"] input:focus,
    [data-testid="stTextArea"] textarea:focus {
        border-color: var(--teal); box-shadow: 0 0 0 2px rgba(46, 158, 143, 0.12);
    }
    [data-testid="stSelectbox"] > div > div,
    [data-testid="stSlider"] > div > div {
        border-radius: 10px;
    }
    div[data-testid="stExpander"] { background: #ffffff; }
    [data-testid="stAlert"] { border-radius: 12px; }

    /* 指标卡（全文数据分析面板的 st.metric）卡片化 */
    [data-testid="stMetric"] {
        background: #ffffff; border: 1px solid var(--line); border-radius: 14px;
        padding: 0.8rem 1rem;
        box-shadow: 0 2px 10px rgba(46, 158, 143, 0.045);
    }

    /* ---------- 结果面板（摘要 / 全文分析输出区） ---------- */
    div[class*="st-key-panel"] {
        background: #ffffff; border: 1px solid var(--line); border-radius: 18px;
        padding: 1.15rem 1.35rem;
        box-shadow: 0 3px 14px rgba(46, 158, 143, 0.05);
    }

    /* ---------- 输入区卡片（来源选择 / 参数设置） ---------- */
    div[class*="st-key-step-card"] {
        background: #fbfefd; border: 1px solid var(--line); border-radius: 18px;
        padding: 1.15rem 1.35rem;
    }
    /* 卡片内控件标题加粗，选项标题保持常规 */
    div[class*="st-key-step-card"] label[data-testid="stWidgetLabel"] p {
        font-weight: 600; margin-bottom: 0.15rem;
    }
    /* 参数子面板：语言 / 长度并排放进浅青底圆角块，与上方文献选择区分 */
    div[class*="st-key-step-params"] {
        background: #f2faf8; border: 1px solid #d9ece7; border-radius: 14px;
        padding: 0.85rem 1.1rem 0.35rem; margin-top: 0.9rem;
    }
    div[class*="st-key-step-params"] [data-testid="stRadio"] { margin-bottom: 0.45rem; }
    /* 卡片内控件组之间的呼吸感 */
    div[class*="st-key-step-card"] [data-testid="stVerticalBlock"] > div { margin-top: 0.55rem; }

    /* ---------- 其它 ---------- */
    div[data-testid="stExpander"] { border-radius: 12px; border-color: var(--line); }
    .kw-chip {
        display: inline-block; background: var(--teal-s); color: var(--teal-d);
        border-radius: 999px; padding: 2px 11px; margin: 2px 5px 2px 0; font-size: 0.79rem;
    }
    /* ---------- 原文定位（v2.1.0） ---------- */
    .loc-val {
        background: #fdeecd; color: #8a5a00; border-radius: 4px;
        padding: 0 4px; font-weight: 600;
    }
    .loc-sent {
        display: block; background: #f7fbfa; border: 1px solid var(--line);
        border-radius: 10px; padding: 0.55rem 0.85rem; margin: 0.35rem 0;
        line-height: 1.65; color: var(--ink); font-size: 0.86rem;
    }
    .loc-tag {
        color: var(--ink3); font-size: 0.76rem; margin-right: 0.5rem;
    }
    .loc-score {
        display: inline-block; background: #eaf7f4; color: var(--teal-d);
        border-radius: 999px; padding: 0 8px; font-size: 0.74rem; margin-left: 6px;
    }
    .loc-kw {
        background: #e3edfd; color: #1d4fb0; border-radius: 4px;
        padding: 0 4px; font-weight: 600;
    }
    /* ---------- 全局页脚：免责声明 / 数据来源 / 隐私（v2.4.0） ---------- */
    .app-footer {
        margin-top: 2.5rem; padding: 0.85rem 1.15rem;
        background: #ffffff; border: 1px solid var(--line); border-radius: 14px;
        font-size: 0.8rem; color: var(--ink2); line-height: 1.8;
    }
    .app-footer .ft-row + .ft-row {
        margin-top: 0.5rem; padding-top: 0.5rem; border-top: 1px dashed var(--line);
    }
    .app-footer b { color: var(--ink); }
    .app-footer .ft-warn { color: #a32d2d; }
    hr { border-color: var(--line); }
    </style>
    """,
    unsafe_allow_html=True,
)


def header(title: str = "🩺 医学文献智能摘要与检索系统", subtitle: str = "PubMed 检索 · 摘要长度可选 · 图表解析 · PDF 全文链接 · 收藏管理"):
    st.markdown(
        f"""
        <div class="page-head">
            <h1>{title}</h1>
            <p>{subtitle}</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

# ---------------- 页面路由 ----------------
NAV_ITEMS = [
    ("系统首页", "house"),
    ("文献检索", "search"),
    ("智能摘要", "file-earmark-text"),
    ("PDF 全文分析", "file-earmark-pdf"),
    ("综述工作台", "journal-text"),
    ("批量全文", "stack"),
    ("我的文献库", "star"),
    ("检索历史", "clock-history"),
    ("隐私与数据", "shield-lock"),
    ("更新日志", "clipboard-data"),
]
NAV_LABELS = [label for label, _ in NAV_ITEMS]

# 页面改名兼容（v3.1.0）：「我的收藏」升级为「我的文献库」（新增分组 / 标签 / 笔记）。
# 侧边栏显示新名，但已发出的深链（?page=我的收藏）与用户书签必须仍落到同一页，
# 否则改名会把这些链接一次性变成死链。此处集中映射，调用点在下方 pending 与深链解析。
PAGE_ALIAS = {"我的收藏": "我的文献库"}


def _norm_page(name):
    """把历史页面名归一到当前页面名；None 与未知名字原样返回。"""
    if not name:
        return name
    return PAGE_ALIAS.get(name, name)


def goto(target: str):
    """供首页卡片按钮跳转使用（回调方式；配合 manual_select 实现外部导航）"""
    st.session_state["pending_page"] = target


pending = _norm_page(st.session_state.pop("pending_page", None))
if pending is None:
    # 深链：?page=文献检索 可直接落到指定页（分享链接、落地页直达）
    try:
        _q = st.query_params.get("page")
        if isinstance(_q, list):
            _q = _q[0] if _q else None
        _q = _norm_page(_q)
        if _q and _q in NAV_LABELS:
            pending = _q
            # 记下深链目标：首次进入会先撞到隐私同意门，勾选后要回到这个页面而不是首页
            st.session_state["deeplink_page"] = _q
    except Exception:
        pass
pending_idx = NAV_LABELS.index(pending) if pending in NAV_LABELS else None

# ---------------- 后台任务（v1.7.0）：全文摘要 / 图表解析 / 批量全文 ----------------
def _ft_job(job, article: dict, title: str, max_sentences: int):
    """后台任务：单篇全文抓取（PMC → 浏览器兜底）+ 章节化摘要 + 数据分析"""
    jobs.update(job, 0.08, "正在抓取全文（PMC → 网页兜底）…")
    with logger.span("全文抓取", pmcid=article.get("pmcid", ""), title=title[:40]):
        secs, src = pubmed.fetch_fulltext_any(article)
    jobs.update(job, 0.55, "正在生成章节化摘要与数据分析…")
    with logger.span("章节化摘要", title=title[:40]):
        summary = summarizer.fulltext_summary(secs, title=title, max_sentences=max_sentences)
    analytics = summarizer.analyze_fulltext(secs)
    logger.info(f"全文摘要完成：{len(secs)} 章节 · 来源 {src}")
    return {"secs": secs, "summary": summary, "analytics": analytics,
            "pmcid": article.get("pmcid", ""), "title": title, "source": src}


def _fig_job(job, pmcid: str):
    """后台任务：图表抓取 + 图片下载（非 OA 走浏览器截图兜底）"""
    jobs.update(job, 0.15, "正在从 PMC 抓取图表…")
    with logger.span("图表元数据抓取", pmcid=pmcid):
        figures = pubmed.fetch_pmc_figures(pmcid)
    if figures:
        jobs.update(job, 0.5, "正在获取图表图片（非开放获取文献走浏览器截图兜底，约需 20-60 秒）…")
        with logger.span("图表图片获取", pmcid=pmcid, 图表数=len(figures)):
            pubmed.fetch_figure_images(pmcid, figures)
    got = [f for f in figures if f.get("data")]
    logger.info(f"图表解析完成：{len(got)}/{len(figures)} 张拿到图片 (pmcid={pmcid})")
    return got


def _batch_job(job, articles: list[dict], max_sentences: int, workers: int, lang: str):
    """后台任务：多篇文献并行全文抓取 + 摘要 +（可选）中文翻译"""
    from concurrent.futures import ThreadPoolExecutor

    total = len(articles)
    results: list[dict | None] = [None] * total

    def one(i: int, a: dict):
        try:
            with logger.span("批量·单篇", 序号=i, 标题=a["title"][:40]):
                secs, src = pubmed.fetch_fulltext_any(a)
                summary = summarizer.fulltext_summary(secs, title=a["title"], max_sentences=max_sentences)
                analytics = summarizer.analyze_fulltext(secs)
            results[i] = {"article": a, "ok": True, "summary": summary, "analytics": analytics, "source": src}
        except Exception as e:  # noqa: BLE001 —— 单篇失败不拖垮整批
            logger.error(f"批量·单篇失败 序号={i}", e)
            results[i] = {"article": a, "ok": False, "error": str(e)}
        done = sum(1 for r in results if r is not None)
        jobs.update(job, progress=done / total * 0.85, note=f"{done}/{total} 篇完成")

    logger.info(f"批量任务开始：{total} 篇 · {workers} 线程 · 输出 {lang}")
    with ThreadPoolExecutor(max_workers=workers) as ex:
        for f in [ex.submit(one, i, a) for i, a in enumerate(articles)]:
            f.result()

    if lang == "中文":
        for i, r in enumerate(results):
            if not (r and r["ok"]):
                continue
            jobs.update(job, progress=0.85 + 0.15 * i / max(total, 1), note=f"翻译为中文：{i + 1}/{total}…")
            for p in r["summary"].get("sections", []):
                body = " ".join(p["sentences"])
                if summarizer.is_mostly_english(body):
                    try:
                        body = summarizer.translate_text(body, "en|zh-CN")
                    except Exception:
                        pass
                p["body_zh"] = body
    jobs.update(job, 1.0, f"完成（{sum(1 for r in results if r and r['ok'])}/{total} 篇成功）")
    return {"items": results, "lang": lang, "max_sentences": max_sentences}


def _apply_job_result(job: dict):
    """UI 主线程：把已完成任务的结果写入 session_state 并触发整页刷新"""
    if job["status"] == "error":
        logger.error(f"后台任务失败：{job['name']} | {job.get('error') or ''}")
        st.toast(f"❌ {job['name']} 失败：{(job['error'] or '')[:80]}", icon="❌")
        if job.get("key", "").startswith("fig:"):
            st.session_state["figures"] = []
            st.session_state["fig_error"] = job.get("error") or "未知错误"
        return
    key = job.get("key", "")
    res = job.get("result") or {}
    if key.startswith("ft:"):
        st.session_state["fulltext"] = res["secs"]
        st.session_state["ft_summary"] = res["summary"]
        st.session_state["ft_analytics"] = res["analytics"]
        st.session_state["ft_source"] = res.get("source", "")
        st.toast("✅ 全文摘要已生成", icon="✅")
    elif key.startswith("fig:"):
        st.session_state["figures"] = res
        st.toast(f"✅ 图表解析完成（{len(res)} 张）", icon="✅")
    elif key == "batch":
        st.session_state["batch_results"] = res["items"]
        st.session_state["batch_lang"] = res.get("lang", "中文")
        ok = sum(1 for r in res["items"] if r.get("ok"))
        st.toast(f"✅ 批量全文完成：{ok}/{len(res['items'])} 篇成功", icon="✅")


def _batch_export_md(items: list[dict], lang: str) -> str:
    """把批量全文结果（摘要 + 数据分析）组装成一份 Markdown"""
    from datetime import datetime

    ok_n = sum(1 for r in items if r.get("ok"))
    lines = [
        "# 医学文献批量全文摘要与数据分析",
        "",
        f"> 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')} · 共 {len(items)} 篇（成功 {ok_n}，失败 {len(items) - ok_n}）",
        "",
    ]
    for i, r in enumerate(items, 1):
        a = r["article"]
        lines += [f"## {i}. {a['title']}", ""]
        meta = [", ".join(a.get("authors") or []) if a.get("authors") else "",
                a.get("journal") or "", a.get("year") or ""]
        meta = " · ".join(x for x in meta if x)
        if meta:
            lines += [f"{meta} · PMID: {a.get('pmid', '-')} · {a.get('pmcid', '')}", ""]
        if not r.get("ok"):
            lines += [f"> ❌ 抓取失败：{r.get('error', '')}", "", "---", ""]
            continue
        if str(r.get("source", "")).startswith("网页"):
            lines += ["> 🌐 全文来源：网页提取（浏览器兜底，非 PMC 开放存档）", ""]
        summ, an = r["summary"], r["analytics"]
        if summ.get("sections"):
            for p in summ["sections"]:
                if lang == "中文":
                    body = p.get("body_zh") or " ".join(p["sentences"])
                    lines += [f"**【{p.get('title_zh') or p['title']}】** {body}", ""]
                else:
                    lines += [f"**【{p['title']}】** " + " ".join(p["sentences"]), ""]
        elif summ.get("summary"):
            lines += [summ["summary"], ""]
        lines += ["### 📊 数据分析", ""]
        lines.append(
            f"- 全文词数：{an['total_words']:,}；字符数：{an['total_chars']:,}；"
            f"章节数：{len(an['section_stats'])}"
        )
        kc = an.get("keyword_counts", {})
        kws = an.get("keywords", [])[:10]
        if kws:
            lines.append("- 高频关键词：" + "、".join(f"{k}（{kc[k]}）" if k in kc else k for k in kws))
        for name, m in (an.get("metrics") or {}).items():
            lines.append(f"- {name}（{m['count']} 处）：" + "；".join(m["samples"][:5]))
        lines += ["", "---", ""]
    return "\n".join(lines)


def _batch_export_xlsx(items: list[dict], lang: str = "中文") -> bytes:
    """批量结果统计表：Excel 三工作表（文献汇总 / 章节明细 / 统计指标），带表头样式与筛选"""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F6E54")
    head_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(vertical="top", wrap_text=True)

    def _style(ws, widths: dict, freeze: str = "A2"):
        for c, wdt in widths.items():
            ws.column_dimensions[c].width = wdt
        for cell in ws[1]:
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = freeze
        ws.auto_filter.ref = ws.dimensions

    # ---------- Sheet1 文献汇总 ----------
    ws = wb.active
    ws.title = "文献汇总"
    ws.append(["序号", "标题", "作者", "期刊", "年份", "DOI", "PMID", "PMCID",
               "PubMed 链接", "全文来源", "状态", "摘要章节数", "摘要全文",
               "全文总词数", "全文字符数", "正文章节数", "高频关键词 Top10",
               "统计指标汇总", "失败原因"])
    for i, r in enumerate(items, 1):
        a = r["article"]
        authors = ", ".join(a.get("authors") or [])
        pmid = a.get("pmid") or ""
        if r.get("ok"):
            an, summ = r["analytics"], r["summary"]
            kc = an.get("keyword_counts", {})
            kws = "、".join(f"{k}({kc[k]})" if k in kc else k for k in an.get("keywords", [])[:10])
            msum = "；".join(f"{name}×{m['count']}" for name, m in (an.get("metrics") or {}).items()) or "未检出"
            if summ.get("sections"):
                parts = []
                for p in summ["sections"]:
                    t = (p.get("title_zh") or p["title"]) if lang == "中文" else p["title"]
                    body = (p.get("body_zh") or " ".join(p["sentences"])) if lang == "中文" else " ".join(p["sentences"])
                    parts.append(f"【{t}】{body}")
                preview = " ".join(parts)
                n_secs = len(summ["sections"])
            else:
                preview, n_secs = summ.get("summary", ""), 0
            # 摘要完整导出（仅防 Excel 单元格 32767 字符上限做保护性截断）
            if len(preview) > 32000:
                preview = preview[:32000] + "…"
            ws.append([i, a.get("title", ""), authors, a.get("journal", ""), a.get("year", ""),
                       a.get("doi", ""), pmid, a.get("pmcid", ""),
                       f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/" if pmid else "",
                       r.get("source", ""), "成功", n_secs, preview,
                       an.get("total_words", ""), an.get("total_chars", ""),
                       len(an.get("section_stats") or []), kws, msum, ""])
            ws.cell(ws.max_row, 11).font = Font(color="1F6E54", bold=True)
        else:
            ws.append([i, a.get("title", ""), authors, a.get("journal", ""), a.get("year", ""),
                       a.get("doi", ""), pmid, a.get("pmcid", ""), "", "", "失败",
                       "", "", "", "", "", "", "", r.get("error", "")])
            ws.cell(ws.max_row, 11).font = Font(color="C0392B", bold=True)
    _style(ws, {"A": 6, "B": 45, "C": 28, "D": 20, "E": 8, "F": 22, "G": 13, "H": 13,
                "I": 34, "J": 12, "K": 8, "L": 11, "M": 60, "N": 11, "O": 11, "P": 11,
                "Q": 40, "R": 30, "S": 30})
    for row in ws.iter_rows(min_row=2):
        for cell in (row[1], row[12], row[16]):  # 标题/预览/关键词列换行
            cell.alignment = wrap

    # ---------- Sheet2 章节明细 ----------
    ws2 = wb.create_sheet("章节明细")
    ws2.append(["文献序号", "文献标题", "章节", "章节词数", "章节字符数"])
    for i, r in enumerate(items, 1):
        if not r.get("ok"):
            continue
        t = r["article"].get("title", "")
        for st in r["analytics"].get("section_stats", []):
            ws2.append([i, t, st.get("title", ""), st.get("words", ""), st.get("chars", "")])
    _style(ws2, {"A": 9, "B": 45, "C": 22, "D": 11, "E": 12})

    # ---------- Sheet3 统计指标 ----------
    ws3 = wb.create_sheet("统计指标")
    ws3.append(["文献序号", "文献标题", "指标", "出现次数", "示例值"])
    for i, r in enumerate(items, 1):
        if not r.get("ok"):
            continue
        t = r["article"].get("title", "")
        for name, m in (r["analytics"].get("metrics") or {}).items():
            ws3.append([i, t, name, m.get("count", ""), "；".join(m.get("samples", [])[:6])])
    _style(ws3, {"A": 9, "B": 45, "C": 18, "D": 10, "E": 70})
    for row in ws3.iter_rows(min_row=2):
        row[4].alignment = wrap

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@st.fragment(run_every=1.5)
def job_monitor():
    """侧边栏任务进度监视：轮询后台任务，完成后消费结果并整页刷新"""
    snaps = jobs.snapshot()
    running = [j for j in snaps if j["status"] == "running"]
    if not running and not any(not j["consumed"] for j in snaps):
        return
    changed = False
    for j in snaps:
        if not j["consumed"] and j["status"] in ("done", "error"):
            full = jobs.consume(j["id"])
            if full:
                _apply_job_result(full)
                changed = True
    if running:
        st.markdown("⏳ **后台任务**")
        for j in running:
            st.caption(f"{j['name']} · {j['note']}")
            st.progress(max(j["progress"], 0.03))
    if changed:
        st.rerun(scope="app")


# ---------------- 使用体验问卷（v3.7.0：自愿参与，答案只存本机） ----------------
# 项目没有账号体系也没有服务器，做不了（也不该做）静默遥测。这里的取舍是：
# 用满一定次数后在侧边栏放一个**不弹窗的邀请**（用户点「填写」才打开模态框），
# 答案只落 data/feedback/，用户自己决定要不要通过 GitHub 按钮发给开发者。
# 计数、邀请节奏与上限全部由 core/feedback.py 的纯函数控制，可离线断言。
@st.dialog("📝 使用体验问卷（约 1 分钟，自愿）", width="large")
def _survey_dialog():
    st.caption(
        "先说清楚数据去向：这个工具没有账号体系，我们也不知道你是谁——"
        "问卷**只存在你这台设备的 data/feedback/ 目录里**，不会自动上传。"
        "填完后可以通过 GitHub 一键发给我们，或者直接关掉，都不影响任何功能。"
    )
    with st.form("survey_form", clear_on_submit=False):
        freq = st.radio("你大概多久用一次？",
                        ["刚试用", "每周 1–2 次", "每周 3 次以上", "几乎每天"],
                        horizontal=True, key="sv_freq")
        useful = st.multiselect("哪些功能对你最有用？（可多选）",
                                list(feedback.USEFUL_OPTIONS), key="sv_useful")
        sat = st.radio("总体满意度",
                       ["⭐⭐⭐⭐⭐ 好用到想推荐", "⭐⭐⭐⭐ 不错，有小问题", "⭐⭐⭐ 一般",
                        "⭐⭐ 不太顺", "⭐ 有待大改"], key="sv_sat")
        improve = st.text_area("最想改进的一点（可选）", height=80, key="sv_improve",
                               placeholder="哪个环节卡住了、缺什么功能、哪里看不懂……")
        rec = st.radio("愿意推荐给同学 / 同事吗？", ["会", "说不定", "暂时不会"],
                       horizontal=True, key="sv_rec")
        contact = st.text_input("联系方式（可选，仅当你愿意被回访时填写）", key="sv_contact")
        submitted = st.form_submit_button("提交问卷（只存本机）", use_container_width=True)
    if submitted:
        answers = {
            "freq": freq,
            "useful": "、".join(useful) or "（未选）",
            "satisfaction": sat,
            "improve": improve,
            "recommend": rec,
            "contact": contact,
        }
        u = feedback.load_usage()
        path = feedback.save_survey(answers, u.get("count", 0))
        if path:
            feedback.complete_survey(u)
            st.session_state["fb_survey_path"] = path
            st.session_state["fb_survey_text"] = feedback.render_survey_text(
                answers, u.get("count", 0))
            st.session_state["sv_submitted"] = True
        else:
            st.error("保存失败（磁盘或权限问题），内容没有提交，可稍后重试。")
    if st.session_state.get("sv_submitted"):
        st.success("已保存到本机！感谢反馈 🙏")
        st.caption(f"文件位置：{st.session_state.get('fb_survey_path', '')}")
        _sv_text = st.session_state.get("fb_survey_text", "")
        st.code(_sv_text, language="markdown")
        st.link_button("📮 一键通过 GitHub 发给开发者（推荐）",
                       feedback.issue_url("feature", _sv_text))
        if st.button("关闭", use_container_width=True):
            for k in ("fb_dialog", "sv_submitted", "fb_survey_path", "fb_survey_text"):
                st.session_state.pop(k, None)
            st.rerun(scope="app")


# ---------------- 侧边栏 ----------------
# 叙述段可选模型（P3-C6）：只是**常用 OpenAI 兼容模型名的快捷预设**，不是白名单——
# 选了预设但该模型不在用户自己配置的 API 地址下，接口会直接报错。
# 因此下拉默认停留在「跟随侧边栏设置」，由用户对「地址 × 模型」是否配套负责。
#
# 这份清单是**兜底**：侧边栏可点「拉取接口的模型列表」从 /v1/models 取真实列表，
# 拿到就用真的、拿不到才回落这里。写死清单必然随上游发版而过期，动态拉取才是正解。
LLM_MODEL_PRESETS: tuple[str, ...] = (
    "gpt-5.6-sol", "gpt-5.5", "gpt-5-mini",          # OpenAI
    "deepseek-v4-pro", "deepseek-v4.1-flash", "deepseek-flash",  # DeepSeek
    "qwen3.8-max", "qwen3.8-flash",                  # 阿里 Qwen
    "glm-5.2", "glm-4.7",                            # 智谱 GLM
    "kimi-k3",                                       # 月之暗面 Kimi
    "claude-sonnet-5", "claude-haiku-4.5",           # Anthropic
    "gemini-3.5-flash",                              # Google
)
FOLLOW_SIDEBAR = "跟随侧边栏设置"  # 模型下拉的第一项（哨兵值，勿与真实模型名重复）


def _llm_model_choices() -> list[str]:
    """叙述段模型下拉的候选项：优先用从接口拉到的真实列表，否则回落内置预设。"""
    got = st.session_state.get("llm_models") or []
    return [str(m) for m in got] if got else list(LLM_MODEL_PRESETS)
with st.sidebar:
    st.markdown("### 🧭 功能导航")
    if HAS_OPTION_MENU:
        page = option_menu(
            None,
            NAV_LABELS,
            icons=[icon for _, icon in NAV_ITEMS],
            default_index=0,
            manual_select=pending_idx,
            key="nav",
            styles={
                "container": {"padding": "0", "background-color": "transparent"},
                "icon": {"color": "#2e9e8f", "font-size": "0.95rem"},
                "nav-link": {
                    "font-size": "0.93rem",
                    "border-radius": "10px",
                    "padding": "8px 12px",
                    "margin": "2px 0",
                    "color": "#4a6070",
                    "--hover-color": "#eaf7f4",
                },
                "nav-link-selected": {
                    "background-color": "#2e9e8f",
                    "color": "#ffffff",
                    "font-weight": "600",
                },
            },
        )
        if page not in NAV_LABELS:  # 兜底，防止状态异常
            page = NAV_LABELS[0]
    else:
        if pending is not None:
            st.session_state["nav"] = pending
        page = st.radio("页面", NAV_LABELS, label_visibility="collapsed", key="nav")
    st.divider()
    job_monitor()  # 后台任务进度（全文摘要 / 图表解析 / 批量全文）
    st.divider()
    st.markdown("### 🤖 大模型设置（可选）")
    st.caption("配置 OpenAI 兼容接口后，智能摘要页可使用 LLM 生成深度总结；不配置也能使用内置抽取式摘要。")

    llm_base = st.text_input(
        "API Base URL",
        value=st.session_state.get("llm_base", ""),
        placeholder="https://api.openai.com",
    )
    llm_key = st.text_input(
        "API Key",
        value=st.session_state.get("llm_key", ""),
        type="password",
    )
    llm_model = st.text_input(
        "模型名称",
        value=st.session_state.get("llm_model", ""),
        placeholder="deepseek-v4-pro",
    )
    st.session_state["llm_base"] = llm_base
    st.session_state["llm_key"] = llm_key
    st.session_state["llm_model"] = llm_model
    llm_ready = bool(llm_base and llm_key and llm_model)
    if llm_ready:
        st.caption("✅ LLM 已就绪")
    else:
        st.caption("⚪ 未配置，仅使用抽取式摘要")

    # 模型列表动态获取：硬编码的模型名清单必然随上游发版而过期，
    # 改为直接问用户的接口要真实列表；失败不影响手填模型名，也不影响其它功能。
    if st.button("🔄 拉取接口的模型列表", use_container_width=True,
                 disabled=not (llm_base and llm_key),
                 help="从上面 API Base URL 的 /v1/models 读取真实可用模型，供叙述段下拉选用"):
        try:
            _ms = summarizer.list_models(llm_base, llm_key)
            st.session_state["llm_models"] = _ms
            st.success(f"已获取 {len(_ms)} 个模型")
        except Exception as e:
            logger.error("拉取模型列表失败", e)
            st.warning(f"拉取失败：{e}。可继续手填模型名，叙述段下拉仍用内置预设。")
    _got_models = st.session_state.get("llm_models") or []
    if _got_models:
        st.caption(f"已获取 {len(_got_models)} 个模型，叙述段的模型下拉将使用这份真实列表。")

    st.divider()
    st.markdown("### 🌐 翻译设置（可选）")
    st.caption(
        "配置腾讯云机器翻译（每月 500 万字符免费额度）后，中文翻译走腾讯接口，"
        "质量与额度大幅优于默认的 MyMemory 免费接口；未配置或调用失败时自动回退 MyMemory。"
    )
    t_sid = st.text_input(
        "腾讯云 SecretId",
        value=st.session_state.get("tencent_sid", ""),
        placeholder="AKID********",
    )
    t_skey = st.text_input(
        "腾讯云 SecretKey",
        value=st.session_state.get("tencent_skey", ""),
        type="password",
    )
    st.session_state["tencent_sid"] = t_sid
    st.session_state["tencent_skey"] = t_skey
    if t_sid and t_skey:
        st.caption("✅ 腾讯云翻译已就绪")
    elif os.environ.get("TENCENT_SECRET_ID"):
        st.caption("✅ 检测到环境变量凭据")
    else:
        st.caption("⚪ 未配置，使用 MyMemory 免费接口")

    st.divider()
    st.markdown("### ⚡ NCBI 速率（可选）")
    st.caption(
        "NCBI 对每个 IP 的检索限速：无 API Key 约 3 次/秒，申请免费 Key 后约 10 次/秒，"
        "多人同时检索时响应差别很明显。免费申请："
        "https://account.ncbi.nlm.nih.gov/settings/#page=api_keys（填一个邮箱即可，秒批）"
    )
    n_key = st.text_input(
        "NCBI API Key",
        value=st.session_state.get("ncbi_key", ""),
        placeholder="留空则用 3 次/秒的公共额度",
        type="password",
    )
    if n_key != st.session_state.get("ncbi_key_applied", ""):
        http.configure_ncbi(n_key)
        st.session_state["ncbi_key"] = n_key
        st.session_state["ncbi_key_applied"] = n_key
        st.rerun()
    if http.ncbi_api_key():
        st.caption(f"✅ 已启用 API Key · 当前约 {1 / http.ncbi_rate_limit():.0f} 次/秒")
    else:
        st.caption("⚪ 未配置，约 3 次/秒（多人同时检索时会排队）")

    st.divider()
    st.caption("数据源：PubMed E-utilities（NCBI 官方公开 API）\n\n检索结果仅用于研究学习，不构成医疗建议。")

    # ---- 运行诊断（v2.4.0）：出问题时展开这里，截图即可定位 ----
    with st.expander("🩺 运行诊断", expanded=False):
        st.caption(f"版本 {APP_VERSION} · 日志目录 `{logger.LOG_DIR}`（保留 {logger._KEEP_DAYS} 天）")
        _hs = http.stats()
        st.caption(
            f"外部请求：{_hs['requests']} 次 · 重试 {_hs['retries']} 次 · 失败 {_hs['failures']} 次"
            f" · 平均 {_hs['avg_ms']}ms · 下行 {_hs['mb']}MB"
        )
        # P1：依赖成功率与配额（判断是否被限流 / 额度用尽）
        st.caption("📶 " + health.format_summary(7))
        _blk = quota.last_block()
        if _blk.get("message"):
            st.warning(f"配额拦截（{_blk['time']}）：{_blk['message']}")
        else:
            _rem_c = quota.remaining("trans_chars", storage.current_scope(),
                                     own_key=translate.has_tencent())
            if _rem_c is None:
                st.caption("💳 翻译额度：使用你自己的腾讯云密钥，不受宿主额度限制")
            else:
                st.caption(f"💳 本次可用翻译额度：{_rem_c} 字符（自带密钥则不受限）")
        st.caption(
            f"🗂 数据作用域 `{storage.current_scope()}`"
            + ("（云端按会话隔离，刷新或换设备将看不到本次收藏）"
               if storage.current_scope() != "local" else "（本地单机，数据长期保存）")
        )
        # P1 任务 4：缓存命中率。命中率低说明上游被反复打，是额度耗尽的前兆
        st.caption("⚡ 缓存：" + cache.format_summary())
        _c1, _c2 = st.columns(2)
        if _c1.button("🧹 清空缓存", key="cache_clear", use_container_width=True):
            n = cache.clear()
            st.toast(f"已清空 {n} 条缓存", icon="🧹")
            st.rerun()
        if _c2.button("🔄 刷新日志", key="diag_refresh", use_container_width=True):
            st.rerun()
        st.code(logger.tail(60), language="text")
        st.caption("提示：报错时展开此处截图发给开发者，可快速定位问题。")

    # ---- 反馈渠道（P1 任务 6）：本地留档 + 一键生成预填的 Issue ----
    with st.expander("💬 报错 / 建议", expanded=False):
        st.caption("填完点「留档并复制」即可：内容会存在本机，同时复制一段可直接粘贴的文本。")
        _fk = st.selectbox(
            "类型", list(feedback.KINDS),
            format_func=lambda k: {"bug": "🐛 报错", "feature": "💡 建议",
                                    "question": "❓ 疑问", "other": "📝 其他"}[k],
            key="fb_kind",
        )
        _fc = st.text_area("描述", placeholder="哪个功能、什么操作、期望是什么、实际发生了什么",
                           height=90, key="fb_content")
        with st.expander("附带诊断信息（推荐，省去来回追问）", expanded=False):
            st.code(_diag_digest(), language="text")
            st.caption("含版本、请求统计、依赖成功率、配额与缓存状态；不含你的检索内容。")
        _f1, _f2 = st.columns(2)
        if _f1.button("📥 留档并复制", key="fb_save", use_container_width=True):
            _clean = feedback.sanitize(_fc)
            if not _clean.strip():
                st.warning("请先填写描述")
            elif feedback.save(_fk, _clean, _diag_digest()):
                st.session_state["fb_copied"] = _build_issue_text(_fk, _clean)
                st.success(f"已留档 · {feedback.summary_line()}")
        if _f2.button("🔗 生成 Issue 链接", key="fb_issue", use_container_width=True):
            _clean = feedback.sanitize(_fc)
            if not _clean.strip():
                st.warning("请先填写描述")
            else:
                st.link_button("在 GitHub 提交（内容已预填）",
                               feedback.issue_url(_fk, _clean, _diag_digest()))
        if st.session_state.get("fb_copied"):
            st.code(st.session_state["fb_copied"], language="text")
            st.caption("↑ 已复制到剪贴板，可直接粘贴给作者。")
        st.caption(f"本地反馈留档：{feedback.summary_line()}（桌面版可把 data/feedback/ 整个目录发给作者）")

    # ---- 使用体验问卷邀请（v3.7.0）：到次数才出现，自愿参与，绝不自动弹窗 ----
    _fb_u = feedback.load_usage()
    if feedback.should_prompt(_fb_u) and not st.session_state.get("fb_dialog"):
        with st.container(border=True):
            st.markdown(f"📝 **你已经用了 {_fb_u.get('count', 0)} 次**")
            st.caption("愿意花 1 分钟告诉我们用得怎么样吗？问卷自愿填写、"
                       "只存本机 data/feedback/，不会上传任何数据。")
            _q1, _q2, _q3 = st.columns(3)
            if _q1.button("填写", key="fb_sv_yes", use_container_width=True, type="primary"):
                st.session_state["fb_dialog"] = True
                feedback.record_prompt(feedback.load_usage())
                st.rerun()
            if _q2.button("稍后", key="fb_sv_later", use_container_width=True):
                feedback.record_prompt(feedback.load_usage())
                st.toast("好的，之后再问你", icon="👌")
                st.rerun()
            if _q3.button("不再提醒", key="fb_sv_mute", use_container_width=True):
                feedback.mute_prompts(feedback.load_usage())
                st.rerun()
    if st.session_state.get("fb_dialog"):
        _survey_dialog()


# ---------------- 工具函数 ----------------
CHANGELOG = [
    {
        "version": "v4.0.0",
        "date": "2026-10-10",
        "tag": "最新版本",
        "items": [
            ('📖', '字段抽取优先读全文，无全文再回落摘要', '综述工作台的字段抽取过去只读 PubMed 摘要，上传的 PDF 也被压成摘要，抓回来的开放全文从没用上。现统一为单一入口：有全文（PMC 开放全文、本地上传 PDF）就用全文，没有则逐字回落到摘要，无全文时的结果与旧版完全一致。对比表新增「数据来源」列，标注每个字段抽自全文还是摘要。'),
            ('📥', '新增「用全文抽取」入口（可选，不自动抓取）', '工作台新增「用全文抽取」折叠区，点击「抓取所选全文」即走PMC、Europe PMC 与发布商网页三条通道，结果缓存 7 天，同一篇不重复抓取。PubMed 中约两三成为开放获取，付费墙的抓不到，会自动回落摘要并标「未抓到」，对比与骨架不受影响。'),
            ('🔬', '研究类型识别接入 PubMed 权威标引', '过去只靠标题 / 摘要里的措辞猜研究设计，而大量摘要根本不写「随机」「队列」这些字眼，于是大片「未识别」。抓取时其实已带回 NLM 标定的文献类型（如 Randomized Controlled Trial / Cohort Studies），这一版把它作为最高优先级依据，命中处写明「PubMed 文献类型『xxx』」可溯源；措辞线索也一并扩充兜底。'),
            ('📊', '「看证据」页显示研究类型识别率', '新增研究类型识别率与未识别篇目清单（附 PubMed 文献类型原值），改写后未识别率的降幅可在页面直接检验，无需再猜。'),
            ('🧩', '长综述超过 50 篇自动分批生成', '篇数一大，一次生成「结果概述」容易因输出过长被截断。现在超过 50 篇会自动分批（每批 40 篇）并发生成「结果概述」，再基于全部结果统一写一段「讨论」；50 篇以内与旧版逐字一致。界面显示真实进度，不再使用写死的「约 10–40 秒」。'),
            ('⚡', '全文抓取与字段补抽改为并发，并给出覆盖率报告', '「抓取所选全文」与「大模型补抽」过去是一篇篇排队做，几百篇要等很久。现在都改成受控并发（3 路），限速仍按域名严格执行；抓完给出覆盖率报告（尝试 / 成功篇数、来源构成、当前覆盖率、付费墙指引）。'),
            ('📎', '导出「骨架 + 叙述段」时，叙述段嵌进骨架里', '过去叙述段以「附：」挂在骨架外面。现在「结果概述」落到骨架的结果章末尾、「讨论」落到讨论章末尾；标题写法宽松识别，切不开时整体并入讨论章，内容不会丢失。'),
            ('🧠', '抽取有缺口时给出明确补法', '对比表下方统计「有痕迹却没抽到」与「原文确实没写」的格子数，并给出两条补法：在表格里直接手填，或用大模型补抽（附原文引句）；来源仍是摘要且有空缺时，直接指向「用全文抽取」。'),
        ],
    },
    {
        "version": "v3.9.0",
        "date": "2026-10-10",
        "tag": "",
        "items": [
            ('🩹', '修复「结果概述」只吐出一句话', '使用带思考过程的推理模型时，思考 token 也计入输出额度，额度给小了正文会被截断成半句，且这半句还会进入缓存，反复点「撰写叙述段」命中的始终是同一句。现会识别截断（finish_reason=length）并自动加大额度重试；重试被上游拒绝时沿用已拿到的正文；仅当正文彻底为空才报可操作的中文错误。被截断的结果不写入缓存。'),
            ('📈', '叙述段与补抽的额度基线抬高', '叙述段默认额度 2400（严谨）/ 1400（简明）→ 6000 / 3000，字段补抽 2000 → 4000，避免首次调用即因额度不足被截断、浪费一次 token。'),
            ('🗂️', '模型下拉改为拉取接口真实列表', '模型下拉过去是写死的清单，接口一更新就对不上。现直接请求接口的 /v1/models 真实列表；取不到时回退内置预设并如实提示，不声称列表准确。'),
        ],
    },
    {
        "version": "v3.8.1",
        "date": "2026-10-09",
        "tag": "",
        "items": [
            ('🔧', 'CI 语法检查失败的修复（Python 3.11 兼容）', 'v3.8.0 有一处 f-string 表达式写了跨行，属于 Python 3.12+ 才允许的语法，3.11 直接语法报错，CI 因此失败。已改写为兼容写法，并新增 _f11_scan.py 扫描器（内置违规样本自检）。功能与 v3.8.0 完全一致。'),
        ],
    },
    {
        "version": "v3.8.0",
        "date": "2026-10-09",
        "tag": "",
        "items": [
            ('📝', '综述初稿自动成稿：摘要 / 主要发现 / 局限性直接写出', '初稿骨架从待填模板升级为已成稿的底稿。摘要按「目的—方法—结果—结论」写出真实数字（纳入篇数、设计分布、样本量合计、效应方向计数、结论倾向分布、冲突处数）；4.1 主要发现由效应方向与证据等级自动成段；4.3 局限性从检索条件推导（单数据库、检索时限、未检索灰色文献、未获取全文篇数、样本量缺失篇数、纳入量过少提示）。只有语种限制这类工具无从得知的事留空待填。'),
            ('✍️', '待补写部分改为斜体下划线', '填好的事实与待写的部分都是正文黑字，分不清哪句要自己写。所有占位现统一渲染为 *<u>【待补充：…】</u>*（斜体加下划线），成稿与留白一眼可辨。'),
            ('🏷️', '初稿标题不再带检索式符号', '综述主题从检索式自动带出，标题里会混着 AND / OR / NOT、引号与 [tiab] 字段限定。现自动清洗成可读短语（布尔算符与字段限定剥掉、引号只留内容），从检索页进工作台时同样生效。自己填的主题原样保留。'),
            ('🤖', '大模型辅助补抽「未识别」字段（C4.1-C）', '「⚠️ 有摘要未抽到」说明正则没覆盖到该写法，现可把摘要原文交给大模型按相同列名补抽（对比表新入口，需侧边栏配置 Key）。约束：每个补抽值必须附带摘要原文引句，无引句的值直接丢弃；表格以 ⧉ 标注，「逐篇查看抽取依据」可核对引句；人工修正优先，超过 8 篇自动分批；结果只存会话，导出物仍是不带标记的纯数据。'),
        ],
    },
    {
        "version": "v3.7.1",
        "date": "2026-10-09",
        "tag": "",
        "items": [
            ('🔑', '刷新后关键词不再丢失', '检索完成后刷新页面、预览重载或桌面版重启，输入框会全部清空，关键词得重打一遍。现在每次检索成功后把主副关键词、返回条数、检索式与命中总数落盘到本机 data/last_search.json，新会话进检索页自动回填，改两个字再搜即可。同一会话里正在输入的内容不会被回填覆盖。'),
            ('📶', '返回条数上限 50 → 200', '写综述时 50 条常常不够看出结论分歧。滑杆新增 100 / 200 两档，并注明条数越多检索越慢。'),
            ('🧩', '多个副关键词的组合逻辑', '整串副关键词塞进一个括号组时，PubMed 按隐式 AND 处理，NOT 语义会算错。现按逗号、分号分词，各副关键词独立成概念：AND 要求主关键词与全部副关键词都出现，OR 任一出现即可，NOT 排除含任一词的结果。含空格的词组自动加引号按精确短语匹配，界面说明同步更新。'),
        ],
    },
    {
        "version": "v3.7.0",
        "date": "2026-10-09",
        "tag": "",
        "items": [
            ('🔍', '空白单元格区分「原文没写」与「没抽到」', '对比表里过去一律留白的单元格，现在分三种：摘要未提及（原文确实没写）、⚠️ 有摘要未抽到（摘要里有痕迹却没抽出来，建议回原文核对）、无摘要。前一种没什么可做的，后一种说明工具有遗漏、可以补上，混在一起容易误判。'),
            ('✍️', '对比表可直接编辑', '5 个列（样本量、人群、主要终点、关键效应量、结论）可在表格里直接改，回原文核对后填真实值即可，修正值优先生效并落盘保留（刷新或切页不丢）。改样本量会重算证据强度，改结论会重判结论倾向。研究设计、证据等级、偏倚提示、MeSH 读的是原始摘要，不随手工修正变化，这四条界面上已写清。'),
            ('📐', '结构化完整度：提示哪几篇该回原文补', '按 6 个关键字段（研究设计、样本量、人群、主要终点、关键效应量、结论）给每篇算 x/6，顶部给跨文献汇总与最常缺的字段。这不是质量分。抽不到只说明摘要没写或写法罕见，与研究价值无关，用途是提示哪几篇值得回原文补字段。'),
            ('🧹', '「信息缺失」类提示默认折叠', '偏倚提示里「摘要没写」这一类（未提及 ITT、未说明样本量估算、未提及资助等）与「重点核对 / 建议核对」分开，收进默认关闭的开关并只给条数，表 2 新增「信息缺失」计数列。这类提示平铺出来会让每篇都背一串无意义提示，反而稀释真正该看的几条。'),
            ('📤', '导出物为干净数据', '界面上的三态标注与「✎已修正」角标只服务于界面，CSV / Markdown / ZIP 导出的是修正后的纯数据、不带标记，粘进 Word 或导进 Excel 时拿到的是干净内容。'),
            ('📝', '自愿的使用体验问卷', '用满 15 次后侧边栏出现一张邀请卡（不会自动弹窗），愿意就花 1 分钟填一份，不想填可以选「稍后」或「不再提醒」。问卷与使用计数都只存在本机的 data/ 目录，不自动上传。填完想发回，点「一键通过 GitHub 发送」即可（内容已预填、邮箱手机号已自动清洗）；不想发直接关掉，都不影响功能。'),
        ],
    },
    {
        "version": "v3.6.0",
        "date": "2026-10-09",
        "tag": "",
        "items": [
            ('🧭', '典型场景：从「有什么」到「怎么用」', '落地页与 README 新增四个真实用法，每一步都对应界面上真实存在的入口：开题前 30 分钟摸清一个课题的证据家底；写文献综述，从一批文献到一份可改的初稿；只有付费订阅的 PDF 也要读透并引用；科室小讲课一周内跟上一个新进展。每条标出适用人群、要解决的问题、产出物与耗时。'),
            ('🧰', '使用技巧：七条容易漏掉的用法', '包括「召回太少先看实际执行的检索式」（加引号或字段限定会静默关掉 PubMed 的自动词表映射）、「收藏 ≠ 管理」（分组管课题、标签管跨组维度、笔记放判断，各管一层，导出 Markdown 时一并带上）、「大模型额度怎么省」（抽取式摘要、对比表、证据等级、偏倚提示全部离线不花额度，叙述段按风格与语言分开缓存）。'),
            ('📐', '场景只写有对应功能的做法', '场景里的每一步都能在界面上找到对应按钮，未实现的能力一律不写。该提醒的边界直接写进卡片，例如「骨架是初稿，【待补充：…】必须自己回原文补」「扫描件会明确提示先 OCR，不会给出假装能读的结果」。'),
            ('🩹', '修正三处与实现不符的文案', '「十项能力」实际已是 13 项；「带新角标为 v3.0 新增」实为 v3.x 各版本迭代新增；综述板块原写「全部由离线规则引擎完成、不调用大模型」，v3.5.0 起叙述段会调用自己配置的模型，已改为「对比表、冲突核查、骨架离线完成，只有可选的叙述段会用到模型」。'),
        ],
    },
    {
        "version": "v3.5.0",
        "date": "2026-10-09",
        "tag": "",
        "items": [
            ('🎨', '叙述段可选风格 / 语言 / 模型', '大模型撰写的「结果概述」与「讨论」不再只有一种写法。风格可选学术严谨（书面语体、按研究设计分组、逐条标注来源）或简明扼要（直陈要点、篇幅约为严谨版的一半，仍保留来源与关键数值）；语言可选中文或英文，投英文期刊时不必自己翻译。'),
            ('🔒', '换风格不放宽约束', '换风格不会让模型变松：只能使用给定事实，不得编造样本量、效应量、P 值、结论；事实缺失必须写占位符（中文「摘要未提供」，英文 not reported in the abstract）；不做临床推荐。这四条在任何风格与语言组合下都不变，已用离线断言逐条锁死。'),
            ('🈯', '语言切换覆盖整条链路', '提示词本身、事实行句式（Smith J 等（2023，随机对照试验，样本量 1200…）与 Smith J et al. (2023, randomized controlled trial, n=1200…)）、研究设计标签与冲突类型标签都随语言切换，避免英文段落里夹一个「随机对照试验」。'),
            ('🧠', '叙述段可单独选模型', '新增「叙述段使用模型」下拉（9 个常用 OpenAI 兼容模型预设，默认跟随侧边栏设置），可以便宜模型做批量摘要、强模型写正文。预设只是常用模型名的快捷方式、不是白名单，所选模型必须存在于自己配置的 API 地址下，否则接口会直接报错。'),
            ('🧾', '叙述段标明生成设置', '叙述段上方注明「风格 · 语言（模型）」；改了设置却没重新生成时，会提示需重新点击按钮，避免把上一版结果当成新设置的产物。缓存键也带上风格与语言，切换后不会命中旧结果。'),
        ],
    },
    {
        "version": "v3.4.0",
        "date": "2026-10-09",
        "tag": "",
        "items": [
            ('🧬', 'MeSH 词表联动', '接入 NLM 官方词表，展示识别到的主题词、命中方式（标题精确匹配 / 入口词命中）、官方定义、树号及全部同义词。无法确定对应主题词时直接标注「未找到」，不做推测。'),
            ('⛔', '字段限定会关闭自动映射', 'aspirin[tiab] 及引号短语会使 PubMed 绕过自动词表映射，导致同义词丢失且界面无提示。现于检索完成后折叠展示实际执行的检索式，若为原话回显即说明映射未生效，无需额外请求。'),
            ('🔀', '同义词自动扩展（可选）', '提供关闭 / 仅提示（默认）/ 自动扩展三档。自动扩展将 lung cancer 展开为 ("Lung Neoplasms"[MeSH Terms] OR "Lung Cancer"[tiab] OR "Cancer of the Lung"[tiab] …)；多词关键词按最长匹配切分（immunotherapy lung cancer → immunotherapy + lung cancer）。同义词均取自 NLM 官方词表，未增删；含布尔运算符或字段限定的手写检索式不改写。'),
            ('🏷️', '文献卡与综述对比表新增 MeSH 主题词', '解析 efetch 主题词列表，★ 标记主要主题，保留副主题词及标记；对比表新增「MeSH 主要主题」列。'),
            ('⚠️', '已知限制：最新文献尚无 MeSH 主题词', 'NLM 人工标引存在数月至一年的滞后，新发表文献主题词字段为空属正常现象，文献卡与对比表均已注明。'),
        ],
    },
    {
        "version": "v3.3.1",
        "date": "2026-10-09",
        "tag": "",
        "items": [
            ('🔧', '对比表不再大片「未明确 / 未识别」', '横向对比表里人群、主要终点等字段常常整片空白或标「未明确」，结论倾向也有近半数判不出来。根因是 PubMed 摘要的段落标签在解析时被抹平，抽取器只能在整段文字里靠单一正则去猜。现把段落标签保留下来（Background / Methods / Patients / Results / Conclusions 等），抽取器改为到对应段落里取。'),
            ('📑', '摘要分段器兼容三种标签写法', '支持全大写（BACKGROUND）、首字母大写（Background，BMJ / Lancet 风格）与中文（背景 / 目的 / 方法 / 结果 / 结论）标签；复合标签（METHODS AND RESULTS、DESIGN, SETTING, AND PARTICIPANTS）自动拆开；没有标签的摘要仍回落整段扫描，不影响已有行为。本地 PDF 全文解析产物同样受益。'),
            ('🎯', '修掉「结论」取错段落的缺陷', '摘要把标签与正文连排时（… RESULTS: xxx. CONCLUSIONS: yyy.），结论会误取成 RESULTS 段末句，「结论」列显示的其实是结果。现结论一律取自 CONCLUSIONS 段，结论倾向的判定也随之准确。'),
            ('🔎', '人群 / 终点 / 样本量抽取扩容', '人群新增年龄限定（adults aged 75 years or older）、状态限定（community-dwelling adults）与中文模式；主要终点放宽以命中 primary composite endpoint，抽不到时给出相关表述并标注「摘要未标注主要终点」；样本量新增数字前置模式（19,114 community-dwelling adults…）。同一批演示文献：人群空 5 → 3 篇，主要终点空 6 → 3 篇，结论倾向「未明确」3 → 0 篇。'),
            ('⚖️', '极性判定补裸动词与保守兜底', '补上 reduced / reduces / lowered / fewer / 降低 / 减少 等不带 significantly 的方向动词；文字线索全落空时，仅当效应量 < 1 且结局为不良事件才做方向兜底，其余仍保留「未明确」。否定式（did not reduce / no reduction）有专门护栏，不会被误判为有效。'),
        ],
    },
    {
        "version": "v3.3.0",
        "date": "2026-10-09",
        "tag": "",
        "items": [
            ('🧰', '结构化评价工具（P3-C4）', '证据提示原本只能看摘要，摘要没写的只能标「信息缺失」，而方法学评价必须回全文逐条回答。现按研究设计自动匹配工具：随机对照试验用 RoB 2（5 个域）；队列研究用 NOS 队列版，病例对照研究用 NOS 病例对照版，横断面研究用改良 NOS；系统评价 / Meta 分析用 AMSTAR-2（16 项，标出 7 个关键域）；临床指南用 AGREE II；叙述性综述用 SANRA；病例报告用 CARE；基础 / 动物实验用 SYRCLE。工具以可勾选清单呈现，判定方式与出处一并给出。'),
            ('⚖️', 'GRADE 分级自查', '按研究设计给出起始等级（RCT 与含 RCT 的系统评价为高，观察性研究为低，病例报告为极低，动物 / 体外与叙述性综述标为不适用），再列出 5 个降级因素（偏倚风险、不一致性、间接性、不精确性、发表偏倚）与 3 个升级因素（效应量大、剂量-反应、混杂方向）的判据与查法。工具不给最终等级，只把判据摆开避免漏项。'),
            ('🔬', '偏倚提示规则扩到 32 条', '新增十类摘要层面线索：未提及意向性分析（ITT）、未说明样本量估算、疑似企业资助或利益相关、未提及资助与利益冲突、含事后 / 亚组 / 探索性分析、提及基线不均衡、主要终点为复合终点、观察性研究未提及混杂调整、失访或退出比例偏高（支持两种语序）、预试验 / 可行性研究。摘要没写造成的提示一律落在「信息缺失」档，只有摘要明确出现值得警惕的表述才升级。'),
            ('📦', '两份新导出物纳入 ZIP', '「结构化评价自查清单」（按工具分组列出信号问题 + 逐篇适用工具）与「GRADE 分级自查表」（起始等级表 + 降级 / 升级因素表 + 逐篇起始等级），均已纳入「打包导出全部产出」，ZIP 内编号为表 4、表 5。'),
            ('✋', '只给问题，不造权威错觉', '所有信号问题都注明真实出处（BMJ / Ottawa Hospital Research Institute / AMSTAR 团队等）并声明为便于快速核对做了大幅简化，不是原版量表的完整复现；每条题项都提示答案应取自全文，不能仅凭摘要作答。摘要没写不等于研究没做。'),
        ],
    },
    {
        "version": "v3.2.0",
        "date": "2026-10-09",
        "tag": "",
        "items": [
            ('📄', 'PDF 全文分析（P3-C3）', '全文分析原本只覆盖 PMC 开放获取与网页抓取兜底，手里下载的 PDF 无处可用，而它恰恰是最常见的形态。现可上传本地 PDF，自动解析出标题、作者、期刊、年份、DOI 等元数据与分章节正文（摘要、引言、方法、结果、讨论、参考文献），并接上后续全部能力：章节化全文摘要、原文定位、数据分析、图表表格解析与引用导出。'),
            ('🧩', '分栏感知的版面重建', 'PDF 的文本按坐标散落，直接抽取会把左右双栏并成一行，页边元数据栏（Citation / Editor / 版权行）与右栏正文在同一高度被合并，正文句子被从中截断。现按词覆盖度曲线自动识别分栏边界，先左栏后右栏重建阅读顺序，并把页边元数据整块剔除。'),
            ('✂️', '扫描件不支持，并给出可操作提示', '纯图片型 PDF 没有文本层，不接 OCR 的解析会得到空正文。与其返回满是噪声的摘要，不如明确提示「未检测到文本层」，并建议改用带文本层的版本或用 OCR 工具处理后重传。加密 PDF、损坏文件、超页数 / 超体积也都有对应提示。'),
            ('🔒', '版权确认门与内存、页数上限', '上传区前有版权与合规确认勾选（未确认不开放上传），因上传与后续使用产生的责任由使用者自行承担。解析全程在内存中进行，文件不保存到磁盘、不上传第三方；单篇默认上限 20 MB / 200 页（MEDLIT_PDF_MAX_MB / MEDLIT_PDF_MAX_PAGES 可调），面向单篇论文而非整本书。'),
            ('🔗', '下游能力全部打通', '解析结果与 PMC 全文同构，因此零改动接入已有链路：章节化全文摘要 + 原文定位（摘要句子与关键数值回溯原文）+ 全文数据分析（篇幅分布、高频关键词、统计指标）；图表与表格解析，表格结构化提取并可导出 CSV，图形按整页渲染查看，可选多模态 LLM 逐页看图解读；引用导出六种格式；一键「加入文献架」，到综述工作台与检索结果、收藏一起参与横向对比、结论冲突核查、PRISMA 记录与初稿骨架。'),
            ('🖥️', '桌面版同步纳入 PDF 解析依赖', '打包配置补齐 pdfminer 的 CMap 数据（缺了中文 PDF 会解析失败）与 pypdfium2 的原生渲染库（页面转 PNG 用），并显式声明这些运行时才 import 的模块，避免打包版 PDF 功能在 exe 里失效。'),
            ('🧪', '测试补齐 141 条断言', '新增 _test_pdfdoc.py（99 条：自带极简 PDF 生成器，覆盖元数据、启发式标题、表格、扫描件、错误边界、双栏与页边、与下游四模块对接、导出与渲染）与 _test_pdfpage.py（42 条：注入假上传器，真跑通上传 → 解析 → 摘要 → 定位 → 图表表格 → 引用导出 → 文献架 → 综述纳入整条链路），并都纳入一键发布的测试清单。'),
        ],
    },
    {
        "version": "v3.1.1",
        "date": "2026-10-08",
        "tag": "",
        "items": [
            ('🐞', '修一个会静默清空分组的缺陷', '文献库卡片里的「分组」下拉，无论这篇属于哪个分组都显示「（未分组）」：首次渲染时控件状态还不存在，被一律归零，覆盖了真实分组。不改下拉、直接点保存，分组就被清掉了。现下拉正确回显已存分组，并加了冒烟断言把这条锁住。'),
            ('🔄', '批量操作后卡片同步刷新', '批量移动分组、加标签、移出标签，以及标签管理里的重命名与删除，过去不会更新卡片编辑区里的旧值，之后顺手点一下卡片的保存就会把刚做好的批量结果覆盖回去。现这些操作完成后会清掉对应卡片的控件状态，重新按存储值初始化。'),
            ('📐', '引用导出挪到文献列表之后', '引用导出自带一段很长的代码预览，原先夹在筛选区与文献卡片之间，会把真正要看的文献卡片整段挤出首屏。现顺序改为筛选 → 卡片列表 → 导出，更贴近先筛、再看、后导的用法。'),
            ('🖼️', '落地页补上文献库管理', '在线介绍页的能力清单增至十一项，并新增两张真实截图：一张是卡片上的「分组 / 标签 / 有笔记」与展开后的编辑区，一张是分组与标签管理、批量整理与三档筛选。'),
        ],
    },
    {
        "version": "v3.1.0",
        "date": "2026-10-08",
        "tag": "",
        "items": [
            ('🗂️', '我的收藏升级为我的文献库（P3-C1）', '收藏夹原本是一个平铺列表，文献一多就没法用。现补上三件整理文献的基础设施：分组（一篇文献归入一个分组，对应一个课题或一篇综述）、标签（多对多，用于研究类型 / 干预 / 人群这类跨分组维度）、笔记（纯文本，记录纳入与排除的理由、样本量疑问、待复核的点）。四张统计卡（收藏数、分组数、标签数、已写笔记数）一眼看清文献库状态。'),
            ('📁', '分组的新建、重命名与删除', '分组是显式实体，可以改名、可以删除。删除分组不会删除文献，组内文献退回「未分组」，标签与笔记原样保留。分组名唯一（忽略大小写与连续空白差异），重名会被拦下，不会悄悄建两个「综述选题」。'),
            ('🏷️', '标签：跨分组筛选，改名即合并', '支持一次给多篇批量打标签，也可以在一处把标签改名或删除（改名会自动合并到已有同名标签，不留重复项）。文献列表可按分组 + 标签 + 关键词三者叠加筛选，关键词同时匹配标题、作者、标签、笔记与分组名，记在笔记里的判断也能被搜到。'),
            ('🧰', '批量整理与按筛选结果导出', '勾选多篇即可批量移动分组、加标签、移出收藏。Markdown 导出跟随当前筛选结果，并把分组、标签、笔记一并写进去，可以直接当综述的文献清单底稿；引用导出（BibTeX / RIS 等）同样只导出筛选出的部分。'),
            ('🔗', '旧链接继续可用', '侧边栏入口由「我的收藏」改名为「我的文献库」，但 ?page=我的收藏 这类已发出的深层链接与用户书签仍然有效（内部保留了旧名映射），不会因为改名变成死链。'),
            ('💾', '数据仍只存本机或本会话', '分组、标签、笔记单独存于 data/library.json（在线版按会话分片：data/users/<会话>/library.json），与 favorites.json 分离，因此旧数据零迁移；取消收藏时同步清掉该文献的标注，不留孤儿数据。'),
        ],
    },
    {
        "version": "v3.0.1",
        "date": "2026-10-08",
        "tag": "",
        "items": [
            ('📇', '引用格式导出（P3-C2）', '把检索结果、收藏、综述纳入文献一键导出成参考文献管理器能直接导入的格式：BibTeX（Zotero / JabRef / LaTeX）、RIS（Zotero / EndNote / Mendeley / NoteExpress）、EndNote 的 .enw 标签格式、PubMed 原生 MEDLINE、医学期刊常用的 Vancouver，以及中文毕业论文常用的 GB/T 7714-2015。六个入口共用同一个导出区：选格式 → 就地预览 → 下载。'),
            ('🔎', '元数据补全', '原来只抽取标题、作者、期刊、年份、DOI，缺卷、期、页码，导出的 .bib 导进 Zotero 后还要手工补。现于检索阶段就一并取出卷、期、页码、ISSN、文献类型与语种，并保留团体作者（研究协作组），否则像 The ESPRIT Study Group 这类署名会整条丢失。'),
            ('✋', '缺字段就整条省略', '电子优先发表（online ahead of print）本来就没有卷期页码，这是正常情况。工具宁可少写一项，也不填「[待补充]」之类的假值，参考文献写错了比缺了更难被发现。导出区会单独提示哪些记录缺了引用必需字段（作者 / 标题 / 期刊 / 年份），建议回检索页重新抓取。'),
            ('📦', '综述 ZIP 附带 .bib 与 .ris', '「打包导出全部产出」现在除了对比表、冲突核查、偏倚清单、适用性对照、筛选记录、初稿骨架与 Vancouver 参考文献，还多两份可直接导入文献管理器的 .bib 与 .ris，省去从 Markdown 里手抄一遍。'),
            ('🖼️', '落地页补入综述工作台与证据化', '在线介绍页还停留在「八项能力」，与 v3.0 的实际能力脱节。现扩到十项（新增项带「新」角标），并新增「从一批文献，到一篇综述初稿」专章，用横向对比表、证据与适用性、结论冲突核查、初稿骨架四张真实截图说明新链路；生成脚本改为按基名查找截图并自动降采样，落地页体积可控。'),
            ('🧹', '工程侧清理', '删除 6 个历史打包目录（约 880 MB），仅保留 build_v15 / dist_v15；清理早期隧道脚本与工作区调试残留；docs/shots 统一为 JPEG。'),
        ],
    },
    {
        "version": "v3.0.0",
        "date": "2026-10-08",
        "tag": "",
        "items": [
            ('🩺', '证据与适用性核查（P2 证据化）', '面向临床医生，在综述工作台内新增第三个标签页「证据与适用性」。不替医生判断文献该不该用（工具读不到全文，也不认识患者），而是把摘要层面能核实的线索摆出来：样本量够不够、有没有对照、随访多长、终点是硬终点还是替代终点、结果是在谁身上得到的，并逐条附上判定理由与原文依据。'),
            ('🔬', '偏倚风险提示（约 16 条规则）', '覆盖小样本、单臂 / 无对照、观察性设计使用因果表述、回顾性设计、横断面设计、未提及盲法、自报结局、未报告区间估计或 P 值、结论缺效应量、未见试验注册号、单中心或未说明中心数、替代终点、随访过短或未说明、系统评价未见异质性说明、基础 / 动物实验等。级别只用重点核对 / 建议核对 / 信息缺失三档，刻意不用高 / 中 / 低风险，那是 RoB 2 等工具的专有判定，需要逐条回答信号问题才能给出。'),
            ('📐', '研究类型分层与证据等级参考', '研究设计归入干预性研究、观察性研究、描述性研究、二次研究、基础研究五大类，并给出牛津 CEBM 2011 的简化等级对应（1a 系统评价 → 5 机制研究）。等级仅按设计映射，未考虑偏倚风险、间接性与不一致性，页面与导出文件中均标明不是正式分级、不能写进方法学部分。'),
            ('🧭', '临床适用性五维对照', '自动抽出人群、干预、主要终点性质、随访时长、研究场景（单中心 / 多中心 / 未说明），再把外推到具体患者前需要逐维回答的问题列全，避免漏项。工具不知道使用者面对的患者是谁，因此只列维度、不给适用或不适用的结论。'),
            ('📄', '两份新导出物', '「偏倚风险提示清单」（含跨文献频次汇总，可写进讨论与局限性）与「临床适用性对照」（含逐维比对清单），均已纳入 ZIP 打包；横向对比表新增「证据等级」「偏倚提示」两列；综述初稿骨架新增 3.4 偏倚风险与适用性概览，并在 4.3 局限性中自动补入本次纳入研究集中出现的问题。'),
            ('✋', '抽取不到就写「未提及」', '「摘要未提及」本身就是需要标注的信息：摘要没写研究中心数，不等于它是单中心；没写随访时长，不等于随访短。这类条目单独归入「信息缺失」档，不会与真正的风险提示混在一起。'),
            ('💳', '版本号升至 v3.0.0', 'P2 双主线（综述化 + 证据化）完成，功能面向的两类用户（医学生 / 研究生、临床医生）均已覆盖，故由 2.x 升为 3.0.0 标记产品形态定型；本次同时把 v2.9.0 的综述工作台一并纳入正式发版。'),
        ],
    },
    {
        "version": "v2.9.0",
        "date": "2026-10-07",
        "tag": "",
        "items": [
            ('🧾', '综述工作台（P2 综述化）', '面向医学生 / 研究生的综述写作流程，独立成页而不是塞进单篇摘要页。综述本质是多篇之间的工作：选好一组文献后，自动抽取研究设计、样本量、人群、主要终点、效应量与结论，生成可导出的横向对比表（CSV / Markdown），并附逐篇抽取依据，每条字段都能回原文核对。'),
            ('⚖️', '结论冲突自动识别', '同一主题词下结论相反时高亮提示，分结论极性冲突与效应方向冲突（HR / OR / RR 方向相反）两类，标注可信度并给出可能原因（人群、剂量、终点定义、随访时长差异）。工具只提示需人工核对，不做对错仲裁；未检出冲突时也会明确说明这不等于结论一致。'),
            ('🧾', 'PRISMA 式筛选记录', '检索式与数据库命中总数从文献检索页自动带出（NCBI 的 count 字段），按 PMID / 标题自动检测重复条数，逐级推导命中 → 去重 → 题摘筛查 → 全文评估 → 纳入并保证表内数字自洽，可导出成可直接写进综述方法学部分的 Markdown。'),
            ('📝', '一键生成综述初稿骨架', '按引言 → 资料与方法 → 结果 → 讨论 → 结论 → 参考文献组织已读文献，事实来自自动抽取并标注来源，需要作者判断的地方一律写成【待补充：…】占位，避免把机器填的内容误当成自己写好的结论；参考文献按 Vancouver 格式生成。可打包导出 ZIP（对比表 + 冲突核查 + 筛选记录 + 初稿）。'),
            ('🤖', '可选：大模型撰写叙述段', '只把工具已抽取好的结构化事实交给模型，提示词中硬性禁止编造样本量、效应量与结论，事实缺失处要求写「摘要未提供」；使用自带 Key，离线规则引擎部分零成本。'),
            ('🔬', '配套抽取引擎（core/review.py，纯离线）', '研究设计识别按优先级匹配 9 类设计（系统评价 / RCT / 指南 / 队列 / 病例对照 / 横断面 / 病例报告 / 基础实验 / 叙述性综述），样本量支持 n=、共纳入 N 例等 4 种写法，结论倾向判定先排除否定式表述（no significant…）再判阳性，避免把「无显著改善」读成「有效」；证据强度用设计基准分 + 样本量 + 是否报告区间估计的可解释加权，明确标注不是 GRADE 分级。'),
            ('🗂️', '综述工作区可续写', '综述会跨多次刷新：主题、勾选文献、筛选记录数字按会话作用域落盘，刷新或来回切页不丢进度。'),
            ('🐛', '修复「检索历史」页清空按钮报错', '该按钮引用了并不存在的常量 storage.HIST_FILE，点击会抛 AttributeError；现改为调用 storage.clear_history()。'),
        ],
    },
    {
        "version": "v2.8.2",
        "date": "2026-10-07",
        "tag": "",
        "items": [
            ('🔒', '隐私说明改为正式政策页', '隐私说明原来塞在侧边栏设置区里，位置突兀、篇幅零散。现改为左侧导航中的独立页面「隐私与数据」，按适用范围、不收集的信息、本地与会话数据、结果缓存、第三方服务、API 凭据、运行日志、医疗免责、数据来源、权利、政策更新共十一条组织，并标注版本与生效日期。'),
            ('✅', '首次使用需确认数据处理方式', '首次进入时页面顶部会出现一条确认条，说明本工具不收集个人身份信息、以及第三方传输的例外情况，勾选后方可使用。不做强制弹窗打断，确认状态仅存在本次会话，不写盘、不上传。'),
            ('🔗', '支持页面深链', 'URL 追加 ?page=文献检索 可直接落到指定页面，便于分享链接与落地页直达。'),
        ],
    },
    {
        "version": "v2.8.1",
        "date": "2026-10-07",
        "tag": "",
        "items": [
            ('🚀', '一键启动在端口被占用时不再直接失败', '只要目标端口已被占用（本机常是自己上次没关的服务），Streamlit 会立刻退出，双击窗口一闪而过。现启动器先判断占用者是谁：是自己就直接复用并打开浏览器；是别的程序就自动往后找空闲端口并在窗口里明示；都没有则给出明确处理办法。'),
            ('🌐', '一键发布在网络不佳时不再推送失败', '本机到 GitHub 的链路会随机回 502、被 reset 或 TLS 中断，之前每次发版都得手动加 --resolve 参数。现发布会依次尝试轮换直连 IP → 清空代理直连 → 系统默认，自动找到可用链路；标签与分支分别推送，全部失败时给出补救命令。'),
            ('🧭', '启动脚本的解释器查找更可靠', '找不到 WorkBuddy 环境就退化成裸 python，而本机 PATH 里没有 python，双击后同样是一闪而过。现依次尝试 WorkBuddy 环境、D:\\python、PATH 上的 python.exe，全都没有才报错并给出安装指引。'),
            ('🛑', '新增「一键停止」', '启动器会复用已有实例，因此需要明确的关闭入口才能重启。新增 stop.py / 一键停止.bat：按端口定位本项目的进程并结束，认进程命令行避免误杀无关程序。'),
        ],
    },
    {
        "version": "v2.8.0",
        "date": "2026-10-07",
        "tag": "",
        "items": [
            ('🔒', '隐私说明页', '侧边栏新增可查阅的完整隐私说明：哪些数据存在哪、谁能看到、声明不收集姓名 / 手机号 / 邮箱等身份信息、不做用户画像与行为追踪、不用于训练模型，并特别提示翻译与大模型请求会经过第三方服务，请勿粘贴患者身份信息。'),
            ('💬', '反馈渠道闭环', '侧边栏「报错 / 建议」：填完一键留档到本机并复制可粘贴文本，或直接生成已预填标题与诊断信息的 GitHub Issue 链接。诊断信息含版本、请求统计、依赖成功率、配额与缓存状态，报错时不用再反复描述环境。反馈内容会自动隐去邮箱与手机号。'),
            ('🌐', '产品落地页', '新增单文件 index.html 落地页：一句话定位、真实界面截图、八项能力清单、三步上手、FAQ 与合规声明，图片以 base64 内嵌，可直接用浏览器打开或作为 GitHub Pages 发布。截图由 _shot.py 通过 CDP 实拍，重拍后跑 _build_landing.py 即可同步更新。'),
            ('⚡', 'NCBI API Key 已启用', '已申请并配置免费 Key，检索限速从约 3 次/秒提升到约 10 次/秒。实测 8 路并发冷检索从 7.67 秒降到 3.9–6.1 秒。'),
        ],
    },
    {
        "version": "v2.7.0",
        "date": "2026-10-07",
        "tag": "",
        "items": [
            ('⚡', '持久化缓存', '同一篇文章被反复检索、反复翻译、反复抓取全文，每次都要重新调用 PubMed 与翻译接口，而 PubMed 是按 IP 限速的公共配额，重复消耗等于把所有人的额度一起花掉。现检索结果、PMC 全文、抽取摘要、中文译文、大模型摘要全部走本地缓存：同一检索式的 8 路并发从 7.7 秒降到几乎瞬时，摘要重算从 50 毫秒降到 0 毫秒，命中即不消耗翻译额度与 NCBI 限速配额。带 TTL 与条数上限自动淘汰，不会把磁盘撑爆。'),
            ('🧹', '缓存可控可观测', '侧边栏「🩺 运行诊断」显示命中率、命中 / 未命中、写入与淘汰条数，并提供一键清空。命中率偏低说明上游在被反复打，是额度耗尽的前兆信号。'),
        ],
    },
    {
        "version": "v2.6.1",
        "date": "2026-10-07",
        "tag": "",
        "items": [
            ('⚡', '支持 NCBI API Key', 'NCBI 对每个 IP 的检索限速：无 Key 约 3 次/秒，申请免费 Key 后约 10 次/秒。这是多人同时使用时唯一的硬天花板。侧边栏「⚡ NCBI 速率」填入 Key 即刻生效，所有检索请求自动附带，不需要改任何代码。免费申请地址已放在设置旁，填一个邮箱即可。'),
        ],
    },
    {
        "version": "v2.6.0",
        "date": "2026-10-07",
        "tag": "",
        "items": [
            ('🔒', '会话级数据隔离', '原先所有访问者共用同一份历史与收藏，A 用户的检索记录会出现在 B 用户的界面上。现按会话分片存储：在线版每人一份独立数据，互不可见；桌面版仍存本机，路径不变、旧数据平滑迁移。'),
            ('📊', '用量配额与成本封顶', '新增配额层：单次会话的翻译字符数与大模型调用次数均有上限，超出时给出明确提示与解决建议（自带密钥或稍后再试）。宿主额度仅作体验池，可用环境变量一键关闭（MEDLIT_HOST_QUOTA=0），保证传播时不会被刷穿。'),
            ('🔑', '自带密钥即不受限', '在侧边栏填入自己的翻译或大模型密钥后，用量不再受共享额度限制；未填密钥时走体验池并受配额保护。这样既降低传播门槛，又不会让一个链接烧穿账单。'),
            ('💓', '外部依赖健康监控', '按天统计 PubMed、Europe PMC、图表、翻译、大模型、Unpaywall 六大外部依赖的成功率与错误次数，请求层自动埋点。侧边栏「🩺 运行诊断」可查看近 7 天成功率，出现大范围失败时能立刻判断是自家代码问题还是上游接口波动。'),
        ],
    },
    {
        "version": "v2.5.0",
        "date": "2026-10-07",
        "tag": "",
        "items": [
            ('🌐', '统一外部请求层', '原先所有对外请求（PubMed 检索、PMC 全文、图表包、翻译接口、大模型接口）各写各的超时与重试，部分通道完全没有重试，偶发网络抖动或限流就直接失败。现统一由一层接管：连接 / 读取超时分离（任何请求都有硬上限，不再把界面卡死）、5xx 与 429 自动指数退避重试并尊重服务端的 Retry-After、按域名控制请求速率（遵守 NCBI 与 Europe PMC 的公开接口配额）、统一埋点。侧边栏「🩺 运行诊断」新增请求统计（次数 / 重试 / 失败 / 平均耗时）。'),
            ('🚀', '修复本地启动后浏览器打不开', '旧启动脚本固定等 4 秒就打开浏览器，而服务冷启动常需更久，浏览器实际打在尚未监听的端口上，表现就是连接很慢。改为轮询端口真实就绪后再打开浏览器，并显示本次启动耗时。'),
            ('🏷', '版本号单一来源与发布流程', '版本号原来散落在多处、手工同步，容易出现页面显示与实际代码不一致。现统一由 version.py 提供；新增 release.py 一键发布（跑测试 → 改版本 → 同步部署目录 → 提交打标签 → 推送）与 GitHub Actions 持续集成，每次推送自动跑冒烟测试。'),
        ],
    },
    {
        "version": "v2.4.0",
        "date": "2026-10-07",
        "tag": "",
        "items": [
            ('📋', '运行日志系统', '新增按天落盘的运行日志：记录检索、全文抓取、摘要生成、图表解析、批量任务与 LLM 调用的关键步骤与耗时，异常自动记录堆栈。以前线上或桌面端出问题只能靠猜，现在有据可查。侧边栏「🩺 运行诊断」可直接在页面内查看最近日志。日志默认保留 14 天，写入失败不影响正常使用。'),
            ('🛡', '异常兜底与友好报错', '安装全局异常钩子，未捕获异常自动写入日志；抽取式摘要等核心操作失败时给出可读提示与「技术详情」折叠区，不再把程序栈直接抛出来；后台任务失败、批量单篇失败均记录日志。'),
            ('⚕️', '医疗免责声明与数据来源署名', '每页底部常驻页脚：声明摘要由算法或大模型自动生成，不能作为临床诊断与用药依据，诊疗决策须核对原文并由专业医师判断；同时署名数据来源 PubMed / PMC（NCBI 下属 NLM）并遵循其使用条款；附隐私说明，不收集个人身份信息，历史与收藏仅存本机或当前会话。'),
        ],
    },
    {
        "version": "v2.3.1",
        "date": "2026-09-30",
        "tag": "",
        "items": [
            ('🖼', '修复桌面版图表解析必然失败', '桌面版打包配置漏掉了浏览器兜底通道依赖的 websocket 组件，导致非开放获取文献的图表解析在 exe 中必定失败（OA 文献不受影响）。已补全打包依赖并让缺失时优雅降级。'),
            ('🩺', '失败原因不再被吞掉', '图表解析失败时，页面会直接显示真实原因与解决建议。原先只弹一条 80 字提示，随即变成含糊的「未解析出图表」。'),
        ],
    },
    {
        "version": "v2.3.0",
        "date": "2026-09-28",
        "tag": "",
        "items": [
            ('🌐', '翻译引擎可插拔（腾讯云 TMT）', '侧边栏新增翻译设置：配置腾讯云机器翻译的 SecretId/SecretKey 后，中文翻译走腾讯接口（每月 500 万字符免费额度，质量更高），解决 MyMemory 免费额度不够对外发行的问题；未配置或调用失败时自动回退 MyMemory，行为与旧版一致。云端可在 Streamlit Secrets 中配置 TENCENT_SECRET_ID / TENCENT_SECRET_KEY。'),
        ],
    },
    {
        "version": "v2.2.0",
        "date": "2026-09-27",
        "tag": "",
        "items": [
            ('🔎', '全文关键词搜索定位', '在原文定位区输入关键词（中英文均可，多个词空格分隔为同时包含），实时列出所有命中原句：所在章节、章节内句序、命中次数，关键词蓝色高亮。'),
            ('🩹', '统计指标提取修复', '置信区间不再混入引文标记或截断（兼容 Lancet 中点小数 0·64）；百分比排除 95% CI 误抓；计数标签明确「全文 N 处 · 去重 M 种」。'),
        ],
    },
    {
        "version": "v2.1.0",
        "date": "2026-09-27",
        "tag": "",
        "items": [
            ('📍', '摘要句子原文定位', '全文摘要的每一句都可回溯到原文出处：所在章节、章节内句序、带高亮的英文原句；中文翻译与英文原文对照展示，摘要不再来路不明。'),
            ('🔢', '关键数值原文定位', '自动扫描全文的 P 值 / 95%CI / HR·OR·RR / 百分比 / 样本量 n / 均值±SD，逐类列出数值及其所在原句并高亮数值本身。'),
            ('🤖', 'LLM 总结也可溯源', 'LLM 深度总结逐句匹配原文出处。中文总结靠数字与药名 / 指标名等术语跨语言对齐原文，匹配度不足时如实标注「未找到可靠对应」。'),
        ],
    },
    {
        "version": "v2.0.0",
        "date": "2026-09-26",
        "tag": "",
        "items": [
            ('🎉', 'v2.0.0 正式发行', '第一个对外发布的大版本，功能集齐：智能检索 → 全文抓取 → 章节摘要 → 数据分析 → 批量导出全链路。'),
            ('📊', '批量统计表（Excel 三工作表）', '文献汇总 / 章节明细 / 统计指标，DOI、PubMed 直链、摘要预览、指标逐条示例值，打开即用。'),
            ('🌐', '全文三通道抓取', 'PMC 开放全文 → Unpaywall 开放副本 → 无头浏览器网页提取（带反拦截），免费能看的文献基本都能抓到。'),
            ('⚙️', '后台任务与并行架构', '全文摘要与图表解析可同时进行；批量全文多线程并行抓取；NCBI 全局限流防封禁。'),
            ('🔍', '搜索引擎升级', '主副关键词 AND / OR / NOT 组合、期刊筛选（全名自动转缩写）、作者国籍筛选、检索式实时预览。'),
        ],
    },
    {
        "version": "v1.7.2",
        "date": "2026-09-26",
        "tag": "",
        "items": [
            ('📊', '批量统计表升级为 Excel 三工作表', '文献汇总 / 章节明细 / 统计指标，表头配色、冻结首行、自动筛选，打开即用。'),
            ('🧾', '汇总表信息量大幅扩充', '新增 DOI、PubMed 直达链接、摘要章节数、摘要预览、全文字符数、统计指标汇总（如 P值×5；95%CI×3）、状态红绿标注。'),
            ('📐', '章节明细与指标逐条成表', '每篇文献各章节的词数 / 字符数分布一张表，P 值、置信区间等统计指标连同示例值逐条列出。'),
        ],
    },
    {
        "version": "v1.7.1",
        "date": "2026-09-26",
        "tag": "",
        "items": [
            ('🌐', '全文抓取第三通道：网页提取兜底', 'PMC 双通道失败时自动改用无头浏览器打开 Unpaywall 开放副本 / DOI 出版社页，直接提取渲染正文（非截图识别，零损失）。内置反拦截（正常 UA + 隐藏自动化特征）与正文清理（掐导航、去参考文献）。'),
            ('🧹', '失败原因更透明', '付费墙文献会明确提示：免费渠道（含网页抓取）均无法获取其全文。例如 MMR 那篇 Elsevier 文章经 Unpaywall 确认无任何合法开放副本。'),
            ('✅', '修复批量任务 TypeError', '批量结果返回结构与页面消费端不一致导致的崩溃（list 按 dict 解析），已端到端实测修复。'),
        ],
    },
    {
        "version": "v1.7.0",
        "date": "2026-09-26",
        "tag": "",
        "items": [
            ('🚀', '批量全文摘要（新页面）', '从检索结果或收藏中勾选多篇文献，多线程并行抓取 PMC 开放全文并逐篇生成章节化摘要与数据分析；单篇失败不影响整批。内置 NCBI 限流保护（约 3 请求/秒），并行再快也不会触发封禁。'),
            ('⬇️', '一键导出摘要与分析', '批量结果一键导出：完整 Markdown（每篇含章节化摘要、词数统计、高频关键词、统计指标）+ CSV 统计表（Excel 直接打开，UTF-8 BOM 不乱码）。'),
            ('⚡', '全文摘要与图表解析可同时进行', '智能摘要页的全文抓取与图表抓取改为后台任务执行：点击后页面立即可操作，进度实时显示在左侧「后台任务」，完成自动刷新展示，两项功能互不阻塞。'),
            ('🧵', '并行架构升级', '新增轻量后台任务管理器（工作线程 + UI 轮询），所有耗时操作不再冻结页面。'),
        ],
    },
    {
        "version": "v1.6.0",
        "date": "2026-09-25",
        "tag": "",
        "items": [
            ('🔎', '搜索引擎升级：主副关键词组合', '主关键词 + 副关键词自由组合，支持 AND（同时包含）/ OR（任一包含）/ NOT（排除）三种逻辑，检索式实时预览。'),
            ('📰', '来源期刊筛选', '按期刊过滤检索结果，支持全名（Nature Medicine）或缩写（Nat Med）：全名经本地主流期刊词典 + NLM Catalog 自动转为 MEDLINE 缩写，解决全名检索为 0 条的问题。'),
            ('🌍', '作者国籍 / 地区筛选', '新增 20 个常用国家 / 地区下拉（按作者单位 Affiliation 匹配），可结合日期范围精准定位某国团队的研究。'),
            ('🛠', '检索式优先级修复', 'OR 逻辑与期刊 / 国籍 / 日期过滤组合时自动整体加括号，避免 PubMed 运算符优先级（AND 高于 OR）把过滤条件吞进 OR 分支导致结果错误。'),
        ],
    },
    {
        "version": "v1.5.1",
        "date": "2026-09-25",
        "tag": "",
        "items": [
            ('🌐', '英文摘要不再混入中文小标题', '全文摘要结果同时保留原始章节标题与中文标题：中文模式用中文小标题，英文模式用原文标题（Background / Methods / Results 等），切换语言无需重新生成。'),
            ('🎨', '智能摘要页输入区分区强化', '摘要语言与长度并排放入浅青色参数子面板，控件标题加粗、组间距加大，来源选择 / 参数设置 / 生成按钮三段更分明。'),
        ],
    },
    {
        "version": "v1.5.0",
        "date": "2026-09-25",
        "tag": "",
        "items": [
            ('🖥️', '新增 Windows 桌面版', 'PyInstaller + pywebview 打包为独立应用：双击 exe 起本地服务并弹出独立窗口，无需安装 Python；数据（收藏 / 历史 / 图表缓存）写入 %APPDATA%\\MedLitSummary，程序目录只读。'),
            ('🔧', '数据目录支持环境变量覆盖', 'storage 与图表缓存路径支持 MEDLIT_DATA_DIR 环境变量重定向，为打包发行做准备；网页版行为不变。'),
        ],
    },
    {
        "version": "v1.4.1",
        "date": "2026-09-25",
        "items": [
            ('🔗', '修复：有时无法解析出 PDF 全文', '全文抓取升级为双通道：NCBI efetch 无正文（非 OA 文献返回 200 + 错误 XML）时自动改走 Europe PMC 全文 XML，两源互补；正文开头未分章节的段落不再丢失；全部失败时给出具体原因与建议（该文献可能不属于 PMC 开放获取子集，可走 PDF / DOI 链接），不再静默返回空。已批量实测 13 篇文献全部解析成功。'),
            ('🎨', '各功能页界面与主页风格协同', '检索结果、摘要来源、全文分析、图表解析、收藏列表、检索记录、更新日志等区块标题统一为主页的竖条标题样式；摘要 / LLM 总结 / 图表概括结果卡片化呈现；输入框、下拉、滑杆、指标卡、提示框统一主题色与圆角；收藏页新增统计卡（收藏数 / 可用摘要数）。'),
        ],
    },
    {
        "version": "v1.4.0",
        "date": "2026-09-25",
        "tag": "",
        "items": [
            ('🏠', '主页焕新', '简约清新视觉主题：Hero 区配内联 SVG 插画、核心能力卡片点击直达、快速开始三步引导、数据概览与最近动态；侧边栏升级为图标导航（streamlit-option-menu），全站按钮 / 卡片 / 步骤条统一新样式。'),
        ],
    },
    {
        "version": "v1.3.6",
        "date": "2026-09-25",
        "tag": "",
        "items": [
            ('🤫', '修复：LLM 结果导出了思考过程', '推理模型把输出额度全部用于思考、正文为空时，旧版会把思考过程（reasoning_content）当结果展示。现改为自动加大输出额度重试一次，仍失败则给出可操作提示；思考过程绝不作为结果输出。已实测 deepseek-flash 多图视觉分析输出干净。'),
        ],
    },
    {
        "version": "v1.3.5",
        "date": "2026-09-25",
        "tag": "",
        "items": [
            ('📑', '修复：全文摘要章节输出乱序', '章节原按得分高低输出（讨论 → 结果 → 引言 → 方法）。现名额分配仍按章节权重，但展示顺序改为遵循论文逻辑（引言 → 方法 → 结果 → 讨论 → 结论）；同时清理抽取句开头的承接连接词（此外，/ 然而，/ 因此，…），避免脱离上下文后语义悬空。'),
            ('🧹', '关键词去泛词与译文校验', 'percentage change、difficult、use 等无主题区分度的泛词不再入选；关键词译文增加合理性校验，拒绝翻译记忆库返回的整句垃圾与词典释义串。'),
        ],
    },
    {
        "version": "v1.3.4",
        "date": "2026-09-25",
        "tag": "",
        "items": [
            ('🌐', '网络报错中文化 + 重试加固', 'NCBI 请求重试升级为指数退避（2 / 4 / 6 / 8 秒），可度过瞬时 DNS 污染窗口；检索失败等网络异常改为可操作的中文提示（SSL 证书校验失败 / 连接失败 / 超时分别说明原因与处理建议），技术详情折叠展示。'),
        ],
    },
    {
        "version": "v1.3.3",
        "date": "2026-09-24",
        "tag": "",
        "items": [
            ('🖼', '修复：非开放获取文献图表抓取失败', 'Europe PMC 图片包对非 OA / 作者手稿文献不可用（实测 PMC5727893 即此场景）。新增三级兜底：逐图接口 → 无头浏览器抓取原图 → 图表页整页截图（带失败重试与本地缓存），实测 4/4 全部成功。'),
            ('🌐', '关键句支持中文翻译', '「句子重要性得分」中的关键句在输出语言为中文时自动翻译（带缓存），(Fig. 2) 等引用一并正确转换。'),
        ],
    },
    {
        "version": "v1.3.2",
        "date": "2026-09-24",
        "tag": "",
        "items": [
            ('🔬', '修复：视觉分析对推理模型输出为空', 'deepseek-flash 等推理模型的思考过程（reasoning）也计入 max_tokens，复杂看图请求会耗尽额度导致正文为空。现提高 token 上限至 8000，且正文为空时自动回退显示推理内容；已实测 deepseek-flash 看图分析正常。'),
            ('💡', '提示更新', '视觉分析按钮与错误提示标注 deepseek-flash 为可用模型（已实测支持图片输入）。'),
        ],
    },
    {
        "version": "v1.3.1",
        "date": "2026-09-24",
        "tag": "",
        "items": [
            ('🐛', '修复：摘要结果 KeyError 崩溃', '长时间运行的服务热重载界面但缓存旧版引擎模块，导致 source_count 键缺失而崩溃。现改为兜底读取，并重启本地服务。'),
        ],
    },
    {
        "version": "v1.3.0",
        "date": "2026-09-24",
        "tag": "",
        "items": [
            ('🔬', 'LLM 视觉分析图表图片', '新增「用 LLM 视觉分析图表图片」：把 PMC 图表图片缩放编码后连同说明文字一起送多模态模型（如 gpt-4o / qwen-vl），模型直接看图解读数据趋势与结论，不再只依赖说明文字；需配置支持图片输入的模型。'),
            ('🧮', '修复：摘要句子数与所选档位不符', '冗余句过滤可能少给句子，现改为三轮回填（严格去冗余 → 放宽阈值 → 按分数补齐），原文句子数足够时输出句数必然与所选档位一致；原文本身不足档位句数时，结果旁会明示实际输出 N 句（原文共 M 句，已全部纳入）。'),
            ('🈶', '修复：部分关键词未翻译为中文', '免费翻译接口对部分词会原样返回英文、或返回繁体（如 stem cells → 幹細胞）。现在会扫描接口的候选译文兜底，并将繁体自动转为简体（幹細胞 → 干细胞）。'),
        ],
    },
    {
        "version": "v1.2.3",
        "date": "2026-09-24",
        "tag": "",
        "items": [
            ('🧩', '修复：摘要出现无意义碎片', '切句时缩写保护失效，句中引用如 (Fig. 3) 会在 Fig. 处被误切断，留下悬空的「（图」，后半截被丢弃。现已正确保护 Fig. / et al. / e.g. 等缩写，超长句翻译分段也改为优先在标点处断开。'),
            ('🈶', '修复：关键词未翻译为中文', '输出语言选「中文」时，摘要页与全文分析的关键词现自动翻译为中文（内置 60+ 高频医学术语对照表优先命中，避免免费接口的词典误译如 cell → 單元格；未命中走翻译接口，失败回退英文）。'),
            ('🔧', '配套增强', '关键词译文带缓存，同一词不重复请求；翻译接口限速保护避免触发限制。'),
        ],
    },
    {
        "version": "v1.2.2",
        "date": "2026-09-24",
        "tag": "",
        "items": [
            ('🈶', '修复：全文摘要小标题语言错误', '「【Background】【Methods】」等英文章节小标题现统一显示为「背景 / 方法 / 结果 / 讨论 / 结论」等规范中文标题（支持带编号章节如 1. Introduction），且翻译时只翻译正文、不再把小标题混入翻译接口，输出语言不再错乱。'),
            ('🖼', '修复：无法抓取图表', 'Europe PMC 对部分文献返回 200 + XML 错误体而非图片包，旧版把错误体当图片包缓存导致永远报错。现在校验有效图片包后才落盘、坏缓存自动删除重下，并对非开放获取文献给出明确中文提示。'),
            ('🔧', '配套增强', '图片文件名匹配升级为大小写不敏感 + 归一化模糊匹配（fig1 ↔ F1.jpg 也能对上）；NCBI 接口请求增加自动重试，网络抖动不再直接失败；图表信息概括在选「中文」输出时同样自动翻译。'),
        ],
    },
    {
        "version": "v1.2.1",
        "date": "2026-09-23",
        "tag": "",
        "items": [
            ('🔑', '优化：关键词提取质量', '新增完整英文停用词表（功能词 / 代词 / 学术套话，如 the、her、new、study、figure 等一律不再入选）、名词复数归并（mutations 与 mutation 不再重复占位）、自动识别 overall survival、pd-l1 expression 这类医学短语；关键词后标注出现次数。'),
            ('🈶', '中文关键词升级', '改用极大频繁 n-gram 提取，输出「总生存期」这样的完整术语，不再出现「存期」「总生」半截词。'),
            ('📉', '配套优化', '摘要页关键词与全文关键词统一走同一套过滤与加权逻辑，全文分析与摘要结果口径一致。'),
        ],
    },
    {
        "version": "v1.2.0",
        "date": "2026-09-21",
        "tag": "",
        "items": [
            ('📚', '全文摘要', '对有 PMC 开放全文的文献，抓取论文正文（引言 / 方法 / 结果 / 讨论，数万字符）做章节化摘要，不再局限于摘要本身。'),
            ('📊', '全文数据分析面板', '全文词数统计、各章节篇幅分布图表、高频关键词提取、P 值 / 百分比 / 风险比 / 置信区间 / 样本量等统计指标自动提取。'),
            ('🈶', '全文摘要中文输出', '英文全文摘要自动翻译为中文（延续 v1.1.0 的翻译能力）。'),
        ],
    },
    {
        "version": "v1.1.1",
        "date": "2026-09-21",
        "tag": "",
        "items": [
            ('🧮', '修复：文本长度统计不准', '字符统计不再把空格和换行计入，英文文献同时显示词数，两种口径一目了然。'),
            ('🧠', '优化：摘要算法全面升级', '新增结构区块权重（Results / Conclusion 优先）、TF-IDF 词权重、标题相关性、数字结果句加分、冗余句去除，提取的核心信息更关键。'),
        ],
    },
    {
        "version": "v1.1.0",
        "date": "2026-09-21",
        "tag": "",
        "items": [
            ('🐛', '修复：内置摘要不支持中文', '英文文献的抽取式摘要现会自动翻译为中文输出（选择「中文」输出语言时自动生效），翻译失败时回退英文原句并提示。'),
        ],
    },
    {
        "version": "v1.0.1",
        "date": "2026-09-21",
        "tag": "",
        "items": [
            ('📏', '摘要长度可选', '短（约 3 句）/ 中（约 6 句）/ 长（约 10 句）三档自由切换，抽取式摘要与 LLM 总结均支持。'),
            ('📎', '原文文档链接', '检索结果与摘要页附加 PDF 全文（PMC）、DOI 原文、PubMed 页面直达链接。'),
            ('🖼', '文献图表解析', '抓取 PMC 开放全文中的图表图片与说明文字：图片视图逐图展示、内置引擎自动概括、LLM 逐图解读，图片本地缓存加速。'),
        ],
    },
    {
        "version": "v1.0.0",
        "date": "2026-09-21",
        "tag": "首个正式版",
        "items": [
            ('🔍', 'PubMed 智能检索', '关键词 / 作者 / 日期过滤，相关性或最新排序，拼写纠错建议。'),
            ('⚡', '抽取式智能摘要', '词频-位置加权引擎，离线可用，关键词提取与句子得分展示。'),
            ('🤖', 'LLM 深度总结', '可选配置 OpenAI 兼容接口，输出结构化中文总结。'),
            ('⭐', '收藏与历史', '收藏管理、Markdown 导出、检索历史自动留存。'),
        ],
    },
]


def render_footer():
    """全局页脚（v2.4.0）：医疗免责声明 + 数据来源署名 + 隐私说明。

    医疗类产品的底线件套——常驻显示在每页底部，避免任何「可直接用于诊疗」的误读。
    """
    st.markdown(
        """
        <div class="app-footer">
            <div class="ft-row">
                <b>⚕️ 医疗免责声明</b>　
                本工具生成的摘要、翻译与数据分析均由算法或大模型自动生成，可能存在遗漏、偏差或曲解。
                <span class="ft-warn"><b>不能作为临床诊断、用药或治疗方案的依据</b></span>；
                诊疗决策请务必核对文献原文，并以专业医师的判断为准。
            </div>
            <div class="ft-row">
                <b>📚 数据来源</b>　
                文献检索与全文来自 <b>PubMed / PMC</b>（美国国立卫生研究院 NCBI 下属国家医学图书馆 NLM），
                遵循 NCBI 使用条款，仅供学习与研究使用。
            </div>
            <div class="ft-row">
                <b>🔒 隐私说明</b>　
                本工具不收集任何个人身份信息；检索历史与收藏仅保存在本机（云端版为当前会话），不会上传至任何服务器。
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_changelog():
    for i, rel in enumerate(CHANGELOG):
        badge_color = "#2e9e8f" if i == 0 else "#8a97a5"
        tag_html = ""
        if rel.get("tag"):
            tag_html = (
                f'<span style="background:{badge_color}; color:#fff; border-radius:12px;'
                f'padding:2px 12px; font-size:0.8rem;">{rel["tag"]}</span>'
            )
        with st.container(border=True):
            badge_html = (
                f'<span style="background:{badge_color}; color:#fff; border-radius:12px;'
                f'padding:2px 12px; font-size:0.8rem; margin:0 4px;">{rel["tag"]}</span>'
                if rel.get("tag")
                else ""
            )
            st.markdown(
                '<div style="display:flex; align-items:center; gap:10px; flex-wrap:wrap;">'
                f"<span style=\"font-size:1.25rem; font-weight:700;\">{rel['version']}</span>"
                f"{badge_html}"
                f"<span style=\"color:#888; font-size:0.85rem;\">{rel['date']}</span>"
                "</div>",
                unsafe_allow_html=True,
            )
            for icon, title, desc in rel["items"]:
                st.markdown(f"- {icon} **{title}**：{desc}")


def _library_badges(pmid: str) -> None:
    """文献卡片上的分组 / 标签 / 笔记角标（没有标注时不占位）。"""
    if not pmid:
        return
    m = library.get_meta(pmid)
    if not m:
        return
    parts = []
    if m.get("folder"):
        parts.append("📁 " + (library.folder_name(m["folder"]) or "（已删除的分组）"))
    if m.get("tags"):
        parts.append(" ".join("#" + t for t in m["tags"]))
    if m.get("note"):
        parts.append("📝 有笔记")
    if parts:
        st.caption("　|　".join(parts))


def _drop_card_widgets(pmids) -> None:
    """清掉这些文献在卡片编辑区留下的控件状态（v3.1.1）。

    卡片上的分组下拉 / 标签框 / 笔记框都用 key 存自己的值，Streamlit 会在后续运行里
    沿用这份状态。批量整理或标签管理改了同一篇的标注之后，卡片上仍是**旧值**；
    用户顺手点一下卡片里的「保存」，就会把刚做好的批量结果覆盖回去。
    在批量改动后清掉状态，让它们在本轮重新按存储值初始化。
    """
    for p in pmids or []:
        for pre in ("lib_f_", "lib_t_", "lib_n_"):
            st.session_state.pop(f"{pre}{p}", None)


def library_editor(pmid: str) -> None:
    """文献库页卡片内的编辑区：分组 / 标签 / 笔记一次保存（v3.1.0，P3-C1）。"""
    m = library.get_meta(pmid)
    folders = library.list_folders()
    opts = [library.UNGROUPED] + [f["id"] for f in folders]
    names = {library.UNGROUPED: "（未分组）", **{f["id"]: f["name"] for f in folders}}
    cur = m.get("folder", library.UNGROUPED)
    # 控件状态的初始化**必须区分两种"不在选项里"**：
    # - 首次渲染时 session_state 里根本没有这个 key（读到 None）→ 要用已存的分组当初值。
    #   若这里一律归零成「未分组」，就会覆盖掉下面的 index=，卡片永远显示「未分组」，
    #   用户不改下拉直接点「保存」就会把分组**静默清空**（v3.1.1 修，截图时发现）。
    # - key 已存在但指向的分组被别处删了 → 才需要归位，否则 Streamlit 会因
    #   「控件值不在选项内」直接抛异常。
    fkey = f"lib_f_{pmid}"
    if fkey not in st.session_state:
        st.session_state[fkey] = cur if cur in opts else library.UNGROUPED
    elif st.session_state[fkey] not in opts:
        st.session_state[fkey] = library.UNGROUPED
    with st.expander("🏷️ 分组 / 标签 / 笔记", expanded=False):
        sel = st.selectbox(
            "分组", opts, index=opts.index(cur) if cur in opts else 0,
            format_func=lambda x: names.get(x, x), key=fkey,
        )
        tag_text = st.text_input(
            "标签（逗号分隔多个）", value=", ".join(m.get("tags", [])),
            placeholder="如：RCT, 心血管, 高证据等级", key=f"lib_t_{pmid}",
        )
        note = st.text_area(
            "笔记", value=m.get("note", ""), height=110,
            placeholder="记下纳入 / 排除的理由、样本量疑问、待复核的点……", key=f"lib_n_{pmid}",
        )
        if st.button("💾 保存", key=f"lib_save_{pmid}"):
            library.set_folder(pmid, sel)
            library.set_tags(pmid, library.parse_tags(tag_text))
            library.set_note(pmid, note)
            st.toast("已保存", icon="💾")
            st.rerun()
        if m.get("note_updated"):
            st.caption(f"笔记最后修改：{m['note_updated']}")


def article_card(a: dict, show_actions: bool = True, manage: bool = False):
    with st.container(border=True):
        col1, col2 = st.columns([5, 1])
        with col1:
            st.markdown(f"**{a['title']}**")
            meta = []
            if a.get("authors"):
                meta.append(", ".join(a["authors"][:4]) + (" 等" if len(a["authors"]) > 4 else ""))
            if a.get("journal"):
                meta.append(f"*{a['journal']}*")
            if a.get("year"):
                meta.append(a["year"])
            if meta:
                st.caption(" · ".join(meta))
            if a.get("pmid"):
                links = pubmed.get_pdf_links(a)
                link_md = "  |  ".join(f"[{name}]({u})" for name, u in links)
                st.caption(f"PMID: {a['pmid']}  |  {link_md}")
            _library_badges(a.get("pmid", ""))
        with col2:
            if show_actions and a.get("pmid"):
                if storage.is_favorited(a["pmid"]):
                    if st.button("取消收藏", key=f"unfav_{a['pmid']}"):
                        storage.remove_favorite(a["pmid"])
                        library.clear_meta([a["pmid"]])
                        st.rerun()
                    st.caption("⭐ 已收藏")
                elif st.button("⭐ 收藏", key=f"fav_{a['pmid']}"):
                    storage.add_favorite(a)
                    st.toast("已加入收藏", icon="⭐")
                    st.rerun()
        if manage and a.get("pmid"):
            library_editor(a["pmid"])
        if a.get("abstract"):
            with st.expander("📖 摘要全文", expanded=False):
                st.write(a["abstract"])
        mesh_block(a)


def mesh_block(a: dict):
    """展示文献的 MeSH 主题词（C5）。

    为什么要单独讲一句「最新文献通常还没有」：NLM 的人工标引滞后数月到一年，
    刚上线的文章 MeSH 字段是空的。不说明的话，用户会以为是解析失败。
    """
    headings = a.get("mesh") or []
    if not headings:
        return
    major = [m["heading"] for m in headings if m.get("major")]
    with st.expander(f"🏷️ MeSH 主题词（{len(headings)} 个，其中主要主题 {len(major)}）",
                     expanded=False):
        if major:
            st.markdown("**主要主题（Major Topic）**：" + "、".join(f"`{m}`" for m in major))
        rows = []
        for m in headings:
            quals = [q["name"] + ("*" if q.get("major") else "") for q in (m.get("qualifiers") or [])]
            rows.append({
                "主题词": ("★ " if m.get("major") else "") + m["heading"],
                "副主题词": "、".join(quals) if quals else "—",
                "MeSH UI": m.get("ui", ""),
            })
        st.dataframe(rows, hide_index=True, use_container_width=True,
                     column_config={"主题词": st.column_config.TextColumn(width="large")})
        st.caption(
            "带 ★ 的是本文的主要主题（Major Topic）；副主题词后带 * 表示该副主题词也是主要主题。"
            "NLM 的人工标引滞后数月到一年，**最新发表的文献通常还没有 MeSH 主题词**，属正常现象。"
        )


def ensure_results():
    return st.session_state.get("results", [])


# ---------------- 引用导出（v3.0.1，P3-C2） ----------------
def cite_export_block(articles: list[dict], key: str, *, expanded: bool = False,
                      plain: bool = False, title: str = "📇 引用导出",
                      hint: str = ""):
    """通用引用导出区：选格式 → 实时预览 → 下载。

    为什么不只给一个下载按钮：导出的参考文献是直接贴进论文的，用户需要先看到
    长什么样、缺不缺字段，再决定用哪种格式——所以预览是主功能，下载是顺手的。

    ``plain=True`` 时不套折叠面板（用于「我的文献库」这类导出本来就是主操作的页面）。
    """
    if not articles:
        return
    _LANG = {"bibtex": "bibtex", "ris": "text", "endnote": "text",
             "medline": "text", "vancouver": "text", "gbt7714": "text"}

    def _body():
        opts = {f"{f['label']}　·　{f['hint']}": f["key"] for f in cite.FORMATS}
        c1, c2 = st.columns([3, 2], gap="medium")
        with c1:
            pick = st.selectbox("导出格式", list(opts.keys()), key=f"cite_fmt_{key}")
        with c2:
            st.caption("　")
            st.caption(hint or f"共 {len(articles)} 篇，导出文件可直接导入参考文献管理器。")
        fmt = opts[pick]
        text = cite.render(articles, fmt)
        bad = [(i, m) for i, a in enumerate(articles, 1)
               if (m := cite.completeness(a))]
        if bad:
            st.warning(
                "以下记录缺少引用必需字段，导出后需要手动补齐（或回检索页重新抓取）："
                + "；".join(f"第 {i} 条缺「{'/'.join(m)}」" for i, m in bad[:6])
                + ("…" if len(bad) > 6 else "")
            )
        shown = text if len(text) <= 4000 else text[:4000] + "\n\n…（预览已截断，完整内容请下载）"
        st.code(shown, language=_LANG.get(fmt, "text"))
        st.download_button(
            f"⬇️ 下载 {cite.label(fmt)}（{len(articles)} 篇）",
            text.encode("utf-8"),
            file_name=cite.filename(articles, fmt),
            mime=cite.mime(fmt),
            use_container_width=True,
            key=f"cite_dl_{key}_{fmt}",
        )

    if plain:
        sec_title(title, hint or f"共 {len(articles)} 篇，可直接导入参考文献管理器")
        _body()
    else:
        with st.expander(title, expanded=expanded):
            _body()


# ---------------- 首页组件 ----------------
def sec_title(text: str, hint: str = ""):
    st.markdown(
        f'<div class="sec-title"><div class="bar"></div><div class="tx">{text}</div>'
        f'<div class="hint">{hint}</div></div>',
        unsafe_allow_html=True,
    )


# ---------------- 原文定位渲染（v2.1.0） ----------------
def _render_located_sentences(index: list[dict], sentences: list[str], show_heading_tags: bool = True):
    """逐句展示摘要句的原文出处：章节 · 句序 · 高亮原句（数值高亮）"""
    import html as _h

    for sent in sentences:
        s = sent.strip()
        if len(s) < 12:  # 过短的多为标题/列表行，不做定位
            continue
        if s.startswith(("#", "|", "- [ ]")):
            continue
        matches = locate.locate_sentence(s, index)
        if matches:
            m = matches[0]
            sc = float(m.get("score", 0.0))
            sc_tag = (
                f'<span class="loc-score">{"精确" if sc >= 0.99 else f"匹配 {sc:.2f}"}</span>'
                if show_heading_tags else ""
            )
            st.markdown(
                f'<div class="loc-sent"><span class="loc-tag">📍 {_h.escape(m["section"]) or "正文"} · 第 {m["pos"] + 1} 句</span>{sc_tag}'
                f'<br>{locate.highlight_values(m["text"])}</div>',
                unsafe_allow_html=True,
            )
        else:
            st.markdown(
                f'<div class="loc-sent"><span class="loc-tag">⚠️ 未找到可靠原文对应（可能为模型改写 / 概括）</span>'
                f'<br>{_h.escape(s)}</div>',
                unsafe_allow_html=True,
            )


def _render_value_locator(index: list[dict]):
    """关键数值 → 原文出处：按类型分组，数值在原句中高亮"""
    vals = locate.find_values(index)
    if not vals:
        st.caption("未在全文中扫描到常见统计数值（P 值 / CI / 百分比等）。")
        return
    total = sum(len(v) for v in vals.values())
    st.caption(f"共定位 {total} 处关键数值，每类最多展示 40 处；数值在原句中以橙色高亮。")
    for vtype, items in vals.items():
        with st.expander(f"{vtype} — 全文 {len(items)} 处", expanded=False):
            for it in items:
                st.markdown(
                    f'<div class="loc-sent"><span class="loc-tag">📍 {it["section"] or "正文"} · 第 {it["pos"] + 1} 句</span>'
                    f'{locate.highlight_values(it["sentence"])}</div>',
                    unsafe_allow_html=True,
                )


def _render_keyword_search(index: list[dict]):
    import html as _h

    kw_q = st.text_input(
        "输入关键词定位原文（中英文均可；多个词用空格分隔 = 需同时包含）",
        key="kw_locate_q",
        placeholder="例如：metformin HbA1c / 低血糖 / 95% CI",
    )
    q = (kw_q or "").strip()
    if not q:
        st.caption("输入后在全文原句中实时定位：显示所在章节、章节内句序与命中次数，关键词蓝色高亮。")
        return
    terms = locate.parse_terms(q)
    hits = locate.search_keyword(index, q, limit=80)
    if not hits:
        st.info(f"全文原句中未找到同时包含「{'」和「'.join(terms) if len(terms) > 1 else terms[0]}」的内容。"
                "可尝试换用英文原词（正文多为英文）。")
        return
    total_occ = sum(h["count"] for h in hits)
    shown = hits[:50]
    st.caption(
        f"共 {len(hits)} 句命中、出现 {total_occ} 次"
        + (f"，已展示前 {len(shown)} 句" if len(hits) > len(shown) else "")
        + "；按原文出现顺序排列。"
    )
    for h in shown:
        st.markdown(
            f'<div class="loc-sent"><span class="loc-tag">📍 {_h.escape(h["section"]) or "正文"} · 第 {h["pos"] + 1} 句 · 命中 {h["count"]} 次</span>'
            f'<br>{locate.highlight_query(h["text"], terms)}</div>',
            unsafe_allow_html=True,
        )


def stat_card(label: str, value, unit: str = ""):
    st.markdown(
        f'<div class="stat-card"><div class="k">{label}</div>'
        f'<div class="v">{value}</div><div class="u">{unit}</div></div>',
        unsafe_allow_html=True,
    )


def hero_illustration() -> str:
    """首页插画：纯内联 SVG 线描，随主题配色，无外部资源依赖"""
    return """
<svg viewBox="0 0 340 250" xmlns="http://www.w3.org/2000/svg" role="img" aria-label="文献检索与智能摘要示意图">
<defs>
<linearGradient id="hBlob" x1="0" y1="0" x2="1" y2="1">
<stop offset="0%" stop-color="#d6f1ea"/><stop offset="100%" stop-color="#dbe9fa"/>
</linearGradient>
<linearGradient id="hDoc" x1="0" y1="0" x2="0" y2="1">
<stop offset="0%" stop-color="#ffffff"/><stop offset="100%" stop-color="#f5fbfa"/>
</linearGradient>
<linearGradient id="hBar" x1="0" y1="1" x2="0" y2="0">
<stop offset="0%" stop-color="#2e9e8f"/><stop offset="100%" stop-color="#72cbbd"/>
</linearGradient>
</defs>
<circle cx="166" cy="126" r="104" fill="url(#hBlob)" opacity="0.5"/>
<circle cx="304" cy="46" r="15" fill="#dcecfb" opacity="0.8"/>
<circle cx="32" cy="200" r="10" fill="#d9f2ec" opacity="0.9"/>
<g transform="translate(52,42)">
<rect x="0" y="0" width="138" height="170" rx="16" fill="url(#hDoc)" stroke="#cfe7e3" stroke-width="1.6"/>
<rect x="22" y="30" width="56" height="10" rx="5" fill="#8ed3c6"/>
<rect x="22" y="52" width="94" height="6" rx="3" fill="#e2eef2"/>
<rect x="22" y="66" width="82" height="6" rx="3" fill="#e2eef2"/>
<rect x="22" y="80" width="94" height="6" rx="3" fill="#e2eef2"/>
<rect x="22" y="94" width="62" height="6" rx="3" fill="#e2eef2"/>
<rect x="22" y="118" width="94" height="6" rx="3" fill="#e2eef2"/>
<rect x="22" y="132" width="74" height="6" rx="3" fill="#e2eef2"/>
<g transform="translate(22,152)">
<rect x="0" y="-16" width="9" height="16" rx="3" fill="url(#hBar)"/>
<rect x="15" y="-26" width="9" height="26" rx="3" fill="url(#hBar)" opacity="0.82"/>
<rect x="30" y="-11" width="9" height="11" rx="3" fill="url(#hBar)" opacity="0.6"/>
<rect x="45" y="-21" width="9" height="21" rx="3" fill="url(#hBar)" opacity="0.72"/>
</g>
</g>
<g transform="translate(178,132)">
<circle cx="0" cy="0" r="42" fill="#ffffff" stroke="#2e9e8f" stroke-width="4.5"/>
<line x1="30" y1="30" x2="58" y2="58" stroke="#2e9e8f" stroke-width="9" stroke-linecap="round"/>
<path d="M-22 12 L-8 -6 L4 6 L18 -18" fill="none" stroke="#3d8fd1" stroke-width="3.4" stroke-linecap="round" stroke-linejoin="round"/>
<circle cx="-8" cy="-6" r="3.2" fill="#3d8fd1"/>
<circle cx="18" cy="-18" r="3.2" fill="#3d8fd1"/>
</g>
<g transform="translate(200,38)">
<rect x="0" y="0" width="92" height="30" rx="15" fill="#ffffff" stroke="#dcecea" stroke-width="1.4"/>
<text x="46" y="20" text-anchor="middle" font-family="system-ui,-apple-system,Segoe UI,sans-serif" font-size="12.5" font-weight="600" fill="#2e8f83">P &lt; 0.001</text>
</g>
<g transform="translate(248,190)">
<rect x="0" y="0" width="78" height="30" rx="15" fill="#ffffff" stroke="#dcecea" stroke-width="1.4"/>
<text x="39" y="20" text-anchor="middle" font-family="system-ui,-apple-system,Segoe UI,sans-serif" font-size="12.5" font-weight="600" fill="#3d8fd1">HR 0.59</text>
</g>
</svg>
"""


FEATURES = [
    {
        "icon": "🔍",
        "bg": "#eaf7f4",
        "title": "文献检索",
        "desc": "接入 PubMed 官方接口，支持关键词、作者、发表日期过滤与拼写纠错；"
                "可选 MeSH 词表联动，自动补齐同义词、提升召回，并展示 PubMed 实际执行的检索式。",
        "btn": "去检索",
        "target": "文献检索",
    },
    {
        "icon": "📝",
        "bg": "#eef6fd",
        "title": "智能摘要",
        "desc": "内置抽取式引擎离线可用，按章节权重提取关键句；配置 LLM 后可生成结构化深度总结。",
        "btn": "去摘要",
        "target": "智能摘要",
    },
    {
        "icon": "📄",
        "bg": "#eaf7f4",
        "title": "PDF 全文分析",
        "desc": "上传本地 PDF 论文，解析出章节与题录信息，接上全文摘要、原文定位、图表表格与引用导出。",
        "btn": "去上传",
        "target": "PDF 全文分析",
    },
    {
        "icon": "🧾",
        "bg": "#eef7f1",
        "title": "综述工作台",
        "desc": "写综述：多篇横向对比、结论冲突提示、PRISMA 筛选记录、初稿骨架（叙述段可选学术严谨 / 简明扼要、中文 / 英文与所用模型）。"
                "看证据：研究类型分层、证据等级参考、偏倚提示、临床适用性。",
        "btn": "去综述",
        "target": "综述工作台",
    },
    {
        "icon": "📊",
        "bg": "#f3f0fd",
        "title": "全文数据分析",
        "desc": "抓取论文正文，输出章节篇幅分布与高频关键词，自动提取 P 值、样本量、风险比等指标。",
        "btn": "去分析",
        "target": "智能摘要",
    },
    {
        "icon": "🖼",
        "bg": "#fdf4ea",
        "title": "图表解读",
        "desc": "解析全文中的图表图片与说明文字，逐图视图展示并自动概括，快速把握研究图示信息。",
        "btn": "去解读",
        "target": "智能摘要",
    },
]


def render_home():
    favs = storage.list_favorites()
    hist_all = storage.list_history(200)
    results = ensure_results()
    pmc_count = len([a for a in results if a.get("pmcid")])
    lstats = library.stats()

    # ---------- Hero ----------
    with st.container(key="hero"):
        left, right = st.columns([1.32, 1], gap="large")
        with left:
            st.markdown(
                """
                <div class="hero-wrap">
                    <span class="hero-pill">🩺 医学文献工作台 · 免费公开数据源</span>
                    <h1>从 PubMed 检索<br>到论文全文智能摘要</h1>
                    <p class="hero-sub">输入一个关键词，检索、正文抓取、结构化摘要、图表解读与数据分析一次完成，不必在多个网站之间来回切换。</p>
                    <div class="hero-tags">
                        <span>PubMed 官方数据源</span>
                        <span>离线抽取式摘要</span>
                        <span>本地 PDF 上传解析</span>
                        <span>论文全文解析</span>
                        <span>统计指标提取</span>
                        <span>LLM 深度总结（可选）</span>
                    </div>
                </div>
                """,
                unsafe_allow_html=True,
            )
        with right:
            st.markdown(f'<div class="hero-svg-box">{hero_illustration()}</div>', unsafe_allow_html=True)
        b1, b2, b3, _sp = st.columns([1, 1, 1, 1.7])
        with b1:
            st.button("🔍 开始检索", type="primary", width="stretch",
                      on_click=goto, args=("文献检索",))
        with b2:
            st.button("📝 生成摘要", width="stretch",
                      on_click=goto, args=("智能摘要",))
        with b3:
            st.button("📋 更新日志", width="stretch",
                      on_click=goto, args=("更新日志",))

    # ---------- 数据概览 ----------
    s1, s2, s3, s4 = st.columns(4, gap="medium")
    with s1:
        stat_card("已收藏文献", len(favs), "篇 · 支持 Markdown 导出")
    with s2:
        stat_card("累计检索次数", len(hist_all), "次 · 历史自动留存")
    with s3:
        stat_card("当前结果集", len(results), f"篇 · 其中 {pmc_count} 篇有开放全文")
    with s4:
        stat_card("当前版本", APP_VERSION, "新功能见「更新日志」")

    st.write("")

    # ---------- 核心能力 ----------
    sec_title("核心能力", "点击卡片按钮直接进入对应功能")
    # 每行 4 张（卡片可继续增加，超出的自动换行；此前用 zip 截断会静默少渲染）
    for start in range(0, len(FEATURES), 4):
        cols = st.columns(4, gap="medium")
        for col, feat in zip(cols, FEATURES[start:start + 4]):
            with col:
                with st.container(key=f"featcard_{feat['target']}_{feat['title']}"):
                    st.markdown(
                        f'<div class="fc-icon" style="background:{feat["bg"]}">{feat["icon"]}</div>'
                        f'<div class="fc-title">{feat["title"]}</div>'
                        f'<div class="fc-desc">{feat["desc"]}</div>',
                        unsafe_allow_html=True,
                    )
                    st.button(
                        feat["btn"], key=f"featbtn_{feat['title']}", width="stretch",
                        on_click=goto, args=(feat["target"],),
                    )
        st.write("")

    st.write("")

    # ---------- 快速开始 ----------
    sec_title("快速开始", "三步完成一次文献分析")
    STEPS = [
        ("1", "检索文献", "在「文献检索」输入英文关键词（如 immunotherapy lung cancer），按需筛选年份与排序方式。"),
        ("2", "生成摘要", "在「智能摘要」选择文献与长度档位，点「生成抽取式摘要」，英文结果会自动译成中文。"),
        ("3", "深入解读", "带 PMC 开放全文的文献可直接做全文摘要、数据分析与图表解读；已下载的 PDF 到「PDF 全文分析」上传，同样能得到章节摘要、原文定位与表格提取。"),
    ]
    scols = st.columns(3, gap="medium")
    for col, (num, title, desc) in zip(scols, STEPS):
        with col:
            st.markdown(
                f'<div class="step"><div class="num">{num}</div>'
                f'<div class="t">{title}</div><div class="d">{desc}</div></div>',
                unsafe_allow_html=True,
            )

    st.write("")

    # ---------- 最近动态 ----------
    sec_title("最近动态", "")
    r1, r2 = st.columns(2, gap="medium")
    with r1:
        with st.container(border=True):
            st.markdown("**🕘 最近检索**")
            if hist_all:
                for h in hist_all[:4]:
                    st.caption(f"{h['time']} — {h['query']}（{h['n_results']} 条结果）")
            else:
                st.caption("暂无检索记录，点上面「开始检索」试试。")
    with r2:
        with st.container(border=True):
            st.markdown("**⭐ 最近收藏**")
            if favs:
                for f in favs[:4]:
                    st.caption(f"{f.get('saved_at', '')} — {f.get('title', '')[:52]}")
            else:
                st.caption("暂无收藏，检索结果点「⭐ 收藏」即可保存。")


# ---------------- 页面：综述工作台（v3.0.0，P2 主线 A 综述化 + B 证据化） ----------------
# 为什么单独做一页而不是塞进「智能摘要」：综述是**多篇之间**的工作，
# 单篇摘要页的信息结构（一次选一篇）根本装不下横向对比与筛选记录。
RV_DEFAULTS = {
    "rv_topic": "",
    "rv_query": "",
    "rv_date_range": "",
    "rv_total": 0,
    "rv_retrieved": 0,
    "rv_dup": 0,
    "rv_ex_screen": 0,
    "rv_ex_full": 0,
    "rv_reasons": "",
    "rv_notes": "",
    "rv_picked": [],
    # C4.1-B：对比表的人工修正值，形如 {"<pmid或标题前缀>": {"样本量": "120", "结论": "..."}}
    "rv_corrections": {},
}


def _rv_seed():
    """把磁盘上的工作区状态灌进本次会话（只灌尚未存在的键，避免覆盖用户当前输入）。"""
    saved = storage.load_review_state()
    for k, _v in RV_DEFAULTS.items():
        if k in saved and k not in st.session_state:
            st.session_state[k] = saved[k]
    # 刚在「文献检索」跑过检索时，顺手把检索式与命中数带进来，省一次手抄
    if not st.session_state.get("rv_query") and st.session_state.get("last_query"):
        st.session_state["rv_query"] = st.session_state["last_query"]
    if not st.session_state.get("rv_total") and st.session_state.get("last_total"):
        st.session_state["rv_total"] = int(st.session_state["last_total"] or 0)
    if not st.session_state.get("rv_topic") and st.session_state.get("last_query"):
        # 检索式不能直接当主题：字段标签、引号与 AND/OR/NOT 写进标题就是一串
        # 逻辑语言（v3.8.0 用户反馈）。清洗成可读短语再带出，用户仍可自行修改。
        st.session_state["rv_topic"] = review.clean_topic(st.session_state["last_query"])


def _rv_persist():
    storage.save_review_state({k: st.session_state.get(k) for k in RV_DEFAULTS})


def _rv_pool(results: list[dict], favs: list[dict], extra: list[dict] | None = None) -> list[dict]:
    """合并检索结果、收藏与本地 PDF 文献，按 PMID（缺失时用标题前缀）去重，保持原顺序。

    `extra` 用于把「PDF 全文分析」页解析出的本地文献并入池子——它们没有 PMID，
    去重会自然落到「标题前缀」这条分支上，因此多篇标题不同的 PDF 互不干扰。
    """
    seen, out = set(), []
    for a in list(results) + list(favs) + list(extra or []):
        key = (a.get("pmid") or "").strip() or re.sub(
            r"[^a-z0-9\u4e00-\u9fff]", "", (a.get("title") or "").lower())[:60]
        if not key or key in seen:
            continue
        seen.add(key)
        out.append(a)
    return out


def _rv_zip(rows: list[dict], conflicts: list[dict], prisma_rec: dict | None,
            draft: str, topic: str, articles: list[dict] | None = None) -> bytes:
    """把综述工作台的产出打包成一个 zip：对比表 + 冲突核查 + 筛选记录 + 初稿骨架 + 引用文件。"""
    import zipfile

    buf = io.BytesIO()
    stamp = datetime.now().strftime("%Y%m%d")
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"表1_文献对比表_{stamp}.csv", review.comparison_csv(rows))
        z.writestr(f"表1_文献对比表_{stamp}.md", review.comparison_markdown(rows, topic))
        if conflicts:
            z.writestr(f"结论冲突核查_{stamp}.md", review.conflicts_markdown(conflicts, topic))
        z.writestr(f"表2_偏倚风险提示清单_{stamp}.md", review.bias_markdown(rows, topic))
        z.writestr(f"表3_临床适用性对照_{stamp}.md", review.applicability_markdown(rows, topic))
        z.writestr(f"表4_结构化评价自查清单_{stamp}.md", appraisal.appraisal_markdown(rows, topic))
        z.writestr(f"表5_GRADE分级自查表_{stamp}.md", appraisal.grade_markdown(rows, topic))
        if prisma_rec:
            z.writestr(f"文献筛选记录_PRISMA_{stamp}.md", review.prisma_markdown(prisma_rec, rows))
        if draft:
            z.writestr(f"综述初稿骨架_{stamp}.md", draft)
        z.writestr("参考文献_Vancouver.md", review.references_markdown(rows))
        if articles:
            # 直接给可导入文献管理器的成品，省去从 Markdown 里手抄一遍
            z.writestr(f"参考文献_{stamp}.bib", cite.to_bibtex(articles))
            z.writestr(f"参考文献_{stamp}.ris", cite.to_ris(articles))
        z.writestr("说明.txt", (
            "由「医学文献智能摘要与检索系统」自动生成。\n"
            f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
            f"综述主题：{topic or '（未填写）'}\n"
            f"纳入文献：{len(rows)} 篇\n\n"
            "内容由规则引擎从 PubMed 摘要自动抽取，可能存在遗漏或偏差，请核对原文后使用。\n"
            "「证据等级」为按研究设计粗略映射的牛津 CEBM 简化参考，「证据强度」为可解释加权提示，"
            "「偏倚提示」只反映摘要层面可见的线索——三者都不是正式分级或规范偏倚评估"
            "（RoB 2 / NOS / GRADE），不能直接写入方法学部分。\n"
            "表 4「结构化评价自查清单」按研究设计列出 RoB 2 / NOS / AMSTAR-2 等工具的"
            "信号问题，表 5 列出 GRADE 起始等级与降级 / 升级因素：两者都是**待你逐条回答**"
            "的清单，不是评价结论；正式评价请使用原版量表。\n"
        ))
    return buf.getvalue()


def render_review_page():
    import pandas as pd

    header("🧾 综述工作台",
           "多篇文献横向对比 · 结论冲突识别 · 证据与适用性核查 · PRISMA 筛选记录 · 综述初稿骨架")
    st.caption(
        "一份数据源，两种用法：**写综述**（横向对比 / 结论冲突 / 筛选记录 / 初稿骨架）"
        "与**看证据**（研究类型分层 / 证据等级参考 / 偏倚提示 / 临床适用性）。"
        "同一主题下结论打架会自动提示；所有产出都能导出成 Word 可用的 Markdown。"
        "本页全部为**离线规则引擎**，不调用大模型也不消耗任何额度（用 LLM 撰写叙述段需手动触发）。"
    )
    st.caption(
        "⚠️ 证据等级 / 证据强度 / 偏倚提示都是**筛查加速器**，不是正式分级或偏倚评估"
        "（不是 CEBM 分级、不是 GRADE、不是 RoB 2），只能用来决定先看哪几篇，"
        "最终判断请用规范工具评价全文后作出。"
        "「证据与适用性」标签页底部按研究设计提供了 **RoB 2 / NOS / AMSTAR-2 等结构化"
        "信号问题清单与 GRADE 降级 / 升级自查入口**，可直接照着逐条核对。"
    )
    _rv_seed()

    pool = _rv_pool(ensure_results(), storage.list_favorites(),
                    st.session_state.get("pdf_articles"))
    if not pool:
        st.info(
            "还没有可用文献。请先到「文献检索」页检索，或在文献卡片上点「⭐ 收藏」后再回来——"
            "综述工作台的数据来源就是检索结果、收藏，以及「PDF 全文分析」页解析出的本地文献。"
        )
        return

    # 抓到的全文按文献键缓存在会话里（P3-C8）：pool 每次重渲染都会重建对象
    # （收藏尤其如此），直接改对象存不下来，必须在这里回填。
    _rv_ft = st.session_state.setdefault("rv_fulltext", {})
    if _rv_ft:
        for _a in pool:
            _k = review.article_key(_a)
            if _k in _rv_ft and not review.has_fulltext(_a):
                _a["fulltext_sections"] = _rv_ft[_k]

    # ---------- 第 1 步：主题与文献选择 ----------
    sec_title("1️⃣ 确定主题并勾选纳入文献", f"可选文献 {len(pool)} 篇（检索结果 + 收藏 + 本地 PDF，已去重）")
    st.text_input(
        "综述主题 / 研究问题（会写进初稿标题与参考文献说明）",
        key="rv_topic",
        placeholder="例如：PD-1 抑制剂联合化疗在晚期非小细胞肺癌中的疗效与安全性",
    )

    prefill = st.session_state.pop("rv_prefill", None)
    if prefill is not None:
        # data_editor 会把用户编辑保存在自己的会话状态里，换了默认勾选值必须把这个状态丢掉，
        # 否则「全选 / 清空」点了没反应（表格仍显示上一次的勾选）。
        st.session_state.pop("rv_editor", None)
    picked = set(prefill if prefill is not None else (st.session_state.get("rv_picked") or []))
    sel_df = pd.DataFrame({
        "选择": [(a.get("pmid") or a.get("title", "")[:60]) in picked for a in pool],
        "标题": [a.get("title", "")[:90] for a in pool],
        "年份": [str(a.get("year") or "") for a in pool],
        "期刊": [a.get("journal") or "" for a in pool],
        "研究设计": [review.judge_design(a)[0] for a in pool],
        "摘要": ["有" if (a.get("abstract") or "").strip() else "无" for a in pool],
    })
    b1, b2, b3 = st.columns([1, 1, 3])
    if b1.button("✅ 全选", key="rv_all", use_container_width=True):
        st.session_state["rv_prefill"] = [a.get("pmid") or a.get("title", "")[:60] for a in pool]
        st.rerun()
    if b2.button("🧹 清空选择", key="rv_none", use_container_width=True):
        st.session_state["rv_prefill"] = []
        st.rerun()
    with b3:
        st.caption("勾选要纳入综述的文献（建议 5–30 篇：太少看不出差异，太多对比表会失去可读性）。")
    edited = st.data_editor(
        sel_df, hide_index=True, use_container_width=True, key="rv_editor",
        column_config={
            "选择": st.column_config.CheckboxColumn("纳入", help="勾选后进入对比与冲突分析", default=False),
            "标题": st.column_config.TextColumn("标题", width="large", disabled=True),
            "年份": st.column_config.TextColumn("年份", width="small", disabled=True),
            "期刊": st.column_config.TextColumn("期刊", disabled=True),
            "研究设计": st.column_config.TextColumn("研究设计（自动识别）", disabled=True),
            "摘要": st.column_config.TextColumn("摘要", width="small", disabled=True),
        },
    )
    try:
        sel_idx = [i for i, v in enumerate(edited["选择"].tolist()) if bool(v)]
    except Exception:
        sel_idx = []
    selected = [pool[i] for i in sel_idx]
    st.session_state["rv_picked"] = [a.get("pmid") or a.get("title", "")[:60] for a in selected]

    if not selected:
        st.warning("请至少勾选 1 篇文献，下方对比、冲突与初稿才会生成。")
        _rv_persist()
        return
    n_no_abs = len([a for a in selected if not review.source_text(a).strip()])
    if n_no_abs:
        st.caption(f"⚠️ 其中 {n_no_abs} 篇既无摘要也无全文，相关字段只能留空——抽取不到就留空，不做推测。")

    # ---------- P3-C8：可选「抓取全文」（默认只用摘要，用户显式点击才抓） ----------
    _ft_ready = sum(1 for a in selected if review.has_fulltext(a))
    with st.expander(f"📥 用全文抽取（可选）· 已就绪 {_ft_ready}/{len(selected)} 篇", expanded=False):
        st.caption(
            "默认从 **PubMed 摘要**抽取。抓取开放获取（PMC）全文后，样本量、主要终点、"
            "关键效应量、结论会改从**全文相应章节**抽取（通常更完整），并在对比表"
            "「数据来源」列标为 `全文`。**本地上传的 PDF 自带全文，无需抓取**。"
        )
        st.caption(
            "⚠️ 只有开放获取文献抓得到（约两三成，付费墙的抓不到）；抓不到的自动回落摘要，"
            "并列入下方「未抓到」清单。抓取逐篇请求、篇数多时耗时较久；结果有 7 天缓存，"
            "同一篇重复抓取不会再走网络。"
        )
        _fb1, _fb2 = st.columns([1, 3])
        if _fb1.button("📥 抓取所选全文", key="rv_fetch_ft", use_container_width=True):
            _todo = [a for a in selected
                     if not review.has_fulltext(a) and not a.get("is_pdf")
                     and (a.get("pmcid") or a.get("pmid"))]
            if not _todo:
                st.info("所选文献都已有全文来源（本地 PDF 或此前已抓取），无需再抓。")
            else:
                _prog = st.progress(0.0, text=f"准备并发抓取 {len(_todo)} 篇…")
                _ph = st.empty()

                def _ft_prog(done, total, _prog=_prog, _ph=_ph):
                    _prog.progress((done / total) if total else 1.0,
                                   text=f"并发抓取中 {done}/{total}（3 路）…")
                    _ph.caption("提示：命中 7 天缓存的篇目秒回；走无头浏览器兜底的篇目每篇数十秒。")

                _results, _fails = pubmed.fetch_fulltext_many(
                    _todo, workers=3, on_progress=_ft_prog)
                _ok = 0
                _by_src = {}
                for _a, _secs, _src in _results:
                    if _secs:
                        _rv_ft[review.article_key(_a)] = _secs
                        _a["fulltext_sections"] = _secs
                        _a["fulltext_source"] = _src
                        _ok += 1
                        _by_src[_src] = _by_src.get(_src, 0) + 1
                _ph.empty()
                _prog.progress(1.0, text="抓取完成")
                st.session_state["rv_ft_result"] = {
                    "attempted": len(_todo), "ok": _ok, "by_src": _by_src,
                    "fails": [((_a.get("title") or ""), _why) for _a, _why in _fails],
                }
                st.session_state.pop("rv_sig", None)   # 强制按新文本重算对比表
                st.rerun()
        with _fb2:
            st.caption("抓取只改变**抽取来源**，不会改变纳入文献的范围，也不消耗大模型额度；"
                       "3 路并发进行，网络层已按 NCBI / Europe PMC / Unpaywall 域名限速。")
        _res = st.session_state.get("rv_ft_result")
        if _res:
            _ok, _fails = _res.get("ok", 0), (_res.get("fails") or [])
            _att = _res.get("attempted") or 0
            _with_ft = sum(1 for a in selected if review.has_fulltext(a))
            _cov = (_with_ft / len(selected) * 100) if selected else 0.0
            st.success(
                f"本次尝试 **{_att}** 篇，抓到全文 **{_ok}** 篇；未抓到 {len(_fails)} 篇——"
                "这些已自动用**摘要**完成抽取（对比表「数据来源」列为 `摘要`），对比与骨架不受影响。")
            _bysrc = _res.get("by_src") or {}
            if _bysrc:
                st.caption("来源构成：" + "、".join(f"{k} {v} 篇" for k, v in _bysrc.items()))
            st.caption(
                f"**当前全文覆盖率：{_cov:.0f}%**（{_with_ft}/{len(selected)} 篇有全文来源）。"
                "PMC 开放获取仅约两三成，付费墙文献无法免费获取全文——"
                "如需其全文，可在「PDF 全文分析」页上传 PDF（本地上传的 PDF 自带全文、无需抓取）。")
            if _fails:
                with st.expander(f"未抓到全文的文献（{len(_fails)} 篇）", expanded=False):
                    st.caption("这些多为付费墙文献，已回落摘要抽取；想用全文请改为上传 PDF。")
                    for _t, _why in _fails:
                        st.caption(f"· {(_t or '未命名')[:70]} — {_why}")

    # 抽取结果是纯本地计算，但没必要每次控件交互都重算（几十篇全量正则约数百毫秒），
    # 用"所选 PMID 签名 + 人工修正签名 + LLM 补抽签名"做记忆：勾选变化、手工改过
    # 字段或补抽结果变化时才重算。
    corrections = st.session_state.get("rv_corrections") or {}
    llm_fills = st.session_state.get("rv_llm_fills") or {}
    corr_sig = tuple(sorted((k, tuple(sorted((v or {}).items())))
                            for k, v in corrections.items()))
    fill_sig = tuple(sorted((k, tuple(sorted((v or {}).items())))
                            for k, v in llm_fills.items()))
    sig = (tuple(a.get("pmid") or a.get("title", "")[:40] for a in selected),
           corr_sig, fill_sig,
           # 全文抓取会改变抽取来源（摘要 → 全文），必须纳入签名，否则表格不重算
           tuple(sorted(review.article_key(a) for a in selected if review.has_fulltext(a))))
    if st.session_state.get("rv_sig") != sig:
        with st.spinner("正在抽取研究设计、样本量、效应量与结论……"):
            _rows = review.build_comparison(selected)
            _rows = review.apply_corrections(_rows, corrections)   # 人工修正值优先生效
            _rows = review.apply_llm_fill(_rows, llm_fills)        # LLM 补抽只补空格（C4.1-C）
            st.session_state["rv_rows"] = _rows
            st.session_state["rv_conflicts"] = review.detect_conflicts(_rows)
            st.session_state["rv_draft"] = ""
            st.session_state["rv_llm"] = ""
            # 对比表的默认值变了，旧编辑状态必须丢弃，否则表格仍显示上一版内容
            st.session_state.pop("rv_cmp_editor", None)
        st.session_state["rv_sig"] = sig
    rows = st.session_state.get("rv_rows", [])
    conflicts = st.session_state.get("rv_conflicts", {"terms": [], "conflicts": []})
    topic = st.session_state.get("rv_topic", "")

    t1, t2, t3, t4, t5 = st.tabs([
        "📊 横向对比表", "⚖️ 结论冲突核查", "🩺 证据与适用性",
        "🧾 筛选记录（PRISMA）", "📝 综述初稿骨架",
    ])
    # ---------- 表 1：横向对比 ----------
    with t1:
        stats = review.summary_stats(rows)
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            stat_card("纳入文献", stats["total"], "篇 · 可导出 CSV / Markdown")
        with c2:
            stat_card("最新发表年份", stats["year_to"] or "—",
                      f"区间 {stats['year_from'] or '—'}–{stats['year_to'] or '—'}")
        with c3:
            stat_card("样本量合计", f"{stats['n_total']:,}" if stats["n_total"] else "—",
                      f"例 · {stats['n_missing']} 篇未抽取到")
        with c4:
            top_design = next(iter(stats["designs"]), "—")
            stat_card("最常见设计", top_design, f"{stats['designs'].get(top_design, 0)} 篇")

        st.write("")
        cov = review.completeness_overview(rows)
        sec_title("表 1　纳入文献基本特征对比",
                  "字段自动抽取：已有全文的用全文（PMC 开放全文 / 本地上传 PDF），否则用 PubMed 摘要；"
                  "「证据强度」为可解释加权提示，不是正式证据分级")
        st.caption(
            f"**结构化完整度：平均 {cov['avg']}/{cov['max']}**"
            + (f"，其中 {cov['low']} 篇不足一半字段" if cov["low"] else "")
            + "。抽不到只说明来源文本没写或写法罕见，与研究的价值无关——"
              "这一步的用途是让你一眼看出哪几篇值得回原文补字段。"
        )
        if cov["missing"]:
            st.caption("最常缺的字段：" + "　".join(
                f"{k}（{v} 篇）" for k, v in list(cov["missing"].items())[:4]))
        _n_ne = sum(1 for r in rows for v in (r.get("_states") or {}).values()
                    if v == review.CELL_NOT_EXTRACTED)
        _n_nm = sum(1 for r in rows for v in (r.get("_states") or {}).values()
                    if v == review.CELL_NOT_MENTIONED)
        _n_ft = sum(1 for a in selected if review.has_fulltext(a))
        if _n_ne:
            st.caption(
                f"🔎 有 **{_n_ne}** 个格子是「有痕迹未抽到」（原文写了、正则没覆盖）。两种补法："
                "① 表格里直接手工填（带 ✍️ 的列可改）；"
                "② 用下方「🤖 让大模型补抽」按来源原文补抽，并附引句供核对。")
        if _n_nm and _n_ft < len(selected):
            st.caption(
                f"📄 另有 **{_n_nm}** 个格子是「原文确实没写」。若其来源是摘要，"
                "抓开放获取全文或上传 PDF 后常能补上——见上方「📥 用全文抽取（可选）」。")
        st.caption(
            "**空格的三种含义**（过去一律留白，看不出区别）："
            f"`{review.MARK_NOT_MENTIONED}` = 原文确实没写；"
            "`⚠️（有摘要未抽到）` = 原文里有痕迹但工具没抽出来，**建议回原文核对**；"
            f"`{review.MARK_NO_ABSTRACT}` = 没有摘要也没有全文。"
            "若「数据来源」列为 `全文`，同一套标注改用「原文未提及 / 有全文未抽到」措辞。"
        )
        n_corr = sum(len(r["_profile"].get("corrected") or []) for r in rows)
        if n_corr:
            st.caption(f"`✎已修正` = 你手工填的（当前共 {n_corr} 处，导出时按修正后的值输出）")
        shown = pd.DataFrame([
            {c: (review.annotate_cell(r.get(c, ""), (r.get("_states") or {}).get(c, ""),
                                      corrected=c in (r["_profile"].get("corrected") or []),
                                      llm=c in (r["_profile"].get("llm_filled") or {}),
                                      source=(r["_profile"].get("source") or ""))
                 if c in review.EDITABLE_COLUMNS else r.get(c, ""))
             for c in review.COMPARISON_COLUMNS}
            for r in rows
        ])
        edited_cmp = st.data_editor(
            shown, hide_index=True, use_container_width=True,
            height=min(640, 90 + 36 * len(rows)),
            key="rv_cmp_editor",
            column_config={
                "序号": st.column_config.NumberColumn("序号", width="small", disabled=True),
                "标题": st.column_config.TextColumn("标题", width="large", disabled=True),
                "年份": st.column_config.TextColumn("年份", width="small", disabled=True),
                "期刊": st.column_config.TextColumn("期刊", width="small", disabled=True),
                "研究设计": st.column_config.TextColumn("研究设计", width="small", disabled=True),
                "证据等级": st.column_config.TextColumn(
                    "等级（参考）", width="small", disabled=True,
                    help="按研究设计粗略对应牛津 CEBM 分级，未考虑偏倚等降级因素，不是正式分级"),
                "样本量": st.column_config.TextColumn(
                    "样本量", width="small",
                    help="✍️ 可手工修正：直接填数字（如 120）。改后「证据强度」会按新样本量重算"),
                "人群": st.column_config.TextColumn(
                    "人群", width="medium", help="✍️ 可手工修正：回原文核对后直接填写"),
                "主要终点": st.column_config.TextColumn(
                    "主要终点", width="medium", help="✍️ 可手工修正：回原文核对后直接填写"),
                "关键效应量": st.column_config.TextColumn(
                    "关键效应量", width="medium",
                    help="✍️ 可手工修正，例如 HR 0.72；95% CI 0.58–0.90；P=0.004"),
                "结论": st.column_config.TextColumn(
                    "结论（摘要原文）", width="large",
                    help="✍️ 可手工修正：改后「结论倾向」会按新结论重新判定"),
                "结论倾向": st.column_config.TextColumn("结论倾向", width="small", disabled=True),
                "证据强度": st.column_config.TextColumn("证据强度", width="small", disabled=True),
                "偏倚提示": st.column_config.TextColumn(
                    "偏倚提示", width="medium", disabled=True,
                    help="摘要层面可见的线索，不是 RoB 2 / NOS 评估结论；明细见「🩺 证据与适用性」"),
                "MeSH 主要主题": st.column_config.TextColumn(
                    "MeSH 主要主题", width="medium", disabled=True,
                    help="NLM 人工标引的主题词（★ 开头）。摘要里人群 / 疾病写得含糊时，"
                         "它是最可靠的补充——但最新发表的文献尚未标引，此列为空属正常现象"),
                "数据来源": st.column_config.TextColumn(
                    "数据来源", width="small", disabled=True,
                    help="该篇字段抽取所用文本：全文 = 已抓到全文（PMC 开放全文 / 本地上传 PDF，"
                         "提取更完整）；摘要 = 仅有 PubMed 摘要。在「📥 用全文抽取」里可对所选"
                         "文献抓取开放获取全文"),
            },
        )
        st.caption(
            "表格里 **带 ✍️ 的 5 列可以直接改**（回原文核对后填真实值）。"
            "改「样本量」会重算证据强度、改「结论」会重判结论倾向；"
            "研究设计 / 证据等级 / MeSH / 偏倚提示读的是抽取来源文本（全文或摘要），**不随手工修正变化**，"
            "这四条要改请直接修正检索记录或按规范工具评价全文。"
        )
        # 把表格里的手工修改收进 rv_corrections：比对「本次渲染前」与「用户编辑后」两份表，
        # 有改动就落盘并重跑一次，让修正值真正进入对比行（导出 / 冲突识别 / 初稿都跟着走）。
        _new_corr = {k: dict(v) for k, v in (st.session_state.get("rv_corrections") or {}).items()}
        _changed = False
        _prev = shown.to_dict("records")
        if edited_cmp is not None and len(edited_cmp) == len(_prev):
            for _r, _a, _b in zip(rows, _prev, edited_cmp.to_dict("records")):
                _k = review.row_key(_r)
                _cur = dict(_new_corr.get(_k) or {})
                for _col in review.EDITABLE_COLUMNS:
                    _av, _bv = str(_a.get(_col, "")), str(_b.get(_col, ""))
                    if _av == _bv or review.strip_marker(_bv) == review.strip_marker(_av):
                        continue
                    _val = review.strip_marker(_bv)
                    if _val:
                        _cur[_col] = _val
                    else:
                        _cur.pop(_col, None)
                    _changed = True
                if _cur:
                    _new_corr[_k] = _cur
                else:
                    _new_corr.pop(_k, None)
        if _changed:
            st.session_state["rv_corrections"] = _new_corr
            _rv_persist()
            st.rerun()
        if n_corr:
            if st.button("↩️ 清除全部人工修正（回到自动抽取的结果）", key="rv_cmp_reset"):
                st.session_state["rv_corrections"] = {}
                _rv_persist()
                st.rerun()

        # ---------- C4.1-C：LLM 辅助补抽（可选，需侧边栏配置大模型） ----------
        _fill_targets = review.llm_fill_targets(rows)
        _n_fill_fields = sum(len(t["fields"]) for t in _fill_targets)
        if _fill_targets:
            with st.expander(
                f"🤖 让大模型补抽未识别字段（{len(_fill_targets)} 篇 · {_n_fill_fields} 个字段）",
                expanded=False,
            ):
                st.caption(
                    "「⚠️ 有摘要未抽到」说明正则没覆盖到该写法。这一步把**摘要原文**交给大模型"
                    "按相同列名补抽：每个补抽值都必须附带摘要原文引句（没有引句的值直接丢弃），"
                    "表格以 ⧉ 标注、下方「逐篇查看抽取依据」可核对引句；人工修正值始终优先，"
                    "导出物仍是不带标记的纯数据。需要侧边栏已配置大模型，按你自己的 Key 计费。"
                )
                if not llm_ready:
                    st.caption("⚪ 侧边栏未配置大模型，此步不可用；也可以直接在表格里手工填写。")
                else:
                    # 分批并发跑：串行时 200 篇要 25 批、动辄数分钟。这里先如实给出批次与
                    # 时间量级，并在运行时显示真实进度，避免"以为几十秒、结果几分钟没动静"。
                    _n_batches = (len(_fill_targets) + 7) // 8
                    _lo = max(1, round(_n_batches * 8 / 3 / 60))
                    _hi = max(_lo + 1, round(_n_batches * 25 / 3 / 60))
                    st.caption(
                        f"需分 **{_n_batches}** 批请求（每批 ≤8 篇、并发 3 路），"
                        f"预计约 **{_lo}–{_hi} 分钟**；批次结果会缓存，重复点击不重复计费。"
                        "篇数很多时，可先只勾选最需要补的几篇再补抽；"
                        "推理模型（如 deepseek-flash）思考较慢，换非推理模型会明显更快。"
                    )
                    if st.button(
                        f"🤖 补抽 {len(_fill_targets)} 篇文献的 {_n_fill_fields} 个字段",
                        key="rv_llm_fill_run",
                    ):
                        _ph = st.empty()

                        def _on_prog(done, total, _ph=_ph):
                            _ph.info(f"🤖 大模型补抽进行中：已完成 {done}/{total} 批"
                                     f"（每批 ≤8 篇、并发 3 路）…")

                        try:
                            with logger.span("综述补抽", 篇数=len(_fill_targets),
                                             字段数=_n_fill_fields):
                                fills = review.llm_fill_with_llm(
                                    st.session_state["llm_base"], st.session_state["llm_key"],
                                    st.session_state.get("llm_model") or "", rows, selected,
                                    on_progress=_on_prog)
                        except Exception as e:
                            logger.error("综述补抽失败", e)
                            _ph.empty()
                            st.error(f"补抽失败：{e}")
                        else:
                            _ph.empty()
                            if fills:
                                st.session_state["rv_llm_fills"] = {
                                    **(st.session_state.get("rv_llm_fills") or {}), **fills}
                                st.toast(
                                    f"补抽完成：{sum(len(v) for v in fills.values())} 个字段，请核对引句",
                                    icon="🤖")
                                st.rerun()
                            else:
                                st.info("模型没有给出可核对的补抽结果——很可能这些字段确实不在给定文本里。")
        if st.session_state.get("rv_llm_fills"):
            _fills_now = st.session_state["rv_llm_fills"]
            _fc1, _fc2 = st.columns([4, 1])
            with _fc1:
                st.caption(f"⧉ 已有 LLM 补抽 {sum(len(v) for v in _fills_now.values())} 个字段"
                           "（待核对；人工修正优先，引句见下方「逐篇查看抽取依据」）。")
            with _fc2:
                if st.button("🧹 清空补抽", key="rv_llm_fill_reset", use_container_width=True):
                    st.session_state["rv_llm_fills"] = {}
                    st.rerun()

        d1, d2 = st.columns(2)
        d1.download_button("⬇️ 导出对比表 (CSV，Excel 可直接打开)", review.comparison_csv(rows),
                           file_name="纳入文献对比表.csv", use_container_width=True)
        d2.download_button("⬇️ 导出对比表 (Markdown，可直接粘进 Word)",
                           review.comparison_markdown(rows, topic),
                           file_name="纳入文献对比表.md", use_container_width=True)

        with st.expander("🔍 逐篇查看抽取依据（每条字段都能回原文核对）", expanded=False):
            for r in rows:
                p = r["_profile"]
                _cov = review.structure_completeness(p)
                st.markdown(
                    f"**{r['序号']}. {p['title']}**　"
                    f"<span class='kw-chip'>{p['design']}</span>"
                    f"<span class='kw-chip'>{p['design_layer']}</span>"
                    f"<span class='kw-chip'>等级参考 {p['cebm'][0]}</span>"
                    f"<span class='kw-chip'>样本量 {p['n'] or '未抽取'}</span>"
                    f"<span class='kw-chip'>证据强度 {p['evidence']['label']}</span>"
                    f"<span class='kw-chip'>结构化完整度 {_cov['label']}</span>"
                    f"<span class='kw-chip'>{p['bias']['label']}</span>",
                    unsafe_allow_html=True,
                )
                if p.get("corrected"):
                    st.caption("✎ 你手工修正过：" + "、".join(p["corrected"]))
                if p.get("llm_filled"):
                    st.caption("⧉ LLM 补抽（引句供核对，与摘要原文比对后认可；人工修正始终优先）："
                               + "；".join(f"{f}「{q[:80]}」"
                                          for f, q in p["llm_filled"].items()))
                if _cov["missing"]:
                    _st = p.get("states") or {}
                    st.caption("未抽到的字段（" + "；".join(
                        f"{f} → {review.CELL_STATE_LABELS.get(_st.get(f, ''), '—')}"
                        for f in _cov["missing"]) + "）")
                st.caption(f"设计依据：{p['design_evidence']}")
                if p["n_evidence"]:
                    st.caption(f"样本量原文：{p['n_evidence']}")
                if p["population"]:
                    st.caption(f"人群：{p['population']}")
                if p["primary_outcome"]:
                    st.caption(f"主要终点：{p['primary_outcome']}")
                if p["effects"]:
                    st.caption("效应量原文：" + (p["effects"].get("snippet") or review.effect_text(p["effects"])))
                st.caption(f"结论（{p['conclusion_source']}）：{p['conclusion'] or '摘要未写明结论'}")
                st.caption(f"结论倾向判定：{p['polarity']}"
                           + (f"（命中线索「{p['polarity_cue']}」）" if p["polarity_cue"] else "（未命中线索词）"))
                st.caption(f"证据强度算法：{p['evidence']['basis']}")
                if p["bias"]["flags"]:
                    st.caption("偏倚提示：" + "；".join(
                        f"{f['label']}（{f['level']}）" for f in p["bias"]["flags"]))
                st.caption(f"临床适用性：{p['applicability']['note']}")
                st.divider()

    # ---------- 结论冲突核查 ----------
    with t2:
        sec_title("结论冲突核查", "同一主题词下结论相反时高亮提示，供人工核对，工具不做对错判断")
        terms = conflicts.get("terms") or []
        if terms:
            st.caption("本次文献共同关注的主题词：" + "　".join(
                f"`{t}`（{n} 篇）" for t, n in terms))
        dist = stats["polarity"] if rows else {}
        if dist:
            st.caption("结论倾向分布：" + "　".join(f"**{k}** {v} 篇" for k, v in dist.items()))
        if not conflicts.get("conflicts"):
            st.success("✅ 未发现同一主题下的明显结论冲突。")
            st.caption(
                "需要说明的是：未检出冲突 ≠ 结论一致。可能是摘要没写明结论、或主题词覆盖不足，"
                "仍建议人工通读各篇结论段——下方列出了结论倾向分布供快速扫读。"
            )
            for r in rows:
                p = r["_profile"]
                st.markdown(f"- **{r['年份'] or '年份不详'}｜{r['研究设计']}**：{p['conclusion'] or '摘要未写明结论'}")
        else:
            st.warning(
                f"发现 {len(conflicts['conflicts'])} 处可能的结论不一致。"
                "结论打架往往不是「有人做错了」，而是人群、剂量、终点定义或随访时长不同——请回原文核对后再判断。"
            )
            for i, cf in enumerate(conflicts["conflicts"], 1):
                with st.container(border=True):
                    st.markdown(
                        f"**{i}. [{cf['type']}] 主题词「{cf['term']}」**　"
                        f"<span class='kw-chip'>可信度 {cf['confidence']}</span>"
                        f"<span class='kw-chip'>涉及 {cf['count']} 篇</span>",
                        unsafe_allow_html=True,
                    )
                    cols = st.columns(len(cf["sides"]))
                    for col, side in zip(cols, cf["sides"]):
                        with col:
                            st.markdown(f"**{side['label']}**（{len(side['studies'])} 篇）")
                            for s in side["studies"]:
                                st.markdown(
                                    f"- {s['cite']}｜{s['design']}｜样本量 {s['n'] or '未抽取'}"
                                    + (f"｜[PMID {s['pmid']}](https://pubmed.ncbi.nlm.nih.gov/{s['pmid']}/)" if s["pmid"] else "")
                                )
                                st.caption(s["snippet"] or "（摘要未写明结论）")
                    st.caption("提示：" + cf["note"])
            st.download_button("⬇️ 导出结论冲突核查 (Markdown)",
                               review.conflicts_markdown(conflicts["conflicts"], topic),
                               file_name="结论冲突核查.md")

    # ---------- 证据与适用性（P2 主线 B：面向临床医生） ----------
    with t3:
        sec_title(
            "证据与适用性核查",
            "研究类型分层 · 证据等级参考 · 偏倚提示 · 临床适用性——只标出「需要核对什么」，不做「能不能用」的裁决",
        )
        st.caption(
            "本页面向临床场景：不替你判断这篇文献该不该用（工具读不到全文，也不认识你的患者），"
            "而是把**摘要层面能核实的线索**摆出来——样本量够不够、有没有对照、随访多长、"
            "终点是硬终点还是替代终点、以及结果是在谁身上得到的。每条提示都附摘要依据，"
            "**摘要没写的会明确标成「摘要未提及」，不会静默略过**。"
        )
        bov = review.bias_overview(rows)
        top_flag = next(iter(bov["flags"]), "—")
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            stat_card("纳入文献", bov["total"], "篇")
        with c2:
            stat_card("需重点核对", bov["focus_docs"], "篇 · 存在需要优先核对的提示项")
        with c3:
            stat_card("提示项合计", sum(bov["flags"].values()), f"条 · 覆盖 {len(bov['flags'])} 类")
        with c4:
            stat_card("最常见提示", top_flag if top_flag != "—" else "—",
                      f"{bov['flags'].get(top_flag, 0)} 篇" if top_flag != "—" else "本次未触发")

        with st.container(border=True):
            st.markdown(
                "**先说清楚这三样东西是什么，避免误用：**\n\n"
                "- **证据等级**：按研究设计粗略对应牛津 CEBM 2011 分级（1a 系统评价 → 5 机制研究），"
                "**未考虑**偏倚风险、间接性、不一致性等降级因素，不是正式分级；\n"
                "- **证据强度**（对比表）：设计 + 样本量 + 是否报告区间估计的可解释加权，**不是 GRADE**；\n"
                "- **偏倚提示**：只列摘要里看得见的线索，**不是** RoB 2 / NOS 的评估结论。\n\n"
                "三者都只能当作**筛查加速器**，最终判断请用规范工具评价全文后作出。"
            )

        st.write("")
        sec_title("表 2　证据特征与偏倚提示总览", "逐篇明细在下方展开区，每条都附原文依据")
        bio = pd.DataFrame([{
            "#": r["序号"],
            "标题": (r["_profile"]["title"] or "")[:70],
            "研究设计": r["_profile"]["design"],
            "类型层级": r["_profile"]["design_layer"],
            "证据等级": r["_profile"]["cebm"][0],
            "结构化完整度": review.structure_completeness(r["_profile"])["label"],
            "偏倚提示": r["_profile"]["bias"]["label"],
            "提示项": r["_profile"]["bias"]["brief"],
            # 「信息缺失」类只给条数（摘要没写 ≠ 有问题），逐条明细见下方展开区
            "信息缺失项": sum(1 for f in r["_profile"]["bias"]["flags"]
                          if f["level"] == review.BIAS_INFO),
            "主要终点类型": r["_profile"]["applicability"]["outcome_class"],
            "随访": r["_profile"]["applicability"]["followup"],
            "研究场景": r["_profile"]["applicability"]["setting"],
        } for r in rows])
        st.dataframe(
            bio, hide_index=True, use_container_width=True,
            height=min(640, 90 + 36 * len(rows)),
            column_config={
                "标题": st.column_config.TextColumn("标题", width="medium"),
                "类型层级": st.column_config.TextColumn("类型层级", width="small"),
                "证据等级": st.column_config.TextColumn("等级（参考）", width="small"),
                "结构化完整度": st.column_config.TextColumn(
                    "完整度", width="small",
                    help="6 个关键字段（设计 / 样本量 / 人群 / 主要终点 / 效应量 / 结论）抽到了几个"),
                "偏倚提示": st.column_config.TextColumn("偏倚提示", width="small"),
                "提示项": st.column_config.TextColumn("触发的主要提示", width="medium"),
                "信息缺失项": st.column_config.NumberColumn(
                    "信息缺失", width="small",
                    help="「摘要没写」类提示的条数（不等于研究有缺陷）。默认折叠，"
                         "在下方「逐篇证据明细」里按需展开"),
                "主要终点类型": st.column_config.TextColumn("主要终点类型", width="medium"),
                "研究场景": st.column_config.TextColumn("研究场景", width="small"),
            },
        )
        d1, d2 = st.columns(2)
        d1.download_button("⬇️ 导出偏倚风险提示清单 (Markdown)",
                           review.bias_markdown(rows, topic),
                           file_name="偏倚风险提示清单.md", use_container_width=True)
        d2.download_button("⬇️ 导出临床适用性对照 (Markdown)",
                           review.applicability_markdown(rows, topic),
                           file_name="临床适用性对照.md", use_container_width=True)

        with st.expander("🔬 逐篇证据明细（提示级别 + 判定理由 + 原文依据）", expanded=False):
            def _render_bias_flag(fl: dict) -> None:
                st.markdown(f"- **[{fl['level']}] {fl['label']}**：{fl['reason']}")
                if fl["evidence"]:
                    st.caption(f"　原文依据：「{fl['evidence']}」")
                else:
                    st.caption("　原文依据：摘要未提及（该条提示正是基于「未提及」本身）")

            for r in rows:
                p = r["_profile"]
                b = p["bias"]
                strong = [f for f in b["flags"] if f["level"] != review.BIAS_INFO]
                info = [f for f in b["flags"] if f["level"] == review.BIAS_INFO]
                st.markdown(
                    f"**{r['序号']}. {p['title']}**　"
                    f"<span class='kw-chip'>{p['design']}</span>"
                    f"<span class='kw-chip'>{p['design_layer']}</span>"
                    f"<span class='kw-chip'>等级参考 {p['cebm'][0]}</span>"
                    f"<span class='kw-chip'>{b['label']}</span>",
                    unsafe_allow_html=True,
                )
                st.caption(f"等级对应说明：{p['cebm'][1]}")
                if not b["flags"]:
                    st.caption("未自动发现需要提示的项——**不等于没有偏倚**，只是摘要层面看不出线索。")
                for fl in strong:
                    _render_bias_flag(fl)
                # 「信息缺失」档默认**折叠**，只显示条数：这类提示说的是"摘要没写"，
                # 不是"研究有问题"。平铺出来会让每篇都背一串无意义提示，反而把真正
                # 该看的那几条（重点核对 / 建议核对）稀释掉。
                if info:
                    if st.toggle(f"展开 {len(info)} 条「信息缺失」项（摘要没写 ≠ 有问题）",
                                 key=f"rv_binfo_{r['序号']}"):
                        for fl in info:
                            _render_bias_flag(fl)
                a = p["applicability"]
                st.markdown(
                    f"**适用性速览**：人群 `{a['population']}`　干预 `{a['intervention']}`　"
                    f"终点 `{a['outcome_class']}`　随访 `{a['followup']}`　场景 `{a['setting']}`"
                )
                st.caption(f"可及性提示：{a['availability']}")
                st.caption(a["note"])
                st.divider()

        # ---- P3-C4：结构化评价工具入口（必须回全文逐条回答的规范量表）----
        st.write("")
        sec_title(
            "🧰 结构化评价工具（按研究设计自动匹配）",
            "自动提示只能看摘要；下面这些是要回全文逐条回答的规范工具——只列问题，不替你判定",
        )
        st.caption(appraisal.APPRAISAL_CAVEAT)
        aov = appraisal.appraisal_overview(rows)
        k1, k2, k3 = st.columns(3)
        with k1:
            stat_card("匹配到工具", len(aov["toolkits"]), "种 · 按本批研究设计")
        with k2:
            top_tool = aov["toolkits"][0] if aov["toolkits"] else ("—", 0)
            stat_card("最常用工具", top_tool[0],
                      f"{top_tool[1]} 篇" if top_tool[1] else "本次未匹配")
        with k3:
            stat_card("未匹配到工具", aov["unmapped"], "篇 · 研究类型未识别")
        if aov["by_design"]:
            st.caption("本批研究设计分布：" + "、".join(
                f"{k} {v} 篇" for k, v in aov["by_design"].items()))
        _dr = review.design_report(rows)
        if _dr["total"]:
            st.caption(f"**研究类型识别率：{_dr['rate']:.0f}%**"
                       f"（{_dr['recognized']}/{_dr['total']} 篇已识别，"
                       f"{_dr['unrecognized']} 篇未识别）。识别优先取自 PubMed 文献类型，"
                       "其次标题 / 摘要 / 全文措辞。")
        if _dr["unrecognized"]:
            with st.expander(f"未被识别研究类型的文献（{_dr['unrecognized']} 篇）", expanded=False):
                st.caption("以下文献标题 / 摘要未含设计表述，且 PubMed 文献类型不足以判定"
                           "（或本地 PDF 无该字段）。可上传 PDF 后重试，或人工确认设计。")
                for _it in _dr["items"]:
                    _pt = "、".join(_it["pubtypes"]) if _it["pubtypes"] else "无"
                    st.caption(f"· {_it['title'][:70]}（{_it['year'] or '年份不详'}）"
                               f"｜PubMed 文献类型：{_pt}")
        _tk_count = dict(aov["toolkits"])
        _used_ids: list[str] = []
        for _r in rows:
            for _t in appraisal.tools_for_design((_r["_profile"] or {}).get("design", "")):
                if _t["id"] not in _used_ids:
                    _used_ids.append(_t["id"])
        if not _used_ids:
            st.info("本批文献的研究类型均未识别，无法匹配评价工具；请先人工确认研究设计。")
        for _tid in _used_ids:
            _t = appraisal.TOOLKITS[_tid]
            with st.expander(
                    f"📋 {_t['name']}　—　{_t['full']}（本批 {_tk_count.get(_t['name'], 0)} 篇适用）",
                    expanded=False):
                st.markdown(f"**适用**：{_t['scope']}")
                st.markdown(f"**判定方式**：{_t['output']}")
                if _t.get("note"):
                    st.info("提醒：" + _t["note"])
                for _group, _questions in _t["groups"]:
                    st.markdown(f"**{_group}**")
                    for _q in _questions:
                        st.markdown(f"- [ ] {_q}")
                st.caption("出处：" + _t["source"])
        st.download_button(
            "⬇️ 导出结构化评价自查清单 (Markdown)",
            appraisal.appraisal_markdown(rows, topic),
            file_name="结构化评价自查清单.md")

        # ---- P3-C4：GRADE 分级自查入口 ----
        st.write("")
        sec_title(
            "⚖️ GRADE 证据分级自查入口",
            "先按设计定起始等级，再逐条核对 5 个降级因素与 3 个升级因素——工具不代填最终等级",
        )
        st.caption(appraisal.GRADE_CAVEAT)
        gov = appraisal.grade_overview(rows)
        st.markdown("**① 起始等级分布**：" + "、".join(
            f"{k} {v} 篇" for k, v in gov["levels"])
            + (f"（其中 {gov['not_applicable']} 篇为 GRADE 不适用）" if gov["not_applicable"] else ""))
        _gdf = pd.DataFrame([
            {"研究设计": d, "起始等级": appraisal.grade_start_level(d)[0],
             "说明": appraisal.grade_start_level(d)[1]}
            for d in sorted({(_r["_profile"] or {}).get("design") or "未识别" for _r in rows})
        ])
        st.dataframe(_gdf, hide_index=True, use_container_width=True)
        st.markdown("**② 降级因素**（在起始等级上逐条核对，命中即降 1~2 级）")
        st.dataframe(
            pd.DataFrame([{"因素": n, "何时降级": w, "幅度": a, "怎么查": h}
                          for n, w, a, h in appraisal.GRADE_DOWNGRADE]),
            hide_index=True, use_container_width=True,
            column_config={
                "因素": st.column_config.TextColumn("因素", width="small"),
                "何时降级": st.column_config.TextColumn("何时降级", width="large"),
                "幅度": st.column_config.TextColumn("幅度", width="small"),
                "怎么查": st.column_config.TextColumn("怎么查", width="large"),
            })
        st.markdown("**③ 升级因素**（观察性证据体尤其适用，命中即升 1 级）")
        st.dataframe(
            pd.DataFrame([{"因素": n, "何时升级": w, "幅度": a, "怎么查": h}
                          for n, w, a, h in appraisal.GRADE_UPGRADE]),
            hide_index=True, use_container_width=True,
            column_config={
                "因素": st.column_config.TextColumn("因素", width="small"),
                "何时升级": st.column_config.TextColumn("何时升级", width="large"),
                "幅度": st.column_config.TextColumn("幅度", width="small"),
                "怎么查": st.column_config.TextColumn("怎么查", width="large"),
            })
        st.caption("提示：起始等级只是**起点**；同一结局的多项研究要合并到证据体层面调整，"
                   "不能把单篇的起始等级直接当成最终结论。")
        st.download_button(
            "⬇️ 导出 GRADE 分级自查表 (Markdown)",
            appraisal.grade_markdown(rows, topic),
            file_name="GRADE分级自查表.md")

        st.markdown("#### 🧭 向你的患者外推前，请逐维回答")
        st.caption("工具不知道你的患者是谁，所以不给「适用 / 不适用」的结论——只把该比对的维度列全，避免漏项。")
        for dim, q in review.APPLICABILITY_DIMENSIONS:
            st.markdown(f"- **{dim}**：{q}")

    # ---------- 筛选记录（PRISMA） ----------
    with t4:
        sec_title("筛选记录（PRISMA 式）", "自动带出检索与纳入数字，排除理由需人工确认后填写")
        auto_dup = review.auto_duplicates(selected)
        if auto_dup and not st.session_state.get("rv_dup"):
            st.session_state["rv_dup"] = auto_dup
        if not st.session_state.get("rv_retrieved"):
            st.session_state["rv_retrieved"] = len(selected)
        st.caption(
            f"来源说明：检索式与「数据库命中」已从「文献检索」页自动带出；"
            f"「实际纳入题录」默认取当前勾选的 {len(selected)} 篇；"
            f"按 PMID/标题自动检测到的重复为 {auto_dup} 条（已预填，可修改）。"
        )
        f1, f2 = st.columns(2)
        with f1:
            st.text_input("检索数据库", value="PubMed（NCBI E-utilities）", disabled=True, key="rv_db_show")
            st.text_input("检索式", key="rv_query", placeholder="在「文献检索」页检索后会自动带出")
            st.text_input("检索时限（如 2015/01/01–2026/10/01）", key="rv_date_range")
            st.number_input("数据库命中总数", min_value=0, step=1, key="rv_total")
        with f2:
            st.number_input("实际下载题录数", min_value=0, step=1, key="rv_retrieved")
            st.number_input("去重后剔除条数", min_value=0, step=1, key="rv_dup")
            st.number_input("阅读题名/摘要后排除", min_value=0, step=1, key="rv_ex_screen")
            st.number_input("全文评估后排除", min_value=0, step=1, key="rv_ex_full")
        st.text_area(
            "排除原因及篇数（每行一条，例如：非随机对照研究 12 篇）",
            key="rv_reasons", height=90,
            placeholder="重复发表 3 篇\n非目标人群 8 篇\n无法获取全文 2 篇",
        )
        st.text_area("补充说明（可写限定语种、手检补充等）", key="rv_notes", height=70)

        prisma_rec = {
            "topic": topic,
            "query": st.session_state.get("rv_query", ""),
            "database": "PubMed（NCBI E-utilities）",
            "date_range": st.session_state.get("rv_date_range", ""),
            "search_date": datetime.now().strftime("%Y-%m-%d"),
            "total_hits": st.session_state.get("rv_total", 0),
            "retrieved": st.session_state.get("rv_retrieved", 0),
            "duplicates": st.session_state.get("rv_dup", 0),
            "excluded_screening": st.session_state.get("rv_ex_screen", 0),
            "excluded_fulltext": st.session_state.get("rv_ex_full", 0),
            "excluded_reasons": [x.strip() for x in (st.session_state.get("rv_reasons") or "").splitlines() if x.strip()],
            "notes": st.session_state.get("rv_notes", ""),
        }
        cnt = review.prisma_counts(prisma_rec)
        m1, m2, m3, m4 = st.columns(4)
        with m1:
            stat_card("数据库命中", cnt["identified"], "条")
        with m2:
            stat_card("进入筛查", cnt["screened"], f"条 · 去重 {cnt['duplicates']}")
        with m3:
            stat_card("进入全文评估", cnt["assessed"], f"条 · 题摘排除 {cnt['excluded_screening']}")
        with m4:
            stat_card("最终纳入", cnt["included"], f"条 · 全文排除 {cnt['excluded_fulltext']}")
        md = review.prisma_markdown(prisma_rec, rows)
        st.markdown(md)
        st.download_button("⬇️ 导出筛选记录 (Markdown)", md,
                           file_name="文献筛选记录_PRISMA.md", use_container_width=True)
        st.caption("说明：排除理由与篇数由使用者负责填写核对，工具只保证各级数字在表内自洽。")

    # ---------- 综述初稿骨架 ----------
    with t5:
        sec_title("综述初稿骨架", "按「背景—方法—结果—讨论—结论」组织已读文献，事实来自抽取，空缺写成待补充")
        c1, c2 = st.columns([1, 2])
        with c1:
            if st.button("🧩 生成综述初稿骨架", type="primary", use_container_width=True):
                st.session_state["rv_draft"] = review.build_review_draft(
                    topic, rows, conflicts.get("conflicts", []), prisma_rec,
                    extra={"terms": conflicts.get("terms") or []},
                )
                st.toast("初稿骨架已生成（离线，未调用大模型）", icon="🧩")
        with c2:
            st.caption(
                "摘要、主要发现与局限性已由内置算法自动成稿（v3.8.0）；"
                "所有需要你亲自判断的地方一律写成 *<u>【待补充：…】</u>*（斜体下划线），"
                "与工具填好的事实一眼可分——不会误把机器填的内容当成自己写好的结论。"
            )
        draft = st.session_state.get("rv_draft", "")
        if draft:
            st.markdown(draft)
            df1, df2 = st.columns(2)
            df1.download_button("⬇️ 导出初稿骨架 (Markdown)", draft,
                                file_name="综述初稿骨架.md", use_container_width=True)
            df2.download_button(
                "📦 打包导出全部产出 (ZIP)",
                _rv_zip(rows, conflicts.get("conflicts", []), prisma_rec, draft, topic, selected),
                file_name=f"综述产出_{datetime.now().strftime('%Y%m%d')}.zip",
                use_container_width=True,
            )

        cite_export_block(
            selected, "rv",
            title="📇 纳入文献的引用导出（BibTeX / RIS / EndNote / Vancouver / GB/T 7714）",
            hint=f"共 {len(selected)} 篇纳入文献；ZIP 包里也已附带 .bib 与 .ris 两份。",
        )

        st.divider()
        st.markdown("#### 🤖 可选：让大模型撰写「结果概述」与「讨论」叙述段")
        st.caption(
            "只把工具已经抽取好的结构化事实（设计/样本量/效应量/结论/冲突点）交给模型，"
            "并在提示词里硬性禁止编造数据——事实缺失处要求写「摘要未提供」。使用你自己的 Key，按次计费。"
        )
        if not llm_ready:
            st.caption("⚪ 侧边栏未配置大模型，跳过此步也可直接使用上面的骨架。")
        else:
            _batched = len(rows) > review.DRAFT_BATCH_THRESHOLD
            if _batched:
                # 篇数很多时单次生成要覆盖全部研究，最容易触达模型输出上限。
                # 现在改为自动分批：各批并发写「结果概述」，再统一写一段「讨论」。
                _n_b = len(review.split_draft_batches(rows))
                st.caption(
                    f"🧩 本次纳入 **{len(rows)} 篇**，超过 {review.DRAFT_BATCH_THRESHOLD} 篇，"
                    f"将**自动分批生成**（{_n_b} 批「结果概述」并发 3 路撰写，再统一写一段「讨论」），"
                    "避免单次请求覆盖上百篇导致输出被截断。各批结果分别缓存，重复生成不重复计费。"
                )
            else:
                st.caption(
                    f"本次纳入 **{len(rows)} 篇**，单次生成即可覆盖；"
                    "如需更短篇幅，可改用「简明扼要」风格。"
                )
            # ---- P3-C6：风格 / 语言 / 模型 三个选择器 ----
            oc1, oc2, oc3 = st.columns([1, 1, 1.5])
            with oc1:
                style_key = st.selectbox(
                    "写作风格",
                    options=list(review.DRAFT_STYLE_KEYS),
                    format_func=lambda k: next(x["label"] for x in review.DRAFT_STYLES if x["key"] == k),
                    key="rv_llm_style",
                )
            with oc2:
                lang_key = st.selectbox(
                    "输出语言",
                    options=list(review.DRAFT_LANG_KEYS),
                    format_func=lambda k: next(x["label"] for x in review.DRAFT_LANGUAGES if x["key"] == k),
                    key="rv_llm_lang",
                )
            with oc3:
                # 刻意**不给这里加 format_func**：显示串一旦依赖会话状态（如把当前
                # 侧边栏模型名拼进去），两次渲染就会得到不同字符串，Streamlit 按
                # 显示串回查选项下标时会直接报「不在列表中」。当前生效的模型放在
                # 下面的 caption 里说明，不在下拉里动态拼。
                _model_opts = [FOLLOW_SIDEBAR, *_llm_model_choices()]
                if st.session_state.get("rv_llm_model") not in _model_opts:
                    # 候选项会随侧边栏「拉取模型列表」的结果变化。已选值若不在新
                    # 列表里，必须在控件创建前归位，否则回查下标会直接抛异常。
                    st.session_state["rv_llm_model"] = FOLLOW_SIDEBAR
                model_pick = st.selectbox(
                    "叙述段使用模型",
                    options=_model_opts,
                    key="rv_llm_model",
                )
            style_label = next(x["label"] for x in review.DRAFT_STYLES if x["key"] == style_key)
            lang_label = next(x["label"] for x in review.DRAFT_LANGUAGES if x["key"] == lang_key)
            style_hint = next(x["hint"] for x in review.DRAFT_STYLES if x["key"] == style_key)
            model_use = (st.session_state.get("llm_model") or "") if model_pick == FOLLOW_SIDEBAR else model_pick
            st.caption(
                f"当前设置：**{style_label}** · **{lang_label}** · 模型 `{model_use or '未设置'}`　"
                f"（{style_hint}）"
            )
            st.caption(
                "⚠️ 模型列表只是常用模型的快捷预设，**不代表这些模型都能用**——"
                "所选模型必须存在于你上面配置的 API 地址下，否则接口会直接报错；"
                "拿不准就保持「跟随侧边栏设置」。"
            )
            if st.button("🤖 撰写叙述段", use_container_width=False):
                _ph = st.empty()

                def _draft_prog(done, total, _ph=_ph, _batched=_batched):
                    if _batched:
                        _ph.info(f"🤖 叙述段分批生成中：已完成 {done}/{total} 步"
                                 "（各批「结果概述」并发 3 路，最后统一写「讨论」）…")
                    else:
                        _ph.info("🤖 大模型正在撰写叙述段…")

                try:
                    with logger.span("综述叙述生成", 主题=topic[:30], 篇数=len(rows),
                                     风格=style_key, 语言=lang_key, 模型=model_use):
                        st.session_state["rv_llm"] = review.draft_with_llm(
                            st.session_state["llm_base"], st.session_state["llm_key"],
                            model_use, topic, rows,
                            conflicts.get("conflicts", []),
                            style=style_key, language=lang_key,
                            on_progress=_draft_prog,
                        )
                        st.session_state["rv_llm_meta"] = {
                            "style": style_label, "lang": lang_label, "model": model_use,
                        }
                    _ph.empty()
                except Exception as e:
                    _ph.empty()
                    logger.error("综述叙述生成失败", e)
                    st.error(f"生成失败：{e}")
            llm_txt = st.session_state.get("rv_llm", "")
            if llm_txt:
                meta = st.session_state.get("rv_llm_meta") or {}
                stale = meta and (meta.get("style") != style_label or meta.get("lang") != lang_label)
                if meta:
                    st.caption(
                        f"以上叙述段由 **{meta.get('style', '')} · {meta.get('lang', '')}**"
                        f"（模型 `{meta.get('model', '')}`）生成"
                        + ("。⚠️ 你已改动上面的设置，重新点击按钮才会按新设置生成。" if stale else "。")
                    )
                with st.container(key="panel_llm"):
                    st.markdown(llm_txt)
                # 叙述段嵌进骨架对应章节（3.5 结果概述 / 4.4 讨论），不再是文末附录；
                # 骨架还没生成时按原样导出叙述段本身。
                merged = review.embed_narrative(draft or "", llm_txt)
                st.download_button("⬇️ 导出「骨架 + 叙述段」(Markdown)", merged,
                                   file_name=f"综述初稿_含叙述段_{meta.get('style', '')}_{meta.get('lang', '')}.md",
                                   use_container_width=True)
                st.caption("叙述段已按章节嵌入骨架：结果概述 → 3.5，讨论 → 4.4"
                           "（切不开标题时整体并入 4.4），不再是文末附录。")

    _rv_persist()


# ---------------- 页面：PDF 全文分析（v3.2.0，P3-C3 本地 PDF 上传解析） ----------------
def _pdf_shelf_ui():
    """本次会话已解析的 PDF 文献架。

    刻意不写盘：PDF 往往含未公开的全文，落盘会扩大数据面；会话结束即清空，
    与「上传不上传第三方、不收集信息」的隐私口径一致。要长期保留请用引用导出。
    """
    arts = st.session_state.get("pdf_articles") or []
    if not arts:
        return
    sec_title("🧰 本次会话的 PDF 文献架",
              f"共 {len(arts)} 篇 · 可到「综述工作台」一并纳入 · 仅存本次会话，关闭页面即清空")
    for i, a in enumerate(arts):
        with st.container(border=True):
            c1, c2 = st.columns([6, 1])
            with c1:
                st.markdown(f"**{a.get('title') or a.get('pdf_file') or '未命名'}**")
                bits = [x for x in [a.get("journal"), a.get("year"), a.get("pdf_file")] if x]
                st.caption(" · ".join(bits) or "—")
            with c2:
                if st.button("🗑 移除", key=f"pdf_shelf_del_{i}"):
                    st.session_state["pdf_articles"] = [
                        x for j, x in enumerate(arts) if j != i
                    ]
                    st.rerun()
    cite_export_block(arts, key="pdf_shelf", expanded=False,
                      title=f"📇 文献架引用导出（{len(arts)} 篇）",
                      hint=f"共 {len(arts)} 篇本地 PDF 文献，可导出为 BibTeX / RIS / EndNote 等格式。")


def render_pdf_page():
    import pandas as pd

    header("📄 PDF 全文分析",
           "上传本地 PDF 论文 → 章节与元数据解析 · 全文摘要与数据挖掘 · 原文定位 · 图表表格 · 引用导出")
    lim = pdfdoc.limits()

    # ---------- 1. 上传 ----------
    sec_title("1️⃣ 上传 PDF", "解析全程在内存中进行，文件不落盘、不上传第三方")
    with st.container(border=True):
        st.caption(
            f"单篇上限 **{lim['max_mb']} MB / {lim['max_pages']} 页**"
            f"（可用环境变量 `MEDLIT_PDF_MAX_MB` / `MEDLIT_PDF_MAX_PAGES` 调整）。"
            "支持带**文本层**的 PDF（Word / LaTeX / 出版社导出）；"
            "**扫描件与纯图片型 PDF 暂不支持**，上传后会有明确提示。"
        )
        st.caption(
            "🔒 文件只在本机 / 本会话**内存**中解析，不会保存到磁盘、不会上传到任何第三方。"
            "仅当你主动点击「LLM 深度总结」或「视觉分析」时，才会把相应文本 / 图片发送给你自己配置的大模型服务。"
        )
        agreed = st.checkbox(
            "**版权与合规确认**：我确认对所上传 PDF 拥有合法访问权，仅在个人学习 / 研究范围内使用；"
            "因上传、解析及后续使用所产生的一切版权与合规责任，由我自行承担。",
            key="pdf_copyright_ack",
        )
        if not agreed:
            st.info("请先勾选上方的版权与合规确认，再上传文件。")
        up = st.file_uploader(
            "选择 PDF 文件", type=["pdf"], key="pdf_uploader",
            disabled=not agreed, help="单篇论文；扫描件暂不支持",
        )

    _pdf_shelf_ui()

    if not agreed:
        return

    if up is None:
        # 用户移除了文件：连同解析结果一起清掉，避免展示上一次的残留
        if st.session_state.get("pdf_parsed"):
            for k in ("pdf_parsed", "pdf_sig", "pdf_bytes", "pdf_name", "pdf_article",
                      "pdf_error", "pdf_ft", "pdf_an", "pdf_llm", "pdf_vision"):
                st.session_state.pop(k, None)
        if not st.session_state.get("pdf_articles"):
            st.caption("👆 选择 PDF 后自动开始解析。")
        return

    data = up.getvalue()
    sig = f"{up.name}:{len(data)}"
    if st.session_state.get("pdf_sig") != sig:
        try:
            with st.spinner(f"正在解析 {up.name} ……"):
                with logger.span("PDF 解析", 文件=up.name, 大小=f"{len(data) / 1024:.0f}KB"):
                    parsed = pdfdoc.parse_pdf(data, up.name)
        except pdfdoc.PdfError as e:
            st.session_state.update({
                "pdf_sig": sig, "pdf_parsed": None, "pdf_bytes": b"",
                "pdf_name": up.name, "pdf_error": str(e),
            })
        except Exception as e:  # noqa: BLE001
            logger.error("PDF 解析未预期异常", e)
            st.session_state.update({
                "pdf_sig": sig, "pdf_parsed": None, "pdf_bytes": b"",
                "pdf_name": up.name,
                "pdf_error": "解析时发生未预期的错误，已记录到运行日志。请尝试另存为一份新的 PDF 后重试。",
            })
        else:
            st.session_state.update({
                "pdf_sig": sig, "pdf_parsed": parsed, "pdf_bytes": data,
                "pdf_name": up.name, "pdf_error": "",
                "pdf_article": pdfdoc.to_article(parsed, up.name),
            })
            # 换了文件，上一份的派生结果一律作废，防止串场
            for k in ("pdf_ft", "pdf_an", "pdf_llm", "pdf_vision"):
                st.session_state.pop(k, None)

    err = st.session_state.get("pdf_error", "")
    if err:
        st.error("❌ " + err)
        st.caption("可尝试：① 在 Word / LaTeX 中重新导出为带文本层的 PDF；② 用 OCR 工具处理后重传；"
                   "③ 确认文件未加密、未损坏。")
        return

    parsed = st.session_state.get("pdf_parsed")
    if not parsed:
        st.caption("👆 选择 PDF 后自动开始解析。")
        return

    _pdf_render_parsed(parsed, st.session_state.get("pdf_name", up.name))


def _pdf_render_parsed(parsed: dict, name: str):
    import pandas as pd

    meta = parsed.get("meta") or {}
    q = parsed.get("quality") or {}
    sections = parsed.get("sections") or []
    article = st.session_state.get("pdf_article") or pdfdoc.to_article(parsed, name)
    title = meta.get("title") or name

    # ---------- 2. 解析概览 ----------
    sec_title("2️⃣ 解析概览", "元数据 · 质量统计 · 解析提示")
    with st.container(border=True):
        st.markdown(f"### {title}")
        _sub = [x for x in ["、".join(meta.get("authors") or [])[:120],
                            meta.get("journal"), str(meta.get("year") or "")] if x]
        if _sub:
            st.caption(" · ".join(_sub))
        if meta.get("doi"):
            st.caption(f"DOI: {meta['doi']}")
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("页数", q.get("n_pages", 0))
        m2.metric("正文字符", f"{q.get('n_chars', 0):,}")
        m3.metric("章节标题", q.get("n_headings", 0))
        m4.metric("表格", q.get("n_tables", 0))
        for w in q.get("warnings") or []:
            st.info("ℹ️ " + w)
        if q.get("pages_with_images"):
            _pg = ", ".join(map(str, q["pages_with_images"][:20]))
            st.caption(f"🖼 检测到含图片的页：{_pg}" + ("…" if len(q["pages_with_images"]) > 20 else ""))

    b1, b2 = st.columns(2)
    with b1:
        st.download_button("⬇️ 导出解析全文 (Markdown)", pdfdoc.summary_markdown(parsed, name),
                           file_name="pdf_parsed.md", use_container_width=True)
    with b2:
        _shelf = st.session_state.get("pdf_articles") or []
        _dup = any((a.get("title") or "").strip() == (title or "").strip() for a in _shelf)
        if st.button("➕ 加入文献架（供综述工作台纳入）", key="pdf_add_shelf",
                     disabled=_dup, use_container_width=True):
            st.session_state["pdf_articles"] = _shelf + [article]
            st.toast("已加入文献架，可到「综述工作台」纳入该文献", icon="🧰")
            st.rerun()
        if _dup:
            st.caption("✅ 已在文献架中")

    # ---------- 3. 章节浏览 ----------
    sec_title("3️⃣ 章节浏览", f"共 {len(sections)} 节，点击展开阅读")
    for i, s in enumerate(sections):
        with st.expander(f"{i + 1}. {s['title']}（{len(s['text'])} 字符）"):
            st.markdown(s["text"])

    # ---------- 4. 全文摘要与数据分析 ----------
    st.divider()
    sec_title("4️⃣ 全文摘要与数据挖掘", "内置抽取式引擎，离线运行、不消耗任何额度")
    fc1, fc2, fc3 = st.columns([2, 2, 3])
    with fc1:
        pdf_lang = st.radio("输出语言", ["中文", "英文"], horizontal=True, key="pdf_lang")
    with fc2:
        ft_max = st.select_slider("摘要句子数", options=[6, 8, 10, 12, 15], value=8, key="pdf_ftmax")
    with fc3:
        st.caption("　")
        run_ft = st.button("📚 生成全文摘要", type="primary", use_container_width=True)

    if run_ft:
        with st.spinner("正在生成章节化摘要与数据挖掘……"):
            try:
                with logger.span("PDF 全文摘要", 文件=name, 档位=ft_max):
                    st.session_state["pdf_ft"] = summarizer.fulltext_summary(
                        sections, title=title, max_sentences=ft_max)
                    st.session_state["pdf_an"] = summarizer.analyze_fulltext(sections)
            except Exception as e:  # noqa: BLE001
                logger.error("PDF 全文摘要失败", e)
                st.error("全文摘要生成失败，问题已记录到日志。")

    ft = st.session_state.get("pdf_ft")
    an = st.session_state.get("pdf_an")
    if ft and an:
        if not ft.get("sections") and not ft.get("summary"):
            st.warning("未从正文中抽取到可用于摘要的有效章节（可能是正文过短或版面过于特殊）。")
        else:
            if pdf_lang == "中文" and ft.get("sections"):
                parts, failed = [], False
                for p in ft["sections"]:
                    body = " ".join(p["sentences"])
                    if summarizer.is_mostly_english(body):
                        try:
                            body = summarizer.translate_text(body, "en|zh-CN")
                        except Exception:
                            failed = True
                    parts.append(f"【{p.get('title_zh') or p['title']}】 {body}")
                ft_out = "\n\n".join(parts)
                if failed:
                    st.warning("部分章节翻译失败，对应小节已保留英文原句。")
            elif ft.get("sections"):
                ft_out = "\n\n".join(f"【{p['title']}】 " + " ".join(p["sentences"])
                                     for p in ft["sections"])
            else:
                ft_out = ft.get("summary", "")
            sec_title("📚 全文摘要", f"共 {ft.get('used_sections', 0)} 个章节纳入摘要")
            with st.container(key="panel_ft"):
                st.markdown(ft_out)
            st.download_button("⬇️ 导出全文摘要 (Markdown)", ft_out, file_name="pdf_summary.md")

            # ---- 原文定位 ----
            _loc_index = locate.build_index(sections)
            _loc_stat = locate.index_stats(_loc_index)
            sec_title("📍 原文定位",
                      f"摘要句子与关键数值回溯原文（索引 {_loc_stat['sentences']} 句 / {_loc_stat['sections']} 章节）")
            _render_keyword_search(_loc_index)
            with st.expander("📍 摘要句子 → 原文出处", expanded=False):
                for p in ft.get("sections", []):
                    st.markdown(f"**【{p.get('title_zh') or p.get('title', '')}】**")
                    _render_located_sentences(_loc_index, p.get("sentences", []), show_heading_tags=False)
            with st.expander("🔢 关键数值 → 原文出处", expanded=False):
                _render_value_locator(_loc_index)

            # ---- 数据分析 ----
            sec_title("📊 全文数据分析", "词数统计 · 章节篇幅 · 高频关键词 · 统计指标")
            d1, d2, d3 = st.columns(3)
            d1.metric("全文词数", f"{an['total_words']:,}")
            d2.metric("全文字符", f"{an['total_chars']:,}")
            d3.metric("章节数", len(an["section_stats"]))
            st.markdown("**各章节篇幅分布**")
            st.bar_chart({s["title"][:30]: s["words"] for s in an["section_stats"]})
            st.markdown("**🔑 全文高频关键词**（已过滤功能词与套话，括号内为出现次数）")
            kc = an.get("keyword_counts", {})
            ft_kws = an["keywords"]
            if pdf_lang == "中文" and ft_kws:
                with st.spinner("正在翻译关键词为中文……"):
                    ft_kws = summarizer.translate_keywords(ft_kws)
            st.markdown("".join(
                f'<span class="kw-chip">{k} <b>{kc[k]}</b></span>' if k in kc
                else f'<span class="kw-chip">{k}</span>' for k in ft_kws),
                unsafe_allow_html=True)
            st.markdown("**🔬 关键统计指标提取**")
            if an["metrics"]:
                st.caption("「N 处」为全文出现总次数；下方展示去重后的不同取值（最多 12 种）。")
                for _nm, _m in an["metrics"].items():
                    with st.expander(f"{_nm} — 全文 {_m['count']} 处 · 去重 {len(_m['samples'])} 种"):
                        st.markdown("".join(f'<span class="kw-chip">{v}</span>' for v in _m["samples"]),
                                    unsafe_allow_html=True)
            else:
                st.caption("未在全文中提取到常见统计指标。")

            # ---- 可选：LLM 深度总结 ----
            st.divider()
            sec_title("🤖 LLM 深度总结（可选）", "需先在左侧配置大模型；按次消耗你自己的额度")
            if not llm_ready:
                st.caption("⚪ 侧边栏未配置大模型，跳过此步也可使用上面的抽取式摘要。")
            else:
                if st.button("🤖 生成 LLM 深度总结", key="pdf_llm_btn"):
                    _full = "\n\n".join(f"【{s['title']}】{s['text']}" for s in sections)[:24000]
                    try:
                        with st.spinner("LLM 正在生成深度总结……"):
                            with logger.span("PDF LLM 总结", 模型=st.session_state.get("llm_model", ""),
                                             文件=name):
                                st.session_state["pdf_llm"] = summarizer.llm_summary(
                                    _full,
                                    st.session_state["llm_base"],
                                    st.session_state["llm_key"],
                                    st.session_state["llm_model"],
                                    language=pdf_lang,
                                    length_hint="中等，正文约 300–500 字，分点呈现研究设计 / 结果 / 局限",
                                )
                    except Exception as e:  # noqa: BLE001
                        logger.error("PDF LLM 总结失败", e)
                        st.error(f"LLM 调用失败：{e}（请检查 API 地址 / Key / 模型名与网络连通性）")
                _llm_txt = st.session_state.get("pdf_llm")
                if _llm_txt:
                    with st.container(key="panel_llm"):
                        st.markdown(_llm_txt)
                    with st.expander("📍 总结句子 → 原文出处", expanded=False):
                        _li = locate.build_index(sections)
                        _render_located_sentences(_li, summarizer.split_sentences(_llm_txt))
                    st.download_button("⬇️ 导出 LLM 总结 (Markdown)", _llm_txt,
                                       file_name="pdf_llm_summary.md")

    # ---------- 5. 图表与表格解析 ----------
    st.divider()
    sec_title("5️⃣ 图表与表格解析", "表格结构化提取 · 页面渲染看图 · 可选 LLM 视觉解读")
    tables = parsed.get("tables") or []
    if tables:
        st.markdown(f"**📊 表格（共 {len(tables)} 个）**")
        for t in tables:
            rows = t.get("rows") or []
            if not rows:
                continue
            ncol = max(len(r) for r in rows)
            hdr = [(str(c) if c not in (None, "") else f"列{j + 1}") for j, c in enumerate(rows[0])]
            hdr += [f"列{j + 1}" for j in range(len(hdr), ncol)]
            body = [list(r) + [None] * (ncol - len(r)) for r in rows[1:]]
            _label = f"第 {t.get('page')} 页 · 表 {t.get('index')}"
            with st.expander(f"📊 {_label}（{len(rows)} 行 × {ncol} 列）"):
                df = pd.DataFrame(body, columns=hdr)
                st.dataframe(df, use_container_width=True)
                st.download_button(
                    "⬇️ 下载 CSV", df.to_csv(index=False).encode("utf-8-sig"),
                    file_name=f"table_p{t.get('page')}_{t.get('index')}.csv",
                    key=f"pdf_tbl_{t.get('page')}_{t.get('index')}",
                )
    else:
        st.caption("未识别到可结构化提取的表格（部分 PDF 的表格是纯图形，无法直接抽取）。")

    st.markdown("**🖼 图形 / 图片（整页渲染）**")
    st.caption("PDF 中的图常由多个矢量元素拼成，无法直接抠图，这里按**整页渲染**查看，"
               "与「智能摘要」页的截图兜底思路一致。页数较多时请只勾选真正含图的页。")
    img_pages = q.get("pages_with_images") or []
    pick_pages = st.multiselect(
        "选择要渲染的页面",
        options=list(range(1, int(q.get("n_pages", 0)) + 1)),
        default=[p for p in img_pages[:8]],
        key="pdf_fig_pages",
        help="默认选中检测到图片的页；最多渲染 12 页",
    )
    if pick_pages:
        _bytes = st.session_state.get("pdf_bytes") or b""
        figs = []
        for pno in pick_pages[:12]:
            try:
                png = pdfdoc.render_page_png(_bytes, pno, resolution=140)
            except pdfdoc.PdfError as e:
                st.warning(f"第 {pno} 页渲染失败：{e}")
                continue
            figs.append({"label": f"第 {pno} 页", "caption": "", "data": png})
        for f in figs:
            with st.container(border=True):
                st.markdown(f"**{f['label']}**")
                st.image(f["data"], use_container_width=True)
        if figs:
            st.caption(f"已渲染 {len(figs)} 页" + ("（已截断到 12 页）" if len(pick_pages) > 12 else ""))
        if llm_ready and figs:
            if st.button("🔬 用 LLM 视觉分析所选页面", key="pdf_vision_btn"):
                try:
                    with st.spinner("LLM 正在逐页查看并分析（页数较多时约需 1–2 分钟）……"):
                        st.session_state["pdf_vision"] = summarizer.llm_figure_vision(
                            figs,
                            st.session_state["llm_base"],
                            st.session_state["llm_key"],
                            st.session_state["llm_model"],
                            language=pdf_lang,
                        )
                except Exception as e:  # noqa: BLE001
                    st.error(f"视觉分析失败：{e}（请确认所配置模型支持图片输入，如 deepseek-flash / gpt-4o / qwen-vl）")
        if st.session_state.get("pdf_vision"):
            with st.container(key="panel_vision"):
                sec_title("🔬 页面视觉分析", "LLM 多模态逐页看图解读")
                st.markdown(st.session_state["pdf_vision"])

    # ---------- 6. 引用导出 ----------
    st.divider()
    sec_title("6️⃣ 引用导出", "把这篇 PDF 文献导出为可直接导入文献管理器的引用")
    cite_export_block(
        [article], key="pdf_single", expanded=True, plain=True,
        title="📇 本篇文献引用导出",
        hint="字段来自 PDF 内嵌元数据与正文解析，请核对后再用于正式写作。",
    )


# ---------------- 首次进入：隐私条款确认（对应注册页的同意勾选） ----------------
# 设计取舍：不做强制弹窗打断，但第一次进入必须明确勾选才能操作，
# 勾选状态只存在本次会话（session_state），不写盘、不上传——与"不收集个人信息"一致。
if not st.session_state.get("privacy_ack"):
    with st.container(border=True):
        st.markdown("#### 使用前请先确认数据处理方式")
        st.caption(
            "本工具不收集姓名、手机号、邮箱等任何个人身份信息；检索历史与收藏仅保存在"
            "本机（云端为你本次会话专属目录）。唯一的例外是：翻译与大模型摘要请求会"
            "把文本发送至第三方服务，请勿在检索框粘贴患者身份信息。"
            "完整条款见左侧「🔒 隐私与数据」。"
        )
        c1, c2 = st.columns([3, 1])
        ack = c1.checkbox("我已阅读并同意上述数据处理方式", key="privacy_ack_box")
        if c2.button("查看完整政策", key="privacy_open"):
            st.session_state["pending_page"] = "隐私与数据"
            st.rerun()
        if ack:
            st.session_state["privacy_ack"] = True
            # 若本次是通过 ?page= 深链进来的，同意后要回到目标页，而不是被丢回首页
            _target = st.session_state.pop("deeplink_page", None)
            if _target and _target != "系统首页":
                st.session_state["pending_page"] = _target
            st.rerun()
        st.stop()


# ---------------- 页面：我的文献库（v3.1.0，P3-C1 管理化） ----------------
def _lib_export_md(favs: list[dict]) -> str:
    """把（当前筛选出的）收藏导出为 Markdown，附上分组 / 标签 / 笔记。"""
    metas = library.get_meta_bulk([a.get("pmid", "") for a in favs])
    fname = library.folder_index()
    blocks = []
    for a in favs:
        m = metas.get(a.get("pmid", ""), {})
        lines = [
            f"**{a.get('title', '')}**",
            f"- 作者：{', '.join(a.get('authors', []))}",
            f"- 期刊：{a.get('journal', '')} ({a.get('year', '')})",
            f"- PMID: {a.get('pmid', '')}  链接: {a.get('url', '')}",
        ]
        if m.get("folder"):
            lines.append(f"- 分组：{fname.get(m['folder'], '')}")
        if m.get("tags"):
            lines.append(f"- 标签：{'、'.join(m['tags'])}")
        if m.get("note"):
            lines.append(f"- 笔记：{m['note']}")
        lines.append(f"- 摘要：{a.get('abstract', '')}")
        blocks.append("\n".join(lines))
    return "\n\n---\n\n".join(blocks)


def render_library_page():
    header("⭐ 我的文献库",
           "收藏的文献按分组、标签与笔记管理；支持筛选、批量整理与引用格式导出")
    favs = storage.list_favorites()
    if not favs:
        st.info("暂无收藏。去「文献检索」页点击 ⭐ 收藏文献吧——收藏后就能在这里分组、打标签、写笔记。")
        return

    pmids = [f.get("pmid", "") for f in favs if f.get("pmid")]
    library.prune(pmids)            # 顺手清掉已取消收藏的文献留下的标注
    metas = library.get_meta_bulk(pmids)
    folders = library.list_folders()
    fname = {f["id"]: f["name"] for f in folders}
    folder_opts = [library.UNGROUPED] + [f["id"] for f in folders]
    folder_names = {library.UNGROUPED: "（未分组）", **fname}
    tags = library.all_tags()
    stt = library.stats()

    c1, c2, c3, c4 = st.columns(4, gap="medium")
    with c1:
        stat_card("收藏文献", len(favs), "篇")
    with c2:
        stat_card("分组", len(folders), "个 · 可增删改")
    with c3:
        stat_card("标签", len(tags), "种 · 可跨组筛选")
    with c4:
        stat_card("已写笔记", stt["noted"], "篇 · 只存本机")

    # ---------------- 分组管理 ----------------
    with st.expander("📁 分组管理（新建 / 重命名 / 删除）", expanded=False):
        nc1, nc2 = st.columns([3, 1])
        new_name = nc1.text_input("新建分组", placeholder="如：综述选题 A ／ 机制研究 ／ 待精读",
                                  key="lib_new_folder")
        if nc2.button("➕ 新建", key="lib_add_folder", use_container_width=True):
            ok, msg, folder = library.add_folder(new_name)
            if ok:
                st.toast(f"已新建分组「{folder['name']}」", icon="📁")
                st.rerun()
            else:
                st.warning(msg)
        if folders:
            st.caption("删除分组**不会删除文献**——组内文献退回「未分组」，标签与笔记保留。")
            for f in folders:
                fc1, fc2, fc3 = st.columns([3, 1, 1])
                newval = fc1.text_input("重命名分组", value=f["name"], key=f"lib_rn_{f['id']}",
                                        label_visibility="collapsed")
                if fc2.button("重命名", key=f"lib_rnb_{f['id']}", use_container_width=True):
                    ok, msg = library.rename_folder(f["id"], newval)
                    if ok:
                        st.rerun()
                    else:
                        st.warning(msg)
                if fc3.button("删除", key=f"lib_rmf_{f['id']}", use_container_width=True):
                    n = library.delete_folder(f["id"])
                    st.toast(f"已删除分组，{n} 篇文献退回未分组", icon="📁")
                    st.rerun()
        else:
            st.caption("还没有分组。分组是「一个课题 / 一篇综述」这一层，先把文献归到一起，之后再筛选与导出。")

    # ---------------- 标签管理 ----------------
    with st.expander("🏷️ 标签管理（重命名 / 合并 / 删除）", expanded=False):
        if tags:
            tag_names = [t for t, _ in tags]
            if st.session_state.get("lib_tag_pick") not in tag_names:
                st.session_state["lib_tag_pick"] = tag_names[0]
            tc1, tc2, tc3, tc4 = st.columns([2, 2, 1, 1])
            pick = tc1.selectbox(
                "标签", tag_names, key="lib_tag_pick",
                format_func=lambda t: f"{t}（{dict(tags).get(t, 0)} 篇）",
            )
            # 输入框的 key 带上所选标签，切换标签时输入框跟着刷新
            newtag = tc2.text_input("改名 / 合并到", value=pick, key=f"lib_tag_new_{pick}")
            if tc3.button("重命名", key="lib_tag_rn", use_container_width=True):
                ok, msg, cnt = library.rename_tag(pick, newtag)
                if ok:
                    _drop_card_widgets(pmids)      # 卡片的标签框要跟着刷新，否则保存会写回旧标签
                    st.toast(f"已更新 {cnt} 篇文献的标签", icon="🏷️")
                    st.rerun()
                else:
                    st.warning(msg)
            if tc4.button("删除", key="lib_tag_del", use_container_width=True):
                cnt = library.delete_tag(pick)
                _drop_card_widgets(pmids)
                st.toast(f"已从 {cnt} 篇文献移除标签「{pick}」", icon="🏷️")
                st.rerun()
            st.caption("标签总览：" + "　".join(f"{t}（{n}）" for t, n in tags))
        else:
            st.caption("还没有标签。在下面文献卡片的「🏷️ 分组 / 标签 / 笔记」里添加，"
                       "标签适合记录「研究类型 / 干预 / 人群」这类跨分组的维度。")

    # ---------------- 批量整理 ----------------
    with st.expander("🧰 批量整理（移动分组 / 打标签 / 移出收藏）", expanded=False):
        label_to_pmid = {
            f"{a.get('pmid', '')}｜{a.get('title', '')[:60]}": a.get("pmid", "")
            for a in favs if a.get("pmid")
        }
        picked = st.multiselect("选择文献", list(label_to_pmid.keys()), key="lib_bulk_pick",
                                placeholder="输入标题关键词搜索并勾选，可多选")
        picked_pmids = [label_to_pmid[k] for k in picked]
        if picked_pmids:
            b1, b2 = st.columns([2, 1])
            if st.session_state.get("lib_bulk_folder") not in folder_opts:
                st.session_state["lib_bulk_folder"] = library.UNGROUPED
            target = b1.selectbox("目标分组", folder_opts, key="lib_bulk_folder",
                                  format_func=lambda x: folder_names.get(x, x))
            if b2.button("➡️ 移动", key="lib_bulk_move", use_container_width=True):
                n = library.set_folder_bulk(picked_pmids, target)
                _drop_card_widgets(picked_pmids)
                st.toast(f"已移动 {n} 篇到「{folder_names.get(target, target)}」", icon="📁")
                st.rerun()
            t1, t2, t3 = st.columns([2, 1, 1])
            tag_in = t1.text_input("标签（逗号分隔）", key="lib_bulk_tags",
                                   placeholder="如：RCT, 心血管")
            if t2.button("➕ 添加标签", key="lib_bulk_addtag", use_container_width=True):
                n = library.add_tags(picked_pmids, library.parse_tags(tag_in))
                _drop_card_widgets(picked_pmids)
                st.toast(f"已为 {n} 篇添加标签", icon="🏷️")
                st.rerun()
            if t3.button("➖ 移除标签", key="lib_bulk_deltag", use_container_width=True):
                n = library.remove_tags(picked_pmids, library.parse_tags(tag_in))
                _drop_card_widgets(picked_pmids)
                st.toast(f"已从 {n} 篇移除标签", icon="🏷️")
                st.rerun()
            if st.button(f"🗑️ 把这 {len(picked_pmids)} 篇移出收藏", key="lib_bulk_unfav"):
                for p in picked_pmids:
                    storage.remove_favorite(p)
                library.clear_meta(picked_pmids)
                st.rerun()
        else:
            st.caption("先在上方搜索并勾选文献，再执行批量操作。")

    st.divider()
    sec_title("文献列表", "按分组 / 标签 / 关键词筛选；展开卡片上的「🏷️ 分组 / 标签 / 笔记」即可标注")

    fl1, fl2, fl3 = st.columns([2, 2, 3])
    filter_opts = ["__all__"] + folder_opts
    filter_names = {"__all__": "全部分组", **folder_names}
    if st.session_state.get("lib_filter_folder") not in filter_opts:
        st.session_state["lib_filter_folder"] = "__all__"
    cur_f = fl1.selectbox("分组", filter_opts, format_func=lambda x: filter_names.get(x, x),
                          key="lib_filter_folder")
    sel_tags = fl2.multiselect("标签（含全部所选）", [t for t, _ in tags], key="lib_filter_tags")
    kw = fl3.text_input("关键词", placeholder="匹配标题 / 作者 / 标签 / 笔记 / 分组名",
                        key="lib_filter_kw")

    def _hit(a: dict) -> bool:
        m = metas.get(a.get("pmid", ""), {})
        if cur_f == library.UNGROUPED:
            if m.get("folder"):
                return False
        elif cur_f != "__all__" and m.get("folder") != cur_f:
            return False
        if sel_tags:
            have = {t.lower() for t in (m.get("tags") or [])}
            if not all(t.lower() in have for t in sel_tags):
                return False
        if kw:
            blob = " ".join([
                a.get("title", ""), " ".join(a.get("authors") or []),
                " ".join(m.get("tags") or []), m.get("note", ""),
                fname.get(m.get("folder", ""), ""),
            ]).lower()
            if kw.lower() not in blob:
                return False
        return True

    shown = [a for a in favs if _hit(a)]
    st.caption(f"共收藏 {len(favs)} 篇，当前筛选出 **{len(shown)}** 篇。")

    e1, e2, e3 = st.columns([2, 2, 1])
    with e1:
        st.download_button(
            f"⬇️ 导出筛选结果 (Markdown · {len(shown)} 篇)",
            _lib_export_md(shown), file_name="library.md", key="lib_dl_md",
            disabled=not shown, use_container_width=True,
        )
    with e2:
        st.caption("导出内容含分组、标签与笔记，可直接作为综述的文献清单底稿。")
    with e3:
        if st.button("🧹 清空", key="lib_clear_all", use_container_width=True,
                     help="清空全部收藏及其标注（不可撤销）"):
            for a in favs:
                storage.remove_favorite(a.get("pmid", ""))
            library.prune([])
            st.rerun()

    if not shown:
        st.info("当前筛选条件下没有文献，换个分组 / 标签或清空关键词试试。")
    for a in shown:
        article_card(a, manage=True)

    # 导出区放在列表**之后**：引用导出自带一段很长的代码预览，若夹在筛选区与结果列表
    # 之间，会把真正要看的文献卡片整段挤出首屏（v3.1.1 调整顺序）。
    if shown:
        cite_export_block(
            shown, "libs", plain=True,
            title="📇 引用导出（BibTeX / RIS / EndNote / Vancouver / GB/T 7714）",
            hint=f"导出当前筛选出的 {len(shown)} 篇，可直接导入 Zotero / EndNote / NoteExpress",
        )


# ---------------- MeSH 联动 / 检索式说明（供「文献检索」页调用） ----------------
# 这两个函数必须定义在下面的页面分发之前：分发块是模块级的 if/elif 序列，
# 若把 def 放在分发块中间，会同时踩两个坑——后续 elif 变成语法错误，
# 且页面渲染时函数尚未定义（NameError）。
def _render_mesh_report(report: dict, expanded: bool = False):
    """把 MeSH 分析报告渲染成 UI（识别到的主题词 + 同义词 + 扩展后的检索式）。"""
    concepts = report.get("concepts") or []
    matched = report.get("matched", 0)
    title = f"🧬 MeSH 概念识别：{matched}/{len(concepts)} 个词命中主题词"
    with st.expander(title, expanded=expanded):
        for c in concepts:
            rec = c.get("mesh")
            if not rec:
                st.markdown(f"- **{c['text']}** —— 未在 MeSH 中找到有把握的对应主题词，按原词检索")
                continue
            how = "标题精确匹配" if rec.get("matched_by") == "heading" else "命中入口词表（同义词）"
            kind = {"descriptor": "主题词", "supplemental-record": "补充概念记录（药物 / 化学物质）"}
            kind = kind.get(rec.get("type", ""), rec.get("type", ""))
            st.markdown(f"- **{c['text']}** → **{rec['heading']}**　"
                        f"<span style='opacity:.7'>{kind} · {how} · MeSH UI {rec['ui']}</span>",
                        unsafe_allow_html=True)
            if rec.get("scope_note"):
                st.caption(f"　　定义：{rec['scope_note']}")
            entries = rec.get("entry_terms") or []
            if len(entries) > 1:
                st.caption(f"　　同义词（NLM 入口词表，共 {len(entries) - 1} 个）："
                           + "、".join(entries[1:]))
            if rec.get("tree_numbers"):
                st.caption("　　树号：" + "、".join(rec["tree_numbers"][:3]))
        if matched:
            st.markdown("**扩展后的检索式**")
            st.code(mesh.expand_expression(report["keyword"], report), language="text")
            st.caption(
                "这串检索式把 MeSH 主题词与全部入口词显式 OR 进来。"
                "同义词来自 NLM 官方词表，未做任何增删；选择「自动扩展同义词」后它就是实际提交的检索式。"
            )
        else:
            st.caption(
                "没有找到有把握的对应主题词，因此不会做任何改写。"
                "常见原因：① 用的是缩写或口语写法（换成标准英文术语再试）；"
                "② 这是个组合概念或过新的概念，MeSH 尚未收录；"
                "③ 词本身是副主题词（如 therapy / diagnosis），它只能挂在主题词后组合使用。"
            )


def _render_pubmed_translation():
    """展示 PubMed 自动词表映射后「实际执行」的检索式。"""
    trans = pubmed.last_translation()
    if not trans:
        return
    typed = st.session_state.get("last_query", "")
    if trans.strip() == typed.strip():
        return
    with st.expander("🔎 PubMed 实际执行的检索式（自动词表映射）", expanded=False):
        st.code(trans, language="text")
        st.caption(
            "PubMed 会自动把关键词映射到 MeSH 主题词、规范作者与期刊名，上面是它真正执行的检索式。"
            "**如果这里只是把你的原话原样回显，说明自动映射没有生效**——"
            "最常见的原因是检索式里带了字段限定（如 `[tiab]`）或引号短语，"
            "它们会绕过自动词表映射，这时同义词就会丢失。"
        )


# ---------------- 页面：首页 ----------------
if page == "系统首页":
    render_home()

# ---------------- 页面：文献检索 ----------------
elif page == "文献检索":
    # 作者国籍/地区（PubMed 按作者单位 Affiliation 匹配英文国名）
    COUNTRIES = [
        ("不限", ""), ("中国", "China"), ("美国", "United States"), ("英国", "United Kingdom"),
        ("日本", "Japan"), ("韩国", "South Korea"), ("德国", "Germany"), ("法国", "France"),
        ("意大利", "Italy"), ("西班牙", "Spain"), ("加拿大", "Canada"), ("澳大利亚", "Australia"),
        ("印度", "India"), ("巴西", "Brazil"), ("荷兰", "Netherlands"), ("瑞士", "Switzerland"),
        ("瑞典", "Sweden"), ("土耳其", "Turkey"), ("伊朗", "Iran"), ("新加坡", "Singapore"),
    ]
    country_map = dict(COUNTRIES)

    header("🔍 文献检索", "主副关键词组合 · 来源期刊 / 作者 / 国籍筛选 · 日期范围与排序")

    # 会话丢失后自动回填（v3.7.1）：浏览器刷新 / 预览重载 / 桌面版重启会把
    # session_state 清空，输入框跟着清空。这里从最近一次成功检索的落盘条件
    # 恢复主副关键词与条数，用户可以在此基础上直接修改后再次检索。
    # （同一会话内 session_state["kw"] 已存在，注入不会发生，不打扰正在输入的内容）
    _last_search = storage.load_last_search() or {}
    if "kw" not in st.session_state and _last_search.get("keyword"):
        st.session_state["kw"] = _last_search["keyword"]
    if "kw2" not in st.session_state and _last_search.get("secondary"):
        st.session_state["kw2"] = _last_search["secondary"]
    if "kw_n" not in st.session_state:
        st.session_state["kw_n"] = int(_last_search.get("retmax") or 10)

    col_q, col_n = st.columns([4, 1])
    with col_q:
        keyword = st.text_input("主关键词（必填）", placeholder="例如：immunotherapy lung cancer", key="kw")
    with col_n:
        retmax = st.select_slider(
            "返回条数", options=[5, 10, 20, 30, 50, 100, 200], key="kw_n",
            help="综述工作台建议 ≥50：对比表需要足够多的文献才能看出结论分歧；"
                 "条数越多检索越慢（每 50 条约需数秒）。",
        )

    with st.expander("⚙️ 高级检索条件"):
        s1, s2 = st.columns([3, 2])
        with s1:
            secondary = st.text_input(
                "副关键词（可选，多个用逗号或分号分隔）",
                placeholder="例如：PD-1 biomarker, survival",
                key="kw2",
                help=(
                    "多个副关键词用逗号（，或,）或分号（；或;）分隔，各自独立成一个概念：\n\n"
                    "- **AND**：主关键词与**全部**副关键词都要出现\n"
                    "- **OR**：主关键词或任一副关键词出现即可\n"
                    "- **NOT**：含**任一**副关键词的文献全部排除\n\n"
                    "含空格的词组会自动加引号按精确短语匹配。"
                ),
            )
        with s2:
            logic = st.radio(
                "组合逻辑",
                ["AND（同时包含）", "OR（任一包含）", "NOT（排除）"],
                horizontal=True,
                index=0,
            )
        st.caption(
            "组合逻辑说明：多个副关键词用**逗号 / 分号**分隔，各自独立成一个概念——"
            "**AND** 主关键词与全部副关键词都要出现 · "
            "**OR** 主关键词或任一副关键词出现即可 · "
            "**NOT** 含任一副关键词的文献全部排除；含空格的词组自动按精确短语匹配"
        )
        c1, c2, c3 = st.columns(3)
        with c1:
            author = st.text_input("作者（可选）", placeholder="例如：Smith J")
        with c2:
            journal = st.text_input(
                "来源期刊（可选）",
                placeholder="例如：Nature Medicine 或 N Engl J Med",
                help="支持全名或缩写，全名会在检索时自动转换为 MEDLINE 缩写",
            )
        with c3:
            country_zh = st.selectbox(
                "作者国籍 / 地区（可选，按作者单位匹配）",
                [zh for zh, _ in COUNTRIES],
                index=0,
            )
        d1, d2, d3 = st.columns(3)
        with d1:
            start_date = st.text_input("起始日期 YYYY/MM/DD", placeholder="2022/01/01")
        with d2:
            end_date = st.text_input("结束日期 YYYY/MM/DD", placeholder="2026/09/21")
        with d3:
            sort_opt = st.radio(
                "排序方式",
                ["按相关性", "按发表时间（最新优先）"],
                horizontal=True,
            )

    # ---------------- MeSH 词表联动（v3.4.0，P3-C5） ----------------
    mesh_mode = st.radio(
        "🧬 MeSH 词表联动",
        ["关闭", "仅提示", "自动扩展同义词"],
        horizontal=True, index=1,
        help=(
            "PubMed 会对关键词做「自动词表映射」：如果你输入的词正好是 MeSH 主题词或它的入口词，"
            "它会自动把主题词一起搜。但这个机制有两个短板——① 口语写法（heart attack）与缩写常常映射不上；"
            "② 一旦你加了字段限定（如 aspirin[tiab]）或引号，映射会被完全绕过，同义词就此丢失。"
            "本功能把 MeSH 词表显式查出来补齐这部分召回。"
        ),
    )

    def _fresh_mesh_report():
        """只在关键词未变时复用上次的分析结果——否则会拿旧词的结果去构建新词的检索式。"""
        rep = st.session_state.get("mesh_report")
        if rep and st.session_state.get("mesh_report_for") == keyword.strip():
            return rep
        return None

    mesh_report = None
    if keyword.strip() and mesh_mode != "关闭":
        c1, c2 = st.columns([1, 3], gap="medium")
        with c1:
            if st.button("🔬 分析关键词", use_container_width=True,
                         help="查询 NLM MeSH 词表，看看你的词对应哪个主题词、有哪些同义词"):
                with st.spinner("正在查询 MeSH 词表……"):
                    mesh_report = mesh.analyze(keyword)
                st.session_state["mesh_report"] = mesh_report
                st.session_state["mesh_report_for"] = keyword.strip()
        mesh_report = _fresh_mesh_report()
        with c2:
            if mesh_report is None:
                st.caption("　")
                st.caption("点左侧按钮可查看关键词在 MeSH 词表中的对应主题词与同义词。")
            elif mesh_report.get("skipped"):
                st.caption(f"ℹ️ 未做 MeSH 分析：{mesh_report['skip_reason']}")

        if mesh_report and not mesh_report.get("skipped"):
            _render_mesh_report(mesh_report, expanded=(mesh_mode == "自动扩展同义词"))

    # 检索式实时预览（自动扩展模式下预览的就是最终会提交的检索式）
    logic_key = "AND" if logic.startswith("AND") else ("OR" if logic.startswith("OR") else "NOT")
    preview_kw = keyword
    _auto_report = _fresh_mesh_report() if mesh_mode == "自动扩展同义词" else None
    if _auto_report and _auto_report.get("matched"):
        preview_kw = mesh.expand_expression(keyword, _auto_report)
    if keyword.strip():
        st.caption(
            "检索式预览：`"
            + pubmed.build_query(
                preview_kw, secondary, logic_key, author, journal, country_map[country_zh], start_date, end_date
            )
            + "`"
        )

    if st.button("🔍 开始检索", type="primary", use_container_width=True):
        if not keyword.strip():
            st.warning("请输入主关键词")
        else:
            journal_input = journal.strip()
            if journal_input:
                # PubMed [Journal] 字段只匹配 MEDLINE 缩写，全名需先经 NLM Catalog 解析
                with st.spinner("正在解析期刊名称……"):
                    journal_ta = pubmed.resolve_journal_ta(journal_input)
                if journal_ta != journal_input:
                    st.caption(f"期刊「{journal_input}」→ 检索用缩写「{journal_ta}」")
            else:
                journal_ta = ""
            # 自动扩展模式：若还没分析过（用户直接点了检索），这里补一次
            kw_used = keyword
            if mesh_mode == "自动扩展同义词":
                rep = _fresh_mesh_report()
                if rep is None:
                    with st.spinner("正在查询 MeSH 词表……"):
                        rep = mesh.analyze(keyword)
                    st.session_state["mesh_report"] = rep
                    st.session_state["mesh_report_for"] = keyword.strip()
                if rep.get("matched"):
                    kw_used = mesh.expand_expression(keyword, rep)
                    st.caption(f"已用 MeSH 扩展检索式（{rep['matched']} 个概念命中主题词）。")
                else:
                    st.caption("关键词未能在 MeSH 词表中找到有把握的对应主题词，按原词检索。")
            query = pubmed.build_query(
                kw_used, secondary, logic_key, author, journal_ta, country_map[country_zh], start_date, end_date
            )
            with st.spinner("正在检索 PubMed ..."):
                try:
                    with logger.span("PubMed 检索", 检索式=query[:60], 条数上限=retmax):
                        results = pubmed.search_and_fetch(
                            query,
                            retmax=retmax,
                            sort="relevance" if sort_opt.startswith("按相关性") else "pub_date",
                        )
                    logger.info(f"检索完成：{len(results)} 条")
                except Exception as e:
                    logger.error("PubMed 检索失败", e)
                    st.error(f"检索失败：{pubmed.friendly_error(e)}")
                    with st.expander("技术详情"):
                        st.code(str(e)[:500])
                    results = []
            st.session_state["results"] = results
            st.session_state["last_query"] = query
            st.session_state["last_keyword"] = keyword.strip()
            # 记录本次检索在数据库中的命中总数（综述工作台的 PRISMA 记录要用）
            st.session_state["last_total"] = pubmed.last_total()
            # 落盘最近一次检索条件（v3.7.1）：会话丢失后检索页可自动回填
            storage.save_last_search({
                "keyword": keyword.strip(),
                "secondary": secondary.strip(),
                "retmax": int(retmax),
                "query": query,
                "total": int(st.session_state["last_total"] or 0),
            })
            if results:
                storage.add_history(query, len(results))
                feedback.bump_usage()  # 本机使用计数（v3.7.0）：只落盘 data/，绝不上传

    # 拼写建议（espell 接口，与上面的 MeSH 词表联动是两件事）
    if keyword.strip() and not ensure_results():
        for s in pubmed.spelling_suggest(keyword):
            st.info(f"💡 未找到匹配结果，是否想检索：**{s}**？")

    results = ensure_results()
    if st.session_state.get("last_query"):
        sec_title("检索结果", f"共 {len(results)} 篇")
        st.caption(f"检索式：`{st.session_state['last_query']}`")
        _render_pubmed_translation()
        cite_export_block(results, "search", hint="把本次检索结果整体导出，便于在 Zotero 里继续筛选。")
    for a in results:
        article_card(a)
    if not results and st.session_state.get("last_query"):
        st.info("没有检索到文献，试试更宽泛的关键词。")




# ---------------- 页面：智能摘要 ----------------
elif page == "智能摘要":
    header("📝 智能摘要", "抽取式摘要 · 全文摘要与数据分析 · 图表解读 · 可选 LLM 深度总结")
    sec_title("1️⃣ 选择摘要来源", "从最近检索结果中选择，或直接粘贴文本")
    with st.container(key="step-card"):
        source = st.radio(
            "来源",
            ["从最近检索结果中选择", "直接粘贴文本 / 摘要"],
            horizontal=True,
        )
        text = ""
        chosen_title = ""
        chosen_article = None
        run_ext = run_llm = False  # 无文本时保持未触发，供卡片外的结果渲染判断
        if source == "从最近检索结果中选择":
            results = [a for a in ensure_results() if a.get("abstract")]
            fav_results = [a for a in storage.list_favorites() if a.get("abstract")]
            options = [f"{a['title'][:60]}..." for a in results] + [f"⭐ {a['title'][:55]}..." for a in fav_results]
            pool = results + fav_results
            if not pool:
                st.info("暂无可用文献，请先在「文献检索」页检索，或粘贴文本。")
            else:
                idx = st.selectbox("选择文献", range(len(pool)), format_func=lambda i: options[i])
                text = pool[idx]["abstract"]
                chosen_title = pool[idx]["title"]
                chosen_article = pool[idx]
        else:
            pasted = st.text_area("粘贴文献摘要或全文片段", height=200, placeholder="在此粘贴英文或中文医学文献摘要……")
            text = pasted
            chosen_title = "（自定义文本）"

        # 文献全文文档链接
        if chosen_article:
            doc_links = pubmed.get_pdf_links(chosen_article)
            if doc_links:
                st.markdown("📎 **原文文档**：" + "  ·  ".join(f"[{n}]({u})" for n, u in doc_links))

        if text.strip():
            # 统计口径修正：字符数不含空白，英文按词计数
            clean = re.sub(r"\s+", "", text)
            words = len(re.findall(r"[A-Za-z][A-Za-z\-']*|[\u4e00-\u9fff]", text))
            st.caption(f"文本长度：{len(clean):,} 字符（不含空格换行） · 约 {words:,} 词")
            with st.container(key="step-params"):
                pc1, pc2 = st.columns(2)
                lang = pc1.radio("摘要输出语言", ["中文", "英文"], horizontal=True, index=0)
                length_label = pc2.radio(
                    "摘要长度",
                    ["短（约 3 句）", "中（约 6 句）", "长（约 10 句）"],
                    horizontal=True,
                    index=1,
                )
            length_map = {"短（约 3 句）": 3, "中（约 6 句）": 6, "长（约 10 句）": 10}
            llm_length_map = {
                "短（约 3 句）": "简短，正文约 150 字以内",
                "中（约 6 句）": "中等，正文约 300 字",
                "长（约 10 句）": "详细，正文 500 字以上",
            }
            max_sents = length_map[length_label]

            c1, c2 = st.columns(2)
            run_ext = c1.button("⚡ 生成抽取式摘要（内置引擎，离线）", use_container_width=True)
            run_llm = c2.button(
                "🤖 生成 LLM 深度总结" + ("" if llm_ready else "（需先在侧边栏配置）"),
                use_container_width=True,
                disabled=not llm_ready,
            )
        else:
            if source == "直接粘贴文本 / 摘要":
                st.info("👆 粘贴文本后即可生成摘要")

    if run_ext:
        try:
            with st.spinner("正在分析文本……"):
                with logger.span("抽取式摘要", 标题=chosen_title[:40], 档位=max_sents):
                    res = summarizer.extractive_summary(text, ratio=1.0, max_sentences=max_sents, title=chosen_title)
        except Exception as e:
            logger.error("抽取式摘要生成失败", e)
            st.error("摘要生成失败，问题已记录到日志。可稍后重试或换一篇文献。")
            with st.expander("技术详情"):
                st.code(str(e)[:500])
            st.stop()
        summary_out = res["summary"]
        translated = False
        if lang == "中文" and summarizer.is_mostly_english(summary_out):
            with st.spinner("检测到英文摘要，正在自动翻译为中文……"):
                try:
                    summary_out = summarizer.translate_text(summary_out, "en|zh-CN")
                    translated = True
                except Exception as e:
                    st.warning(f"自动翻译失败（{e}），已显示英文原句。也可配置 LLM 获得中文深度总结。")
        lang_tag = " · 中文翻译" if translated else (" · 原文" if lang == "中文" else "")
        with st.container(key="panel_ext"):
            st.markdown(f"#### 📄 摘要结果 — {chosen_title}（{length_label}{lang_tag}）")
            # 明示实际句数：摘要原文不足所选档位时只能全部纳入，避免"句子数与选择不符"的困惑
            # 用 .get 兜底：正在运行的服务若仍缓存旧版引擎模块，也不会 KeyError 崩溃
            src_n = res.get("source_count", 0)
            picked_n = res.get("picked_count", 0)
            lack = src_n < max_sents
            st.caption(
                f"实际输出 {picked_n} 句（原文共 {src_n} 句"
                + ("，原文句子数少于所选档位，已全部纳入）" if lack else "）")
            )
            st.markdown(summary_out)
            key_terms = res["key_terms"]
            if lang == "中文" and key_terms:
                with st.spinner("正在翻译关键词为中文……"):
                    key_terms = summarizer.translate_keywords(key_terms[:8])
            if key_terms:
                st.markdown("**🔑 关键词**")
                st.markdown("".join(f'<span class="kw-chip">{k}</span>' for k in key_terms[:8]), unsafe_allow_html=True)
            with st.expander("📊 句子重要性得分（Top 语句）"):
                top_sents = [s for s, _ in res["scores"][:max_sents]]
                top_scores = [sc for _, sc in res["scores"][:max_sents]]
                if lang == "中文" and top_sents and summarizer.is_mostly_english(" ".join(top_sents)):
                    with st.spinner("正在翻译关键句为中文……"):
                        try:
                            top_sents = summarizer.translate_sentences(top_sents)
                        except Exception:
                            st.caption("（关键句自动翻译失败，已显示英文原文）")
                for s, sc in zip(top_sents, top_scores):
                    st.markdown(f"`{sc}` {s[:160]}")
            with st.expander("🔢 关键数值原文定位", expanded=False):
                _render_value_locator(locate.build_index([{"title": chosen_title or "原文", "text": text}]))
            st.download_button("⬇️ 导出摘要 (Markdown)", summary_out, file_name="summary.md")

    if run_llm:
        try:
            with logger.span("LLM 深度总结", 模型=st.session_state.get("llm_model", ""), 标题=chosen_title[:40]):
                with st.spinner("LLM 正在生成深度总结……"):
                    out = summarizer.llm_summary(
                    text,
                    st.session_state["llm_base"],
                    st.session_state["llm_key"],
                    st.session_state["llm_model"],
                    language=lang,
                    length_hint=llm_length_map[length_label],
                )
            with st.container(key="panel_llm"):
                st.markdown(f"#### 🤖 LLM 深度总结 — {chosen_title}（{length_label}）")
                st.markdown(out)
                with st.expander("📍 原文定位（逐句溯源）", expanded=False):
                    _render_located_sentences(
                        locate.build_index([{"title": chosen_title or "全文", "text": text}]),
                        summarizer.split_sentences(out),
                    )
                st.download_button("⬇️ 导出总结 (Markdown)", out, file_name="llm_summary.md")
        except Exception as e:
            logger.error("LLM 调用失败", e)
            st.error(f"LLM 调用失败：{e}（请检查 API 地址 / Key / 模型名，以及网络连通性）")

    # ---------------- 全文摘要与数据分析（v1.2.0） ----------------
    if chosen_article and chosen_article.get("pmcid"):
        st.divider()
        sec_title("2️⃣ 全文摘要与数据分析", f"基于开放全文（{chosen_article['pmcid']}；PMC 不可用时自动转网页提取兜底）——引言 / 方法 / 结果 / 讨论章节化摘要与统计分析")
        ft_max = st.select_slider("全文摘要句子数", options=[6, 8, 10, 12, 15], value=8)
        if st.button("📚 抓取全文并生成全文摘要", type="primary"):
            # v1.7.0：后台任务执行，页面不阻塞，可与图表解析同时进行
            _pmcid = chosen_article["pmcid"]
            if jobs.is_running(f"ft:{_pmcid}"):
                st.warning("该文献的全文摘要任务正在进行中，进度见左侧「后台任务」。")
            else:
                jobs.start("全文摘要", _ft_job, key=f"ft:{_pmcid}",
                           article={"pmcid": _pmcid, "doi": chosen_article.get("doi") or ""},
                           title=chosen_title, max_sentences=ft_max)
                st.toast("全文抓取已在后台开始，页面可继续操作", icon="⏳")

        if st.session_state.get("ft_summary"):
            ft = st.session_state["ft_summary"]
            an = st.session_state["ft_analytics"]
            summary_text = ft["summary"]

            # 小标题按所选语言取用：中文用 title_zh，英文用原始章节标题（v1.5.1 修复英文摘要混入中文小标题）
            if lang == "中文" and ft.get("sections"):
                parts, failed = [], False
                for p in ft["sections"]:
                    body = " ".join(p["sentences"])
                    if summarizer.is_mostly_english(body):
                        try:
                            body = summarizer.translate_text(body, "en|zh-CN")
                        except Exception:
                            failed = True
                    parts.append(f"【{p.get('title_zh') or p['title']}】 {body}")
                ft_out = "\n\n".join(parts)
                if failed:
                    st.warning("部分章节翻译失败，对应小节已保留英文原句。")
            elif ft.get("sections"):
                ft_out = "\n\n".join(
                    f"【{p['title']}】 " + " ".join(p["sentences"]) for p in ft["sections"]
                )
            else:
                ft_out = summary_text
            if str(st.session_state.get("ft_source", "")).startswith("网页"):
                st.info("🌐 本文不在 PMC 开放存档，正文经网页提取兜底获得（Unpaywall OA 副本 / 出版社页面），章节划分可能与原文略有出入。")
            sec_title("📚 全文摘要", f"共 {ft.get('used_sections', 0)} 个章节纳入摘要")
            with st.container(key="panel_ft"):
                st.markdown(ft_out)

            # ---- 原文定位（v2.1.0）：摘要句子与关键数值回溯原文 ----
            secs_loc = st.session_state.get("fulltext") or []
            if secs_loc:
                _loc_index = locate.build_index(secs_loc)
                _loc_stat = locate.index_stats(_loc_index)
                sec_title(
                    "📍 原文定位",
                    f"摘要句子与关键数值回溯原文（索引 {_loc_stat['sentences']} 句 / {_loc_stat['sections']} 章节）",
                )
                _render_keyword_search(_loc_index)
                with st.expander("📍 摘要句子 → 原文出处", expanded=False):
                    for p in ft.get("sections", []):
                        st.markdown(f"**【{p.get('title_zh') or p.get('title', '')}】**")
                        _render_located_sentences(_loc_index, p.get("sentences", []), show_heading_tags=False)
                with st.expander("🔢 关键数值 → 原文出处", expanded=False):
                    _render_value_locator(_loc_index)

            # ---- 数据分析面板 ----
            sec_title("📊 全文数据分析", "词数统计 · 章节篇幅 · 高频关键词 · 统计指标")
            m1, m2, m3 = st.columns(3)
            with m1:
                st.metric("全文词数", f"{an['total_words']:,}")
            with m2:
                st.metric("全文字符", f"{an['total_chars']:,}")
            with m3:
                st.metric("章节数", len(an['section_stats']))

            st.markdown("**各章节篇幅分布**")
            chart_data = {s["title"][:30]: s["words"] for s in an["section_stats"]}
            st.bar_chart(chart_data)

            st.markdown("**🔑 全文高频关键词**（已过滤功能词与套话，括号内为出现次数）")
            kc = an.get("keyword_counts", {})
            ft_kws = an["keywords"]
            if lang == "中文" and ft_kws:
                with st.spinner("正在翻译关键词为中文……"):
                    ft_kws = summarizer.translate_keywords(ft_kws)
            st.markdown(
                "".join(
                    f'<span class="kw-chip">{k} <b>{kc[k]}</b></span>' if k in kc else f'<span class="kw-chip">{k}</span>'
                    for k in ft_kws
                ),
                unsafe_allow_html=True,
            )

            st.markdown("**🔬 关键统计指标提取**")
            if an["metrics"]:
                st.caption("「N 处」为全文出现总次数；下方展示去重后的不同取值（最多 12 种）。")
                for name, m in an["metrics"].items():
                    with st.expander(f"{name} — 全文 {m['count']} 处 · 去重 {len(m['samples'])} 种"):
                        st.markdown("".join(f'<span class="kw-chip">{v}</span>' for v in m["samples"]), unsafe_allow_html=True)
            else:
                st.caption("未在全文中提取到常见统计指标。")

            st.download_button("⬇️ 导出全文摘要 (Markdown)", ft_out, file_name="fulltext_summary.md")

    # ---------------- 文献图表解析（需要 PMC 开放全文） ----------------
    if chosen_article and chosen_article.get("pmcid"):
        st.divider()
        sec_title("3️⃣ 文献图表解析", f"图片视图 + 说明概括（{chosen_article['pmcid']}）")
        if st.button("🖼 加载并解析文献图表", type="secondary"):
            # v1.7.0：后台任务执行，可与全文摘要同时进行
            _pmcid_f = chosen_article["pmcid"]
            if jobs.is_running(f"fig:{_pmcid_f}"):
                st.warning("该文献的图表解析任务正在进行中，进度见左侧「后台任务」。")
            else:
                st.session_state["fig_error"] = ""
                jobs.start("图表解析", _fig_job, key=f"fig:{_pmcid_f}", pmcid=_pmcid_f)
                st.toast("图表抓取已在后台开始，页面可继续操作", icon="⏳")
        figures = st.session_state.get("figures", [])
        fig_err = st.session_state.pop("fig_error", "")
        if fig_err:
            st.error(f"❌ 图表解析失败：{fig_err}")
        if figures:
            n_shot = sum(1 for f in figures if f.get("is_screenshot"))
            st.success(f"共解析出 {len(figures)} 张图表" + (
                f"（其中 {n_shot} 张为浏览器截图——该文献非完全开放获取，原始图片包不可用，已自动用截图兜底）"
                if n_shot else ""))
            # 视图展示
            for i, f in enumerate(figures):
                with st.container(border=True):
                    st.markdown(f"**{f.get('label') or f'Figure {i + 1}'}**")
                    st.image(f["data"], use_container_width=True)
                    if f.get("caption"):
                        st.caption(f["caption"])
            # 概括：内置抽取式
            captions_text = " ".join(f.get("caption", "") for f in figures if f.get("caption"))
            if captions_text.strip():
                with st.spinner("正在概括图表信息……"):
                    fig_res = summarizer.extractive_summary(captions_text, ratio=0.5, max_sentences=5)
                with st.container(key="panel_figcap"):
                    sec_title("🧩 图表信息概括", "内置抽取式引擎，基于各图表说明文字")
                    fig_out = fig_res["summary"]
                    if lang == "中文" and summarizer.is_mostly_english(fig_out):
                        try:
                            fig_out = summarizer.translate_text(fig_out, "en|zh-CN")
                        except Exception:
                            st.caption("（图表概括自动翻译失败，已显示英文原文）")
                    st.markdown(fig_out)
            # 概括：LLM（可选）
            if llm_ready:
                if st.button("🤖 用 LLM 逐图概括图表信息"):
                    try:
                        with st.spinner("LLM 正在分析图表说明……"):
                            fig_llm = summarizer.llm_figure_summary(
                                figures,
                                st.session_state["llm_base"],
                                st.session_state["llm_key"],
                                st.session_state["llm_model"],
                                language=lang,
                            )
                        with st.container(key="panel_figllm"):
                            sec_title("🤖 图表信息概括", "LLM 逐图解读说明文字")
                            st.markdown(fig_llm)
                    except Exception as e:
                        st.error(f"LLM 调用失败：{e}")
                # v1.3.0：多模态视觉分析——直接把图表图片送 LLM"看图"解读
                if st.button("🔬 用 LLM 视觉分析图表图片（多模态模型如 deepseek-flash / gpt-4o）", use_container_width=True):
                    try:
                        with st.spinner("LLM 正在逐张查看并分析图表图片（图片较多时约需 1-2 分钟）……"):
                            fig_vision = summarizer.llm_figure_vision(
                                figures,
                                st.session_state["llm_base"],
                                st.session_state["llm_key"],
                                st.session_state["llm_model"],
                                language=lang,
                            )
                        with st.container(key="panel_vision"):
                            sec_title("🔬 图表图片视觉分析", "LLM 多模态逐张看图解读")
                            st.markdown(fig_vision)
                    except Exception as e:
                        st.error(
                            f"视觉分析失败：{e}（请确认所配置模型支持图片输入，"
                            "如 deepseek-flash / gpt-4o / qwen-vl 等；纯文本模型无法看图）"
                        )
        elif st.session_state.get("figures") == [] and not fig_err:
            st.info("该文献在 PMC 全文中未解析出图表，或抓取失败。")
    elif chosen_article:
        st.caption("💡 该文献暂无 PMC 开放全文，无法解析图表；可尝试选择带「📄 PDF 全文 (PMC)」链接的文献。")


# ---------------- 页面：PDF 全文分析（v3.2.0，P3-C3） ----------------
elif page == "PDF 全文分析":
    render_pdf_page()


# ---------------- 页面：综述工作台（v3.0.0） ----------------
elif page == "综述工作台":
    render_review_page()


# ---------------- 页面：批量全文（v1.7.0） ----------------
elif page == "批量全文":
    import pandas as pd

    header("📚 批量全文摘要", "多篇文献并行抓取 PMC 开放全文，批量生成章节化摘要与数据分析，一键导出")

    results_list = st.session_state.get("results", [])
    favs = storage.list_favorites()
    src_opts = ["当前检索结果"] + (["我的收藏"] if favs else [])
    src = st.radio("文献来源", src_opts, horizontal=True)
    pool = results_list if src == "当前检索结果" else favs
    pmc_pool = [a for a in pool if a.get("pmcid")]

    if not pmc_pool:
        st.info(
            "当前来源没有带 PMC 开放全文的文献。请先到「文献检索」页检索，"
            "或收藏带「📄 PDF 全文 (PMC)」标记的文献后再来。"
        )
    else:
        sel_df = pd.DataFrame({
            "选择": [False] * len(pmc_pool),
            "标题": [a["title"][:80] for a in pmc_pool],
            "期刊": [a.get("journal") or "" for a in pmc_pool],
            "年份": [a.get("year") or "" for a in pmc_pool],
            "PMCID": [a["pmcid"] for a in pmc_pool],
        })
        edited = st.data_editor(
            sel_df, hide_index=True, use_container_width=True, key="batch_sel",
            column_config={
                "选择": st.column_config.CheckboxColumn("选择", help="勾选要处理的文献", default=False),
                "标题": st.column_config.TextColumn("标题", width="large", disabled=True),
                "期刊": st.column_config.TextColumn("期刊", disabled=True),
                "年份": st.column_config.TextColumn("年份", disabled=True),
                "PMCID": st.column_config.TextColumn("PMCID", disabled=True),
            },
        )
        try:
            chosen_idx = [i for i, v in enumerate(edited["选择"].tolist()) if bool(v)]
        except Exception:
            chosen_idx = []
        selected = [pmc_pool[i] for i in chosen_idx]

        p1, p2, p3 = st.columns(3)
        with p1:
            ft_max = st.select_slider("每篇摘要句子数", options=[6, 8, 10, 12, 15], value=8)
        with p2:
            workers = st.slider("并行线程数", 1, 6, 3, help="NCBI 有限流（约 3 请求/秒），建议 3-4")
        with p3:
            batch_lang = st.radio("输出语言", ["中文", "英文"], horizontal=True)

        if st.button(
            f"🚀 并行抓取并生成摘要（已选 {len(selected)} 篇）",
            type="primary", disabled=not selected, use_container_width=True,
        ):
            if jobs.is_running("batch"):
                st.warning("已有批量任务在进行中，请等待完成（进度见左侧「后台任务」）。")
            else:
                articles = [
                    {k: a.get(k) for k in ("pmid", "pmcid", "title", "journal", "year", "authors", "doi")}
                    for a in selected
                ]
                jobs.start(f"批量全文（{len(articles)} 篇）", _batch_job, key="batch",
                           articles=articles, max_sentences=ft_max, workers=workers, lang=batch_lang)
                st.toast(f"批量任务已开始：{len(articles)} 篇 × {workers} 线程，页面可继续操作", icon="🚀")

    @st.fragment(run_every=1.5)
    def _batch_progress():
        running = [j for j in jobs.snapshot() if j["status"] == "running" and j.get("key") == "batch"]
        if running:
            j = running[0]
            st.info(f"⏳ {j['name']} — {j['note']}")
            st.progress(max(j["progress"], 0.03))

    _batch_progress()

    batch_results = st.session_state.get("batch_results")
    if batch_results:
        batch_lang_out = st.session_state.get("batch_lang", "中文")
        ok_n = sum(1 for r in batch_results if r.get("ok"))
        st.divider()
        sec_title("📦 批量结果", f"共 {len(batch_results)} 篇 · 成功 {ok_n} · 失败 {len(batch_results) - ok_n}")

        md_text = _batch_export_md(batch_results, batch_lang_out)
        xlsx_bytes = _batch_export_xlsx(batch_results, batch_lang_out)
        d1, d2 = st.columns(2)
        with d1:
            st.download_button(
                "⬇️ 一键导出全部摘要与分析 (Markdown)", md_text,
                file_name="batch_fulltext_summary.md", mime="text/markdown",
                type="primary", use_container_width=True,
            )
        with d2:
            st.download_button(
                "⬇️ 导出统计表 (Excel：汇总 / 章节明细 / 统计指标)", xlsx_bytes,
                file_name="batch_fulltext_stats.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

        for i, r in enumerate(batch_results, 1):
            a = r["article"]
            src_note = f"　·　🌐 {r['source']}" if r.get("ok") and str(r.get("source", "")).startswith("网页") else ""
            with st.expander(("✅ " if r.get("ok") else "❌ ") + f"{i}. {a['title']}" + src_note):
                if not r.get("ok"):
                    st.error(f"抓取失败：{r.get('error', '')}")
                    continue
                summ, an = r["summary"], r["analytics"]
                if summ.get("sections"):
                    if batch_lang_out == "中文":
                        for p in summ["sections"]:
                            body = p.get("body_zh") or " ".join(p["sentences"])
                            st.markdown(f"**【{p.get('title_zh') or p['title']}】** {body}")
                    else:
                        for p in summ["sections"]:
                            st.markdown(f"**【{p['title']}】** " + " ".join(p["sentences"]))
                elif summ.get("summary"):
                    st.markdown(summ["summary"])
                m1, m2, m3 = st.columns(3)
                with m1:
                    st.metric("全文词数", f"{an['total_words']:,}")
                with m2:
                    st.metric("章节数", len(an["section_stats"]))
                with m3:
                    st.metric("统计指标类", len(an.get("metrics") or {}))
                kc = an.get("keyword_counts", {})
                kws = an.get("keywords", [])[:8]
                if kws:
                    st.markdown(
                        "".join(
                            f'<span class="kw-chip">{k} <b>{kc[k]}</b></span>' if k in kc else f'<span class="kw-chip">{k}</span>'
                            for k in kws
                        ),
                        unsafe_allow_html=True,
                    )


elif page == "我的文献库":
    render_library_page()


# ---------------- 页面：隐私与数据 ----------------
elif page == "隐私与数据":
    header("🔒 隐私与数据", "本工具的数据处理方式")
    sec_title("隐私政策与数据处理说明",
              f"版本 {APP_VERSION} · 生效日期 2026-10-07 · 本页内容与仓库 README 的「隐私」章节一致")

    st.markdown(
        """
本工具是一款开源的医学文献检索与摘要辅助工具，以本地/会话内存储为主，
不设用户账号体系。以下逐条说明数据如何被处理。

---

#### 一、适用范围

本政策适用于本工具的全部形态：Streamlit Cloud 在线版、Windows 桌面版，
以及用户自行部署的实例。桌面版不与本项目作者共享任何数据；
在线版的运行日志仅用于故障排查，不会用于识别使用者身份。

#### 二、我们不收集的信息

本工具**不收集、不存储、不传输**下列信息：

- 姓名、身份证号、手机号、邮箱地址等直接身份信息
- 用户画像、行为标签、广告标识
- 任何用于跨会话追踪的技术标识

项目代码中不存在第三方统计 SDK、广告脚本或用户行为追踪逻辑。

#### 三、本地与会话数据

| 数据类型 | 存放位置 | 可见范围 | 保留期限 |
|---|---|---|---|
| 检索历史 | 桌面版：本机 `data/history.json`；在线版：`data/users/<会话>/` | 仅本人 | 至用户主动清除 |
| 文献收藏 | 桌面版：本机 `data/favorites.json`；在线版：同上 | 仅本人 | 至用户主动删除 |
| 分组 / 标签 / 笔记 | 桌面版：本机 `data/library.json`；在线版：`data/users/<会话>/library.json` | 仅本人 | 至用户主动删除 |
| 反馈留档 | 本机 `data/feedback/` | 仅本人 | 至用户主动删除 |
| 上传的 PDF | **不落盘**：仅在你浏览器的本次会话内存中解析，关闭页面即释放 | 仅本人 | 会话结束即消失 |

在线版的会话标识由服务端随机生成，仅用于区分不同访问者的数据目录，
不可反向推导身份，且随会话结束失效。

#### 四、结果缓存

为避免对 PubMed 公共接口与翻译额度造成重复消耗，本工具在服务器本地保留一份
**结果缓存**，涵盖检索结果、PMC 全文、抽取摘要、机器译文与大模型摘要。

- 缓存内容**全部来源于 PubMed 公开文献**，不含任何用户身份信息；
- 缓存为**全体用户共享**，以提高命中率并降低上游接口压力；
- 每类缓存设有保留期与条数上限，超出后按最久未用自动淘汰；
- 用户可在侧边栏「🩺 运行诊断」中一键清空缓存。

#### 五、第三方服务

以下操作会**由用户主动触发**并将相应文本发送至第三方服务提供方：

| 功能 | 第三方 | 发送的内容 |
|---|---|---|
| 文献检索 / 全文获取 | NCBI（PubMed / PMC） | 检索式、PMCID、DOI |
| 机器翻译 | 腾讯云机器翻译（或 MyMemory 兜底） | 待翻译的英文摘要文本 |
| 智能摘要 | 用户自填的大模型服务 | 待总结的正文或摘要文本 |

**用户应自行避免在上述输入框中粘贴患者姓名、住院号、身份证号等身份信息。**
本工具不对用户输入内容做身份信息识别与脱敏，该责任由使用者承担。

#### 六、API 凭据

用户填写的腾讯云密钥、NCBI API Key、大模型 API Key 等凭据，
**仅保存在运行进程的内存中**，不写入磁盘文件、不写入运行日志、
不随会话持久化保存。进程重启后需重新填写。

#### 七、运行日志

为便于故障定位，应用会在服务器本地记录运行日志与外部依赖健康统计，
按天滚动，保留 14 天。日志包含功能名称、耗时、请求次数与异常堆栈，
**不包含**用户检索内容与凭据。

#### 八、医疗免责声明

本工具生成的摘要、翻译与数据分析均由算法或大模型自动生成，
可能存在遗漏、偏差或曲解，**不得作为临床诊断、用药或治疗方案的依据**。
使用者应核对文献原文，并以专业医师的判断为准。

#### 九、数据来源与使用限制

文献元数据来自 PubMed / PMC（美国国立卫生研究院 NCBI 下属国家医学图书馆 NLM），
遵循 NCBI 使用条款，仅供学习与研究使用，不得用于批量抓取或商业用途。

#### 十、你的权利

由于本工具不设账号体系、不收集身份信息，**不存在需要向我们申请查询、更正或删除的个人数据**。
你随时可以在侧边栏「🩺 运行诊断」清空缓存、在「我的文献库」「检索历史」中删除记录、
在 `data/` 目录下直接删除全部本地数据。

#### 十一、政策更新

本政策如有变更，会在应用「📋 更新日志」中记录，并在本页顶部更新版本号与生效日期。
本政策随项目源码一同公开，可随时查阅历史版本。
"""
    )

    st.divider()
    st.caption("本页内容与仓库 README 的「隐私」章节保持一致；如发现不一致，以本页为准。")

    st.divider()
    st.markdown("#### 需要反馈？")
    st.caption("如对隐私处理有疑问或发现不一致，请在左侧「💬 报错 / 建议」中提交。")
    if st.button("💬 去反馈", key="privacy_feedback"):
        st.session_state["pending_page"] = "系统首页"
        st.rerun()


# ---------------- 页面：更新日志 ----------------
elif page == "更新日志":
    header("📋 更新日志", "各版本新增与修复内容一览")
    sec_title("更新日志", f"当前版本 {APP_VERSION} · 最新改动在页面顶部，完整说明见「使用说明.md」")
    st.write("")
    render_changelog()


# ---------------- 页面：检索历史 ----------------
elif page == "检索历史":
    header("🕘 检索历史", "回溯过往检索式与结果数量，可一键复用检索条件")
    hist = storage.list_history()
    if not hist:
        st.info("暂无检索历史。")
    else:
        sec_title("检索记录", f"共 {len(hist)} 次 · 点击「重新检索」可直接复用该检索式")
        if st.button("🧹 清空历史"):
            storage.clear_history()
            st.rerun()
        for h in hist:
            with st.container(border=True):
                c1, c2 = st.columns([4, 1])
                with c1:
                    st.markdown(f"**{h['query']}**")
                    st.caption(f"{h['time']} · {h['n_results']} 条结果")
                with c2:
                    if st.button("重新检索", key=f"re_{h['time']}_{h['query'][:10]}"):
                        st.session_state["kw"] = h["query"]
                        st.switch_page("app.py") if False else None
                        st.session_state["page_hint"] = "search"
                        st.toast("已填入关键词，请前往「文献检索」页", icon="🔍")


# ---------------- 全局页脚（v2.4.0）：每页底部常驻 ----------------
render_footer()
